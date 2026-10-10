"""A03–A18 service/API acceptance checks; PostgreSQL concurrency is a separate gate."""

import copy
import io
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.commerce.calculator import (
    convert,
    dec,
    distribute,
    expression,
    validate_cas,
)
from app.commerce.models import Quote
from app.core.errors import DomainError
from app.core.models import AuditEvent
from app.crm.models import Counterparty, Request
from tests.support.commerce import approve, command, fund, get, prepare, wave


def test_a05_cas_and_units():
    assert validate_cas("64-17-5") == "64-17-5"
    assert validate_cas(None, "Смесь") is None
    for value in ("64-17-4", "000-12-3", "abc"):
        with pytest.raises(DomainError):
            validate_cas(value)
    with pytest.raises(DomainError):
        validate_cas(None)
    assert convert(Decimal("1000"), "g", "kg") == Decimal("1")
    with pytest.raises(DomainError):
        convert(Decimal("1000"), "g", "l")


@pytest.mark.parametrize(
    "source",
    [
        "__import__('os').system('id')",
        "x.__class__",
        "sum([1])",
        "1 ** 9999999",
        "[x for x in y]",
        "True",
        "{'a':1}",
    ],
)
def test_untrusted_formula_cannot_execute(source):
    with pytest.raises(DomainError):
        expression(source, {"x": Decimal("1")})


def test_a08_rounding_distribution_and_no_binary_float():
    assert expression("0.1 + 0.2", {}) == Decimal("0.3")
    with pytest.raises(DomainError):
        dec(0.1)
    assert distribute(
        Decimal("1.00"), {"b": Decimal("1"), "a": Decimal("1"), "c": Decimal("1")}, Decimal("0.01"), "half_up"
    ) == {"a": Decimal("0.33"), "b": Decimal("0.34"), "c": Decimal("0.33")}
    with pytest.raises(DomainError):
        distribute(Decimal("1"), {"a": Decimal("0")}, Decimal("0.01"), "half_up")


def test_a11_partial_amount_rounding():
    bases = {"p1": Decimal("1"), "p2": Decimal("1"), "p3": Decimal("1"), "p4": Decimal("1")}
    result = distribute(Decimal("0.02"), bases, Decimal("0.01"), "half_up")
    assert result == {
        "p1": Decimal("0.01"),
        "p2": Decimal("0.00"),
        "p3": Decimal("0.01"),
        "p4": Decimal("0.00"),
    }


def test_a03_rfq_saved_file_not_sent_and_omits_finances(commerce):
    env = commerce
    rfq = command(
        env,
        f"/requests/{env['request_id']}/rfqs",
        {
            "supplier_id": env["supplier_id"],
            "item_ids": [env["item_id"]],
            "response_due": date.today().isoformat(),
        },
    )
    content = env["client"].get(f"/api/v1/rfqs/{rfq['id']}/file")
    assert content.status_code == 200
    book = load_workbook(io.BytesIO(content.content))
    assert book.active.cell(1, 3).value == "CAS"
    assert book.active.cell(2, 3).value == "64-17-5"
    assert book.active.cell(2, 5).value in (None, "")
    assert "цена" not in " ".join(str(cell.value) for cell in book.active[1]).lower()
    assert get(env, f"/requests/{env['request_id']}/rfqs")["items"][0]["sent_at"] is None
    assert env["client"].get(f"/api/v1/rfqs/{rfq['id']}/file").content == content.content


def test_a04_product_variant_deduplication(commerce):
    env = commerce
    body = {
        "name": "Этанол",
        "cas": "64-17-5",
        "manufacturer": "Фабрика",
        "purity": "99%",
        "packaging": "25 г",
        "unit": "g",
        "package_quantity": "25",
    }
    first = env["client"].post("/api/v1/catalog/products", json=body).json()
    second = (
        env["client"].post("/api/v1/catalog/products", json={**body, "manufacturer": " фабрика  "}).json()
    )
    third = (
        env["client"]
        .post("/api/v1/catalog/products", json={**body, "packaging": "100 г", "package_quantity": "100"})
        .json()
    )
    assert first["id"] == second["id"]
    assert third["id"] != first["id"]
    assert not first["verified"]


def test_a08_a09_a10_snapshot_arithmetic_idempotency_and_conflict(commerce):
    env = commerce
    state = prepare(env)
    result = state["calc"]["snapshot"]
    assert result["totals"]["total"] == "3593.75"
    detail = result["lines"][0]["detail"]
    assert Decimal(detail["cost"]) == Decimal("2500")
    assert Decimal(detail["cash_need"]) == Decimal("2731")
    key = str(uuid4())
    body = {**state["calculation_payload"], "idempotency_key": key}
    first = command(env, f"/requests/{env['request_id']}/calculations", body)
    second = command(env, f"/requests/{env['request_id']}/calculations", body)
    assert first["id"] == second["id"]
    assert (
        command(
            env, f"/requests/{env['request_id']}/calculations", {**body, "reason": "Изменение"}, expected=409
        )["code"]
        == "IDEMPOTENCY_CONFLICT"
    )
    command(
        env,
        f"/quotes/{state['quote']['id']}/revise",
        {**state["quote_payload"], "price": "200", "revision_reason": "Новая цена поставщика"},
    )
    saved = get(env, f"/requests/{env['request_id']}/calculations")["items"]
    assert all(row["snapshot"]["totals"]["total"] == "3593.75" for row in saved)
    assert (
        command(
            env, f"/requests/{env['request_id']}/calculations", state["calculation_payload"], expected=409
        )["code"]
        == "QUOTE_REVISION_CONFLICT"
    )


def test_a06_disallows_mixed_supplier_currency(commerce):
    env = commerce
    state = prepare(env)
    second = command(
        env, f"/requests/{env['request_id']}/quotes", {**state["quote_payload"], "currency": "RUB"}
    )
    payload = copy.deepcopy(state["calculation_payload"])
    payload["selections"].append(
        {"quote_id": second["id"], "quote_revision": 1, "quantity": "1", "unit": "pcs"}
    )
    assert (
        command(env, f"/requests/{env['request_id']}/calculations/preview", payload, expected=422)["code"]
        == "QUOTE_GROUP"
    )


def test_a07_expired_quotes_and_analogue_consent(commerce):
    env = commerce
    state = prepare(env)
    expired = command(
        env,
        f"/requests/{env['request_id']}/quotes",
        {**state["quote_payload"], "valid_until": (date.today() - timedelta(days=1)).isoformat()},
    )
    payload = copy.deepcopy(state["calculation_payload"])
    payload["selections"][0]["quote_id"] = expired["id"]
    assert (
        command(env, f"/requests/{env['request_id']}/calculations/preview", payload, expected=422)["code"]
        == "QUOTE_EXPIRED"
    )
    with env["sessions"]() as db:
        quote = db.get(Quote, state["quote"]["id"])
        quote.is_analogue = True
        db.commit()
    assert (
        command(
            env,
            f"/proposals/{state['proposal']['id']}/accept",
            {
                "version": 1,
                "lines": [{"line_id": state["quote"]["id"], "quantity": "1"}],
                "reason": "Частичное принятие",
            },
            expected=422,
        )["code"]
        == "ANALOGUE_APPROVAL"
    )


def test_a11_a12_partial_acceptance_and_invoice_snapshot(commerce):
    env = commerce
    state = prepare(env)
    accepted = command(
        env,
        f"/proposals/{state['proposal']['id']}/accept",
        {
            "version": 1,
            "lines": [{"line_id": state["quote"]["id"], "quantity": "6"}],
            "reason": "Клиент принимает часть",
        },
    )
    assert (
        command(
            env,
            f"/proposals/{state['proposal']['id']}/accept",
            {
                "version": 2,
                "lines": [{"line_id": state["quote"]["id"], "quantity": "5"}],
                "reason": "Превышение потребности",
            },
            expected=422,
        )["code"]
        == "ACCEPTANCE_EXCEEDED"
    )
    execution_id = accepted["executions"][0]["id"]
    invoice_body = {
        "proposal_id": state["proposal"]["id"],
        "lines": [{"execution_id": execution_id, "quantity": "6"}],
        "due_date": date.today().isoformat(),
        "terms": "Оплата части",
    }
    invoice = command(env, f"/requests/{env['request_id']}/invoices", invoice_body)
    assert Decimal(invoice["total"]) == Decimal("2156.25")
    assert (
        command(env, f"/requests/{env['request_id']}/invoices", invoice_body, expected=422)["code"]
        == "QUANTITY_EXCEEDED"
    )
    original = env["client"].get(f"/api/v1/documents/{invoice['id']}/file").content
    assert original.startswith(b"%PDF")
    with env["sessions"]() as db:
        client = db.get(Counterparty, env["client_id"])
        client.name = "Изменённый клиент"
        db.commit()
    assert env["client"].get(f"/api/v1/documents/{invoice['id']}/file").content == original


def test_a13_a14_payments_partial_and_unallocated_overpayment(commerce):
    env = commerce
    state = prepare(env, with_invoice=True)
    body = {
        "invoice_id": state["invoice"]["id"],
        "amount": "4000",
        "currency": "RUB",
        "payment_date": date.today().isoformat(),
        "number": "PAY-001",
        "external_id": "bank-001",
    }
    payment = command(env, f"/requests/{env['request_id']}/payments", body)
    assert get(env, f"/requests/{env['request_id']}/invoices")["items"][0]["payment_status"] == "unpaid"
    assert (
        command(
            env,
            f"/payments/{payment['id']}/allocate",
            {"version": 1, "allocations": [{"invoice_id": state["invoice"]["id"], "amount": "1000"}]},
            expected=422,
        )["code"]
        == "PAYMENT_UNCONFIRMED"
    )
    command(env, f"/payments/{payment['id']}/confirm", {"version": 1})
    payment = command(
        env,
        f"/payments/{payment['id']}/allocate",
        {"version": 2, "allocations": [{"invoice_id": state["invoice"]["id"], "amount": "1000"}]},
    )
    assert get(env, f"/requests/{env['request_id']}/invoices")["items"][0]["payment_status"] == "partial"
    assert (
        command(
            env,
            f"/payments/{payment['id']}/allocate",
            {"version": 3, "allocations": [{"invoice_id": state["invoice"]["id"], "amount": "3000"}]},
            expected=422,
        )["code"]
        == "INVOICE_OVERPAYMENT"
    )
    payment = command(
        env,
        f"/payments/{payment['id']}/allocate",
        {"version": 3, "allocations": [{"invoice_id": state["invoice"]["id"], "amount": "2593.75"}]},
    )
    assert Decimal(payment["unallocated"]) == Decimal("406.25")
    assert (
        command(env, f"/requests/{env['request_id']}/payments", body, expected=409)["code"]
        == "DUPLICATE_PAYMENT"
    )


def test_a15_a16_a17_a18_approval_reversal_and_partial_delivery(commerce):
    env = commerce
    state = prepare(env, with_invoice=True)
    unpaid = command(env, f"/requests/{env['request_id']}/approvals",
        {"execution_ids": [state["execution"]["id"]], "reviewer_id": env["owner_id"]})
    assert Decimal(unpaid['snapshot']['lines'][0]['funding']['ratio']) == 0
    payment = fund(env, state)
    first_approval = unpaid
    command(
        env,
        f"/approvals/{first_approval['id']}/decision",
        {"version": 1, "decision": "returned", "reason": "Проверить сроки"},
    )
    approval = approve(env, state)
    first, second, third = wave(env, "03"), wave(env, "04"), wave(env, "05")
    six = command(
        env,
        f"/waves/{first['id']}/allocations",
        {"execution_id": state["execution"]["id"], "approval_id": approval["id"], "quantity": "6"},
    )
    four = command(
        env,
        f"/waves/{second['id']}/allocations",
        {"execution_id": state["execution"]["id"], "approval_id": approval["id"], "quantity": "4"},
    )
    assert (
        command(
            env,
            f"/waves/{third['id']}/allocations",
            {"execution_id": state["execution"]["id"], "approval_id": approval["id"], "quantity": "1"},
            expected=422,
        )["code"]
        == "WAVE_OVERALLOCATED"
    )
    for kind in ("shipped", "arrived", "delivered"):
        command(
            env,
            f"/allocations/{six['id']}/events",
            {
                "kind": kind,
                "quantity": "6",
                "occurred_at": datetime.now(timezone.utc).isoformat(),
                "reason": "Подтверждена тестовая партия",
            },
        )
    assert (
        command(
            env,
            f"/allocations/{six['id']}/transfer",
            {"version": 1, "target_wave_id": third["id"], "quantity": "1", "reason": "Недопустимый перенос"},
            expected=422,
        )["code"]
        == "TRANSFER_SHIPPED"
    )
    command(
        env,
        f"/allocations/{four['id']}/transfer",
        {
            "version": 1,
            "target_wave_id": third["id"],
            "quantity": "4",
            "reason": "Плановая партия перенесена",
        },
    )
    execution = get(env, f"/requests/{env['request_id']}/executions")["items"][0]
    assert execution["state"] == "partial"
    assert Decimal(execution["delivered"]) == Decimal("6")
    assert Decimal(execution["unallocated"]) == 0
    payment = command(
        env,
        f"/payments/{payment['id']}/reverse",
        {
            "version": 3,
            "amount": "100",
            "allocation_id": payment["allocations"][0]["id"],
            "kind": "unallocate",
            "reason": "Исправление распределения",
        },
    )
    execution = get(env, f"/requests/{env['request_id']}/executions")["items"][0]
    assert execution["financing_deficit"] is True
    assert Decimal(execution["delivered"]) == 6


def test_a20_no_financial_fields_or_idor(commerce):
    env = commerce
    state = prepare(env)
    env["user_id"] = env["other_id"]
    get(env, f"/requests/{env['request_id']}/quotes", expected=404)
    assert env["client"].get(f"/api/v1/documents/{state['proposal']['id']}/file").status_code == 404
    with env["sessions"]() as db:
        req = db.get(Request, env["request_id"])
        req.owner_id = env["other_id"]
        db.commit()
    quote = get(env, f"/requests/{env['request_id']}/quotes")["items"][0]
    assert "price" not in quote
    snapshot = get(env, f"/requests/{env['request_id']}/calculations")["items"][0]["snapshot"]
    assert "profile" not in snapshot and "input" not in snapshot
    assert "detail" not in snapshot["lines"][0] and "quote" not in snapshot["lines"][0]
    with env["sessions"]() as db:
        assert db.scalar(select(AuditEvent.id))


def test_financial_permissions_revocation_and_owner_change(commerce):
    env = commerce
    from uuid import uuid4

    prepare(env, with_invoice=False)

    str(uuid4())

    calc = env["client"].get(f"/api/v1/requests/{env['request_id']}/calculations").json()["items"][0]
    assert "cost" in calc["snapshot"]["lines"][0]["detail"]

    with env["sessions"].begin() as db:
        from app.core.models import PermissionGrant

        db.query(PermissionGrant).filter_by(user_id=env["owner_id"], code="finance.profit.read").delete()
        db.flush()

    replayed = env["client"].get(f"/api/v1/requests/{env['request_id']}/calculations").json()["items"][0]
    assert "cost" not in replayed["snapshot"]["lines"][0]["detail"]
    assert "profit" not in replayed["snapshot"]["lines"][0]["detail"]

    profile_view = env["client"].get("/api/v1/profiles").json()["items"][-1]
    assert "formulas" not in profile_view["definition"]

    with env["sessions"].begin() as db:
        from app.core.models import PermissionGrant
        from app.crm.models import Request as CRMRequest

        db.query(PermissionGrant).filter_by(user_id=env["owner_id"], code="requests.read").delete()
        db.add(PermissionGrant(user_id=env["owner_id"], code="requests.read", scope="own"))
        req = db.get(CRMRequest, env["request_id"])
        req.owner_id = env["other_id"]
        db.flush()

    response = env["client"].get(f"/api/v1/requests/{env['request_id']}/calculations")
    assert response.status_code == 404


def test_r05_partial_amount_small_rounding():
    from decimal import Decimal

    from app.commerce.documents import partial_amount

    # 0.02 / 4 problem:
    original = {"quantity": "4", "net": "0.02", "tax": "0.01", "total": "0.03"}

    profile = {"currency_precision": {"RUB": 2}, "sale_currency": "RUB", "rounding": "half_up"}

    # Issue piece 1 (qty 1)
    # 0.02 * 1 / 4 = 0.005 -> 0.01 (half_up)
    p1 = partial_amount(original, Decimal("1"), [], profile)
    assert p1["net"] == "0.01"
    assert p1["tax"] == "0.00"
    assert p1["total"] == "0.01"

    # Issue piece 2 (qty 1)
    # 0.02 * 1 / 4 = 0.005 -> 0.01
    p2 = partial_amount(original, Decimal("1"), [p1], profile)
    assert p2["net"] == "0.01"
    assert p2["tax"] == "0.00"
    assert p2["total"] == "0.01"

    # Now remaining net is 0.00

    # Issue piece 3 (qty 1)
    p3 = partial_amount(original, Decimal("1"), [p1, p2], profile)
    assert p3["net"] == "0.00"  # Capped because 0.02 - 0.02 = 0
    assert p3["tax"] == "0.00"

    # Issue piece 4 (qty 1) - Final
    p4 = partial_amount(original, Decimal("1"), [p1, p2, p3], profile)
    assert p4["net"] == "0.00"  # Should not be negative!
    assert p4["tax"] == "0.01"  # Because original tax was 0.01, sum(previous) was 0.00

    assert Decimal(p1["net"]) + Decimal(p2["net"]) + Decimal(p3["net"]) + Decimal(p4["net"]) == Decimal(
        "0.02"
    )
    assert Decimal(p1["tax"]) + Decimal(p2["tax"]) + Decimal(p3["tax"]) + Decimal(p4["tax"]) == Decimal(
        "0.01"
    )
    assert Decimal(p1["total"]) + Decimal(p2["total"]) + Decimal(p3["total"]) + Decimal(
        p4["total"]
    ) == Decimal("0.03")

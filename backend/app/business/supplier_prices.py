"""The base of supplier prices gathered from quotes: to look through and to take away as a spreadsheet.

Business remark of 21 October 2026. Every current quote position of the period is one row: supplier, product,
packing, price in its currency and in rubles by the published profile, delivery term and the request it came
for. A second sheet compares suppliers: the best ruble price of each product and packing.
"""

import io
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Any
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select

from app.commerce.rates import in_rubles, profile_rates
from app.core.config import settings
from app.core.db import DB, as_utc
from app.core.security import Actor, request_predicate, require_permission
from app.core.service import audit, plain, prefetch
from app.crm.models import (
    Counterparty,
    Currency,
    Nomenclature,
    Packing,
    ProductGroup,
    QuoteItem,
    QuoteSheet,
    Request,
    RequestItem,
)

from .supplier_stats import period

router = APIRouter(tags=["Аналитика"])
EXPORT_LIMIT = 100000
INK = "173F35"
COLUMNS = [
    ("quoted_at", "Дата квоты", 12), ("quote_number", "Квота", 11), ("request_number", "Заявка", 9),
    ("client_name", "Клиент", 28), ("supplier_name", "Поставщик", 30), ("supplier_code", "Код поставщика", 10),
    ("nomenclature_name", "Номенклатура", 40), ("article", "Артикул", 18), ("manufacturer", "Производитель", 18),
    ("cas", "CAS", 13), ("product_group", "Товарная группа", 18), ("packing", "Фасовка", 10),
    ("quantity", "Кол-во", 8), ("unit_price", "Цена", 12), ("currency", "Валюта", 8),
    ("unit_price_rub", "Цена, ₽", 13), ("delivery_days", "Срок, дней", 9), ("valid_until", "Действует до", 12),
    ("position", "Позиция заявки", 34),
]
BEST_COLUMNS = [
    ("nomenclature_name", "Номенклатура", 40), ("packing", "Фасовка", 10), ("manufacturer", "Производитель", 18),
    ("cas", "CAS", 13), ("product_group", "Товарная группа", 18), ("unit_price_rub", "Лучшая цена, ₽", 14),
    ("supplier_name", "Поставщик", 30), ("unit_price", "Цена", 12), ("currency", "Валюта", 8),
    ("delivery_days", "Срок, дней", 9), ("quoted_at", "Дата квоты", 12), ("article", "Артикул", 18),
    ("highest_rub", "Самая высокая, ₽", 14), ("offers", "Предложений", 11), ("suppliers", "Поставщиков", 11),
]


def price_query(db, user, date_from: date | None, date_to: date | None, q: str, supplier_id: str | None):
    """Current quote positions of the period on requests whose purchase prices the user may see."""
    start, finish = period(date_from, date_to)
    visible = select(Request.id).where(request_predicate(db, user, "finance.purchase.read"),
                                       Request.is_test.is_(False))
    query = (select(QuoteItem).join(QuoteSheet, QuoteSheet.id == QuoteItem.quote_id)
             .join(Nomenclature, Nomenclature.id == QuoteItem.nomenclature_id)
             .where(QuoteSheet.request_id.in_(visible), QuoteSheet.archived.is_(False),
                    QuoteItem.archived.is_(False), QuoteItem.quoted_at >= start, QuoteItem.quoted_at < finish))
    if supplier_id:
        query = query.where(QuoteItem.supplier_id == supplier_id)
    if text := q.strip():
        pattern = f"%{text}%"
        suppliers = select(Counterparty.id).where(Counterparty.name.ilike(pattern))
        query = query.where(or_(Nomenclature.name.ilike(pattern), Nomenclature.article.ilike(pattern),
                                Nomenclature.cas.ilike(pattern), Nomenclature.manufacturer.ilike(pattern),
                                QuoteItem.supplier_id.in_(suppliers)))
    return query.order_by(QuoteItem.quoted_at.desc(), QuoteItem.id)


def local_day(moment) -> date:
    return as_utc(moment).astimezone(ZoneInfo(settings.company_timezone)).date()


def price_rows(db, rows: list[QuoteItem]) -> tuple[list[dict[str, Any]], str | None]:
    rates = profile_rates(db)
    sheets = prefetch(db, QuoteSheet, (row.quote_id for row in rows))
    requests = prefetch(db, Request, (sheet.request_id for sheet in sheets.values()))
    clients = prefetch(db, Counterparty, (request.client_id for request in requests.values()))
    suppliers = prefetch(db, Counterparty, (row.supplier_id for row in rows))
    products = prefetch(db, Nomenclature, (row.nomenclature_id for row in rows))
    groups = prefetch(db, ProductGroup, (product.product_group_id for product in products.values()))
    packings = prefetch(db, Packing, (row.packing_id for row in rows))
    currencies = prefetch(db, Currency, (row.currency_id for row in rows))
    positions = prefetch(db, RequestItem, (row.source_request_item_id for row in rows))
    result = []
    for row in rows:
        sheet = sheets[row.quote_id]
        request = requests[sheet.request_id]
        product = products[row.nomenclature_id]
        group = groups.get(product.product_group_id) if product.product_group_id else None
        currency = currencies[row.currency_id].code
        position = positions.get(row.source_request_item_id) if row.source_request_item_id else None
        result.append({
            "id": row.id, "quoted_at": local_day(row.quoted_at), "valid_until": local_day(row.valid_until),
            "quote_number": sheet.number, "request_id": request.id, "request_number": request.number,
            "client_name": clients[request.client_id].name, "supplier_name": suppliers[row.supplier_id].name,
            "supplier_code": suppliers[row.supplier_id].internal_code, "nomenclature_name": product.name,
            "article": product.article, "manufacturer": product.manufacturer, "cas": product.cas,
            "product_group": group.name if group else None, "packing": packings[row.packing_id].display_name,
            "quantity": row.quantity, "unit_price": row.unit_price, "currency": currency,
            "unit_price_rub": in_rubles(row.unit_price, currency, rates["rates"]),
            "delivery_days": row.delivery_days,
            "position": position.description if position and not position.quote_only else None,
        })
    return result, rates["profile"]


def product_key(row: dict[str, Any]) -> tuple[str, str]:
    """One product and packing across suppliers: names compared by their words, as suppliers spell them."""
    words = re.findall(r"[0-9a-zа-яё]+", row["nomenclature_name"].casefold())
    return " ".join(words), "".join(re.findall(r"[0-9a-zа-яё.]+", row["packing"].casefold()))


def best_prices(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """The lowest ruble price of each product and packing, with how many offers and suppliers stand behind it."""
    offers: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if row["unit_price_rub"] is not None:
            offers[product_key(row)].append(row)
    result = []
    for group in offers.values():
        best = min(group, key=lambda row: (Decimal(row["unit_price_rub"]), row["delivery_days"] or 0))
        result.append({**best, "highest_rub": max(Decimal(row["unit_price_rub"]) for row in group),
                       "offers": len(group), "suppliers": len({row["supplier_name"] for row in group})})
    return sorted(result, key=lambda row: (row["nomenclature_name"].casefold(), row["packing"]))


@router.get("/analytics/supplier-prices")
def supplier_prices(db: DB, user: Actor, date_from: date | None = None, date_to: date | None = None, q: str = "",
                    supplier_id: str | None = None, page: int = Query(1, ge=1),
                    page_size: int = Query(100, ge=1, le=500)):
    require_permission(db, user, "analytics.read")
    require_permission(db, user, "finance.purchase.read")
    query = price_query(db, user, date_from, date_to, q, supplier_id)
    rows = list(db.scalars(query.limit(page_size).offset((page - 1) * page_size)))
    items, profile = price_rows(db, rows)
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery()))
    return plain({"items": items, "total": total, "page": page, "page_size": page_size, "rates_profile": profile})


def sheet(book: Workbook, title: str, columns: list[tuple[str, str, int]], rows: list[dict[str, Any]],
          first: bool = False) -> None:
    page = book.active if first else book.create_sheet()
    page.title = title
    page.append([label for _, label, _ in columns])
    for row in rows:
        page.append([cell(row[key]) for key, _, _ in columns])
    for index, (key, _, width) in enumerate(columns, 1):
        letter = get_column_letter(index)
        page.column_dimensions[letter].width = width
        style = ("DD.MM.YYYY" if key in ("quoted_at", "valid_until")
                 else "#,##0.00" if key in ("unit_price_rub", "highest_rub")
                 else "#,##0.####" if key in ("unit_price", "quantity") else None)
        if style:
            for (value,) in page.iter_rows(min_row=2, min_col=index, max_col=index):
                value.number_format = style
    for value in page[1]:
        value.font = Font(name="Arial", bold=True, color="FFFFFF")
        value.fill = PatternFill("solid", fgColor=INK)
        value.alignment = Alignment(wrap_text=True, vertical="center")
    page.row_dimensions[1].height = 32
    page.freeze_panes = "A2"
    page.auto_filter.ref = page.dimensions


def cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str):
        # A spreadsheet must not run a supplier's text as a formula.
        return "'" + value if value[:1] in ("=", "+", "-", "@") else value
    return value


@router.get("/analytics/supplier-prices.xlsx")
def supplier_prices_xlsx(db: DB, user: Actor, date_from: date | None = None, date_to: date | None = None,
                         q: str = "", supplier_id: str | None = None):
    require_permission(db, user, "analytics.read")
    require_permission(db, user, "finance.purchase.read")
    require_permission(db, user, "exports.download")
    rows, profile = price_rows(db, list(db.scalars(
        price_query(db, user, date_from, date_to, q, supplier_id).limit(EXPORT_LIMIT))))
    for row in rows:
        row["unit_price_rub"] = Decimal(row["unit_price_rub"]) if row["unit_price_rub"] is not None else None
    book = Workbook()
    sheet(book, "Цены поставщиков", COLUMNS, rows, first=True)
    sheet(book, "Лучшие цены", BEST_COLUMNS, best_prices(rows))
    notes = book.create_sheet("О выгрузке")
    start, finish = period(date_from, date_to)
    for line in (
        ["Период квот", f"{date_from or start.date()} — {date_to or finish.date()}"],
        ["Поиск", q or "—"],
        ["Позиций", len(rows)],
        ["Рубли", f"по курсам профиля «{profile}»" if profile else "нет опубликованного профиля с курсами"],
        ["Лучшие цены", "Один товар с одной фасовкой у разных поставщиков сравнивается по словам наименования; "
                        "позиции без курса валюты в сравнение не входят."],
    ):
        notes.append(line)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 90
    for value in notes["A"]:
        value.font = Font(name="Arial", bold=True, color=INK)
    buffer = io.BytesIO()
    book.save(buffer)
    audit(db, user, "analytics", user.id, "supplier_prices_exported",
          after={"date_from": str(date_from), "date_to": str(date_to), "q": q, "rows": len(rows)})
    name = f"База цен поставщиков {date.today():%d.%m.%Y}.xlsx"
    return Response(buffer.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=\"supplier-prices.xlsx\"; "
                                                    f"filename*=UTF-8''{quote(name)}"})

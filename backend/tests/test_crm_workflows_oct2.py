"""October workflow regressions: reusable RFQs, source formatting and counterparty records."""

import io
from datetime import date, timedelta
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from sqlalchemy import delete, select
from test_business_imports import confirm, request
from test_crm import crm as crm
from test_crm import login, post

from app.communication.models import FileRecord
from app.crm.codes import compact_request_number, request_number_value
from app.crm.models import CounterpartyDocument, Request


@pytest.mark.parametrize(('value', 'number'), [(1, '1'), (1000, '1000'), (1001, 'A1'),
                                             (2000, 'A1000'), (2001, 'B1'), (27000, 'Z1000'),
                                             (27001, 'AA1'), (703001, 'AAA1')])
def test_compact_request_number_boundaries(value, number):
    assert compact_request_number(value) == number
    assert request_number_value(number) == value


def test_request_numbers_preserve_legacy_and_do_not_reuse_deleted_numbers(crm):
    login(crm)
    req = request(crm)
    with crm['sessions'].begin() as db:
        db.add(Request(number='З-2026-00005', title='Старая', client_id=req['client_id'], owner_id=crm['admin'].id))
    second = post(crm, '/requests', {'title': 'Вторая', 'client_id': req['client_id']})
    assert req['number'] == '1' and second['number'] == '2'
    with crm['sessions'].begin() as db:
        db.execute(delete(Request).where(Request.id == second['id']))
    third = post(crm, '/requests', {'title': 'Третья', 'client_id': req['client_id']})
    assert third['number'] == '3'
    assert crm['client'].get('/api/v1/requests?q=З-2026-00005').json()['items'][0]['number'] == 'З-2026-00005'


def test_supplier_contacts_owner_names_and_access_scope(crm):
    login(crm)
    supplier = post(crm, '/counterparties', {'name': 'Поставщик', 'kind': 'supplier'})
    assert supplier['owner_name'] == crm['admin'].name
    contact = post(crm, '/contacts', {'client_id': supplier['id'], 'name': 'Алексей', 'department': 'Продажи'})
    assert contact['id'] in {row['id'] for row in crm['client'].get('/api/v1/contacts').json()['items']}
    assert supplier['id'] in {row['id'] for row in crm['client'].get('/api/v1/counterparties?contact_eligible=true').json()['items']}
    changed = crm['client'].patch(f"/api/v1/counterparties/{supplier['id']}", json={
        'version': supplier['version'], 'owner_id': crm['manager'].id, 'reason': 'Передача клиента'
    })
    assert changed.status_code == 200, changed.text
    assert changed.json()['owner_name'] == crm['manager'].name
    hidden = post(crm, '/counterparties', {'name': 'Другой поставщик', 'kind': 'supplier'})
    hidden_contact = post(crm, '/contacts', {'client_id': hidden['id'], 'name': 'Скрытый'})
    login(crm, 'manager@example.com')
    assert crm['client'].get(f"/api/v1/contacts/{contact['id']}").status_code == 200
    assert crm['client'].get(f"/api/v1/contacts/{hidden_contact['id']}").status_code == 404
    post(crm, f"/counterparties/{hidden['id']}/contacts", {'name': 'Запрещено'}, 404)
    assert hidden_contact['id'] not in {row['id'] for row in crm['client'].get('/api/v1/contacts').json()['items']}


def test_counterparty_documents_quarantine_download_archive_replay_and_scope(crm):
    login(crm)
    customer = post(crm, '/counterparties', {'name': 'Документы'})
    path = f"/api/v1/counterparties/{customer['id']}/documents"
    content = b'%PDF-1.4\n1 0 obj <<>> endobj\n%%EOF'
    key = str(uuid4())
    def upload():
        return crm['client'].post(path, files={'file': ('contract.pdf', content, 'application/pdf')},
            data={'category': 'contract'}, headers={'Idempotency-Key': key})
    first = upload()
    assert first.status_code == 201, first.text
    row = first.json()
    assert row['category'] == 'contract' and row['status'] == 'quarantined'
    assert upload().json()['id'] == row['id']
    endpoint = f"/api/v1/counterparty-documents/{row['id']}"
    assert crm['client'].get(endpoint + '/download').status_code == 409
    with crm['sessions'].begin() as db:
        file = db.get(FileRecord, row['file_id'])
        file.status = 'clean'
        assert len(db.scalars(select(CounterpartyDocument)).all()) == 1
    response = crm['client'].get(endpoint + '/download')
    assert response.status_code == 200 and response.content == content
    archived = crm['client'].patch(endpoint, json={'version': row['version'], 'archived': True})
    assert archived.status_code == 200, archived.text
    assert crm['client'].get(path).json()['total'] == 0
    assert crm['client'].get(path + '?archived=true').json()['items'][0]['id'] == row['id']
    restored = crm['client'].patch(endpoint, json={'version': archived.json()['version'], 'archived': False})
    assert restored.status_code == 200
    login(crm, 'manager@example.com')
    assert crm['client'].get(path).status_code == 404
    assert crm['client'].get(endpoint + '/download').status_code == 404
    assert crm['client'].patch(endpoint, json={'version': restored.json()['version'], 'archived': True}).status_code == 404


def import_items(crm, request_id, filename, content):
    response = crm['client'].post(f'/api/v1/requests/{request_id}/table-imports/preview',
                                files={'file': (filename, content)}, data={'kind': 'items'})
    assert response.status_code == 200, response.text
    confirm(crm, request_id, response.json()['id'])
    return crm['client'].get(f'/api/v1/requests/{request_id}/items').json()['items']


def test_common_rfq_keeps_original_borders_widths_header_and_selected_rows(crm):
    login(crm)
    req = request(crm)
    book = Workbook()
    sheet = book.active
    sheet.append(['Name', 'Art', 'Qty'])
    sheet.append(['Название 1', '0012', '1'])
    sheet.append(['', '', ''])
    sheet.append(['Название 2', '0007', '2'])
    border = Border(left=Side(style='medium', color='FF0000'), bottom=Side(style='thin'))
    for row in sheet:
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, horizontal='center')
    sheet['A1'].font = Font(name='Calibri', size=14, bold=True, color='135724')
    sheet.column_dimensions['A'].width = 83
    sheet.row_dimensions[4].height = 47
    stream = io.BytesIO()
    book.save(stream)
    items = import_items(crm, req['id'], 'demand.xlsx', stream.getvalue())
    selected = next(row for row in items if row['source_values'][0] == 'Название 2')
    rfq = post(crm, f"/requests/{req['id']}/rfqs", {
        'idempotency_key': str(uuid4()), 'item_ids': [selected['id']], 'response_due': date.today().isoformat()
    }, 200)
    assert rfq['supplier_id'] is None and rfq['snapshot']['supplier'] is None
    response = crm['client'].get(f"/api/v1/rfqs/{rfq['id']}/file")
    assert response.status_code == 200
    exported = load_workbook(io.BytesIO(response.content))
    result = exported.active
    assert list(result.values) == [('Name', 'Art', 'Qty'), ('Название 2', '0007', '2')]
    assert result['A1'].font.color.rgb == '00135724'
    assert result['A2'].border.left.style == 'medium'
    assert result['A2'].border.left.color.rgb == '00FF0000'
    assert result['A2'].alignment.horizontal == 'center'
    assert result.column_dimensions['A'].width == 83 and result.row_dimensions[2].height == 47
    assert all(cell.data_type != 'f' for row in result for cell in row)
    for name in ('Первый', 'Второй'):
        supplier = post(crm, '/counterparties', {'name': name, 'kind': 'supplier'})
        quote = post(crm, f"/requests/{req['id']}/quotes", {
            'idempotency_key': str(uuid4()), 'item_id': selected['id'], 'item_revision': selected['revision'],
            'supplier_id': supplier['id'], 'supplier_request_id': rfq['id'],
            'product': {'name': 'Этанол', 'cas': '64-17-5', 'manufacturer': 'Maker', 'unit': 'pcs'},
            'price': '10', 'currency': 'RUB', 'price_unit': 'pcs', 'available_quantity': '2',
            'valid_until': (date.today() + timedelta(days=30)).isoformat()
        }, 200)
        assert quote['supplier_request_id'] == rfq['id']


def test_text_rfq_uses_readable_grid_and_keeps_formula_like_values_literal(crm):
    login(crm)
    req = request(crm)
    items = import_items(crm, req['id'], 'demand.txt', b'Name\tQty\n=literal\t1')
    rfq = post(crm, f"/requests/{req['id']}/rfqs", {
        'idempotency_key': str(uuid4()), 'item_ids': [items[0]['id']], 'response_due': date.today().isoformat()
    }, 200)
    response = crm['client'].get(f"/api/v1/rfqs/{rfq['id']}/file")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet['A2'].value == '=literal' and sheet['A2'].data_type == 's'
    assert sheet['A1'].border.bottom.style == 'thin' and sheet['A2'].border.right.style == 'thin'

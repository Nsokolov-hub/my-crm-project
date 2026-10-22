"""Business remarks of 21 October 2026: suppliers that quoted in the request list, quote rows placed by their
product, the manufacturer of a later quote, the quantity of the request in a calculation, clients in Belarus
with VAT 0 %, and the base of supplier prices in Excel."""

import io
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import select

from app.business.models import SupplierMail
from app.commerce.files import BELARUS_VAT_NOTE, read_file
from app.commerce.models import CommercialDocument
from app.core.db import utcnow
from app.crm.models import Nomenclature
from app.crm.table_imports import QUOTE_COLUMNS
from tests.support.business import cmd, confirm, issue, preview, quote_file, request, structured
from tests.support.crm import login, post


def get(crm, path):
    response = crm["client"].get(f"/api/v1{path}")
    assert response.status_code == 200, response.text
    return response.json()


def quotes(crm, req, supplier, rows):
    batch = preview(crm, req["id"], list(QUOTE_COLUMNS.values()),
                    [quote_file(supplier["internal_code"], row) for row in rows], "quotes")
    assert batch["summary"]["errors"] == 0, batch["rows"]
    return batch


def test_the_request_list_counts_the_suppliers_that_gave_a_quote(crm):
    req, _, _, supplier = structured(crm)
    row = next(row for row in get(crm, "/requests")["items"] if row["id"] == req["id"])
    # A supplier that quoted without a request from the CRM counts as asked (remark of 22 October).
    assert (row["quoted_suppliers"], row["answered_suppliers"], row["asked_suppliers"]) == (1, 1, 1)
    other = post(crm, "/counterparties", {"name": "Beijing Asnail", "kind": "supplier"})
    with crm["sessions"].begin() as db:
        db.add_all(SupplierMail(request_id=req["id"], supplier_id=supplier_id, recipient="rfq@example.com",
                                subject="RFQ", body="", html_body="", status="sent", sent_at=utcnow(),
                                author_id=crm["admin"].id) for supplier_id in (supplier["id"], other["id"]))
    row = next(row for row in get(crm, "/requests")["items"] if row["id"] == req["id"])
    assert (row["quoted_suppliers"], row["answered_suppliers"], row["asked_suppliers"]) == (1, 1, 2)


def test_quote_rows_go_to_the_positions_of_their_products(crm):
    """Beijing Asnail numbered the rows of its own quote: the bottle came as row 2, the request holds it as 4."""
    login(crm)
    req = request(crm)
    confirm(crm, req["id"], preview(crm, req["id"], ["Name", "Qty"], [
        ["Narasin", "10"], ["Nicarbazin", "5"], ["Avilamycin A", "5"], ["Density Bottle – calibrated – 5 ml", "5"],
    ])["id"])
    items = get(crm, f"/requests/{req['id']}/items")["items"]
    asnail = post(crm, "/counterparties", {"name": "Beijing Asnail", "kind": "supplier"})
    batch = quotes(crm, req, asnail, [
        {"source_row": "1", "article": "TRC-A794655-250MG", "name": "Avilamycin A", "packing_value": "250",
         "packing_unit": "mg", "unit_price": "4675", "currency": "rub"},
        {"source_row": "2", "article": "023.02.005", "name": "Density Bottle, calibrated, 5 mL", "unit_price": "160"},
        # A name the request does not know: the number in the file stands.
        {"source_row": "2", "article": "NCZ-1", "name": "Никарбазин", "unit_price": "900"},
    ])
    assert [(row["data"]["source_position"], row["data"].get("source_position_stated")) for row in batch["rows"]] == [
        (3, 1), (4, 2), (2, None)]
    confirm(crm, req["id"], batch["id"])
    quoted = get(crm, f"/requests/{req['id']}/quote-items?all=true")["items"]
    placed = {row["nomenclature_name"]: row["source_request_item_id"] for row in quoted}
    assert placed == {"Avilamycin A": items[2]["id"], "Density Bottle, calibrated, 5 mL": items[3]["id"],
                      "Никарбазин": items[1]["id"]}
    # The calculation takes the quantity of the request position.
    assert {row["nomenclature_name"]: row["request_packings"] for row in quoted}["Avilamycin A"] == "5"

    # The same product named by two positions: the number in the file settles it.
    other = post(crm, "/counterparties", {"name": "Shenzhen Sungening", "kind": "supplier"})
    tie = quotes(crm, req, other, [{"source_row": "3", "article": "TRC-A794655-250MG", "name": "Avilamycin A",
                                    "packing_value": "250", "packing_unit": "mg"}])
    assert tie["rows"][0]["data"]["source_position"] == 3


def test_a_later_quote_gives_its_manufacturer_to_a_nomenclature_created_without_one(crm):
    login(crm)
    req = request(crm)
    shenzhen = post(crm, "/counterparties", {"name": "Shenzhen Sungening", "kind": "supplier"})
    row = {"source_row": "", "article": "AMBH2D6EF613-250MG", "name": "Pentanedioic acid derivative",
           "cas": "137281-23-3", "packing_value": "250", "packing_unit": "mg", "unit_price": "2100"}
    first = quotes(crm, req, shenzhen, [row])
    confirm(crm, req["id"], first["id"])
    with crm["sessions"]() as db:
        product = db.scalar(select(Nomenclature).where(Nomenclature.article == "AMBH2D6EF613-250MG"))
        assert (product.manufacturer, product.manufacturer_id) == (None, None)

    second = quotes(crm, req, shenzhen, [{**row, "manufacturer": "Sigma-Aldrich", "unit_price": "2050"}])
    assert second["rows"][0]["data"]["manufacturer_filled"] is True
    [sheet_id] = confirm(crm, req["id"], second["id"])["result"]["quote_sheet_ids"]
    with crm["sessions"]() as db:
        product = db.scalar(select(Nomenclature).where(Nomenclature.article == "AMBH2D6EF613-250MG"))
        assert product.manufacturer == "Sigma-Aldrich" and product.manufacturer_id
    names = {item["manufacturer"] for item in get(crm, f"/requests/{req['id']}/quote-items?all=true")["items"]}
    assert names == {"Sigma-Aldrich"}

    # A file whose quotes were deleted is placed anew instead of being taken for a duplicate.
    sheet = get(crm, f"/quote-sheets/{sheet_id}")
    cmd(crm, f"/quote-sheets/{sheet_id}/delete", {"version": sheet["version"]})
    again = quotes(crm, req, shenzhen, [{**row, "manufacturer": "Sigma-Aldrich", "unit_price": "2050"}])
    assert confirm(crm, req["id"], again["id"])["result"].get("already_imported") is None


def test_a_client_in_belarus_gets_vat_zero_and_the_note_of_the_eaeu_treaty(crm):
    req, data, _, _ = structured(crm, mode="IMPORT")
    before = cmd(crm, f"/requests/{req['id']}/calculations", data)
    client = get(crm, f"/counterparties/{req['client_id']}")
    changed = crm["client"].patch(f"/api/v1/counterparties/{client['id']}",
                                  json={"version": client["version"], "in_belarus": True})
    assert changed.status_code == 200 and changed.json()["in_belarus"] is True
    assert get(crm, f"/requests/{req['id']}")["client_in_belarus"] is True
    # A calculation made before the client was marked cannot be issued: it carries 22 %.
    stale = cmd(crm, f"/requests/{req['id']}/proposals", {
        "calculation_id": before["id"], "valid_until": (date.today() + timedelta(days=30)).isoformat(),
        "terms": "Оплата согласно договору."}, status=422)
    assert stale["code"] == "CALCULATION_VAT_CLIENT"

    current = get(crm, f"/requests/{req['id']}")
    data["request_version"] = current["version"]
    refused = cmd(crm, f"/requests/{req['id']}/calculations/preview", {**data, "vat_deductible": True}, status=422)
    assert refused["code"] == "VAT_DEDUCTION_BELARUS"
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    snapshot = calculation["snapshot"]
    line = snapshot["lines"][0]
    assert (snapshot["sale_vat"], snapshot["vat_deductible"]) == ({"rate": "0", "reason": "belarus"}, False)
    assert Decimal(line["tax"]) == 0 and Decimal(line["total"]) == Decimal(line["net"])
    assert line["tax_category"] == "НДС 0%"
    # The import VAT is still paid at the border and stays in the cost.
    assert Decimal(line["detail"]["import_vat"]) > 0
    assert Decimal(line["detail"]["clean_cost"]) == (Decimal(line["detail"]["customs_base"])
                                                     + Decimal(line["detail"]["duty"])
                                                     + Decimal(line["detail"]["import_vat"]))

    proposal = issue(crm, req, calculation)
    assert (proposal["snapshot"]["vat_rate"], proposal["snapshot"]["vat_note"]) == ("0", BELARUS_VAT_NOTE)
    assert Decimal(proposal["snapshot"]["totals"]["tax"]) == 0
    with crm["sessions"]() as db:
        document = db.get(CommercialDocument, proposal["id"])
        values = [cell.value for row in load_workbook(io.BytesIO(read_file(document.files["xlsx"]))).active
                  for cell in row]
        assert read_file(document.files["pdf"]).startswith(b"%PDF")
    assert BELARUS_VAT_NOTE in values and "В том числе НДС 0 %" in values


def test_the_base_of_supplier_prices_is_shown_and_exported_with_the_best_prices(crm):
    req, _, _, _ = structured(crm)
    period = {"date_from": (date.today() - timedelta(days=1)).isoformat(), "date_to": date.today().isoformat()}
    prices = crm["client"].get("/api/v1/analytics/supplier-prices", params=period).json()
    assert prices["total"] == 1
    row = prices["items"][0]
    assert (row["supplier_name"], row["nomenclature_name"], row["currency"], row["unit_price_rub"],
            row["request_number"]) == ("Aozeal", "Sample", "RUB", "1000.00", req["number"])
    assert crm["client"].get("/api/v1/analytics/supplier-prices",
                             params={**period, "q": "нет такого"}).json()["total"] == 0

    exported = crm["client"].get("/api/v1/analytics/supplier-prices.xlsx", params={**period, "q": "sample"})
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content))
    assert book.sheetnames == ["Цены поставщиков", "Лучшие цены", "О выгрузке"]
    header, first = [[cell.value for cell in line] for line in book["Цены поставщиков"].iter_rows(max_row=2)]
    assert dict(zip(header, first, strict=True))["Цена, ₽"] == 1000
    best = [[cell.value for cell in line] for line in book["Лучшие цены"].iter_rows(max_row=2)]
    assert dict(zip(*best, strict=True))["Поставщик"] == "Aozeal"

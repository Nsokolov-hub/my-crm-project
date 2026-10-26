"""Site requests in Excel, the report by channels, joining a site request to an existing CRM request."""

import io
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.db import utcnow
from app.crm.models import Request, RequestItem
from app.integrations.models import SiteRequest
from tests.integrations.test_site_conversion import command, conversion_body, received
from tests.support.crm import login
from tests.support.site import payload


def test_list_and_one_site_request_in_excel(site):
    site_request = received(site)
    login(site)
    listing = site['client'].get('/api/v1/site-requests/export.xlsx')
    assert listing.status_code == 200, listing.text
    sheet = load_workbook(io.BytesIO(listing.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ('Номер', 'Получена', 'Тип') and rows[1][4] == 'ООО «Пример-Фарма»' and rows[1][11] == 'Реклама'
    one = site['client'].get(f'/api/v1/site-requests/{site_request}/export.xlsx')
    book = load_workbook(io.BytesIO(one.content))
    assert book.sheetnames == ['Заявка', 'Позиции']
    lines = list(book['Позиции'].iter_rows(values_only=True))
    assert lines[1][3] == 'Y0001897' and lines[1][8] == 'Готовы рассмотреть аналоги'


def test_report_by_channels(site):
    first = received(site)
    received(site, attribution={**payload()['attribution'], 'channel': 'organic'})
    third = received(site)
    login(site)
    command(site, f'{third}/reject', {'reason': 'Дубль', 'version': 1})
    command(site, f'{first}/assign', {'assignee_id': site['admin'].id, 'version': 1})
    window = site['client'].get(f'/api/v1/site-requests/{first}/conversion').json()
    created = command(site, f'{first}/convert', conversion_body(window, version=2))['created']
    with site['sessions'].begin() as db:
        request = db.get(Request, created['request_id'])
        request.commercial_stage, request.sale_confirmed_at = 'proposal_sent', utcnow()
    report = site['client'].get('/api/v1/site-requests/report').json()
    channels = {line['channel']: line for line in report['channels']}
    assert channels['cpc']['leads'] == 2 and channels['cpc']['rejected'] == 1 and channels['cpc']['sales'] == 1
    assert channels['cpc']['proposals'] == 1 and channels['organic']['leads'] == 1
    assert report['totals']['leads'] == 3


def test_a_site_request_joins_an_existing_crm_request(site):
    first = received(site)
    second = received(site, site_number='S-2026-000124')
    login(site)
    command(site, f'{first}/assign', {'assignee_id': site['admin'].id, 'version': 1})
    window = site['client'].get(f'/api/v1/site-requests/{first}/conversion').json()
    created = command(site, f'{first}/convert', conversion_body(window, version=2))['created']
    other = site['client'].get(f'/api/v1/site-requests/{second}/conversion').json()
    joined = site['client'].post(f'/api/v1/site-requests/{second}/attach',
                                 json={'version': 1, 'request_id': created['request_id'],
                                       'items': other['proposal']['items']}, headers={'Idempotency-Key': str(uuid4())})
    assert joined.status_code == 200, joined.text
    assert joined.json()['status'] == 'converted' and joined.json()['created']['request_number'] == created['request_number']
    with site['sessions']() as db:
        assert db.scalar(select(func.count()).select_from(RequestItem).where(
            RequestItem.request_id == created['request_id'])) == 2
        assert db.get(SiteRequest, second).request_id == created['request_id']
    again = site['client'].post(f'/api/v1/site-requests/{second}/attach',
                                json={'version': 2, 'request_id': created['request_id'], 'items': []},
                                headers={'Idempotency-Key': str(uuid4())})
    assert again.status_code == 409


def test_the_report_ends_today_in_the_company_timezone(site, monkeypatch):
    from datetime import datetime as real_datetime
    from zoneinfo import ZoneInfo

    from app.integrations import site_reports

    class LateEvening(real_datetime):
        @classmethod
        def now(cls, tz=None):
            # 22:30 UTC on 26 October is already 27 October in Moscow.
            return real_datetime(2026, 10, 26, 22, 30, tzinfo=ZoneInfo('UTC')).astimezone(tz)

    monkeypatch.setattr(site_reports, 'datetime', LateEvening)
    login(site)
    report = site['client'].get('/api/v1/site-requests/report').json()
    assert report['date_to'] == '2026-10-27'

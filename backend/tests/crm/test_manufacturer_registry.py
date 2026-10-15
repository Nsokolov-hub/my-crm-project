"""Registry of manufacturers: one entry per manufacturer whatever the spelling, numbers for Excel, brands."""

import io
from decimal import Decimal
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import select

from app.commerce.models import Execution, Manufacturer
from app.core.models import AuditEvent
from app.crm.models import Nomenclature
from app.crm.table_imports import QUOTE_COLUMNS
from tests.support.business import cmd, confirm, preview, quote_file, release, request, structured
from tests.support.crm import login, post

PACKING = [{"value": "10", "unit": "mg"}]


def start(env):
    """Sign in; the first request also adds the default product group the nomenclature falls into."""
    login(env)
    request(env)


def nomenclature(env, name, manufacturer=None, article=None, **extra):
    body = {"name": name, "article": article or name, "packings": PACKING, **extra}
    if manufacturer is not None:
        body["manufacturer"] = manufacturer
    return post(env, "/nomenclatures", body)


def get(env, path, status=200):
    response = env["client"].get("/api/v1" + path)
    assert response.status_code == status, response.text
    return response.json()


def change(env, path, body, status=200, method="patch"):
    response = getattr(env["client"], method)("/api/v1" + path, json=body)
    assert response.status_code == status, response.text
    return response.json()


def merge(env, source, target, status=200, key=None):
    return change(env, f"/manufacturers/{source['id']}/merge", {
        "version": source["version"], "target_id": target["id"], "reason": "Одно и то же написание",
        "idempotency_key": key or str(uuid4())}, status, "post")


def test_a_spelling_names_one_entry_and_a_number_finds_it(crm):
    start(crm)
    first = nomenclature(crm, "Atropine", "Sigma-Aldrich")
    maker = get(crm, f"/manufacturers/{first['manufacturer_id']}")
    assert (maker["internal_code"], maker["name"], maker["label"], maker["nomenclatures"]) == (
        1, "Sigma-Aldrich", "1 · Sigma-Aldrich", 1)
    again = nomenclature(crm, "Caffeine", "  sigma-ALDRICH ")
    by_number = nomenclature(crm, "Codeine", str(maker["internal_code"]))
    assert {again["manufacturer_id"], by_number["manufacturer_id"]} == {maker["id"]}
    assert again["manufacturer"] == by_number["manufacturer"] == "Sigma-Aldrich"
    assert (again["manufacturer_code"], again["umbrella_name"]) == (1, None)
    unknown = post(crm, "/nomenclatures", {"name": "X", "manufacturer": "77", "packings": PACKING}, status=422)
    assert unknown["code"] == "MANUFACTURER_CODE"
    other = nomenclature(crm, "Morphine", "Supelco")
    assert get(crm, f"/manufacturers/{other['manufacturer_id']}")["internal_code"] == 2

    chosen = change(crm, f"/nomenclatures/{other['id']}", {"version": other["version"], "manufacturer_id": maker["id"]})
    assert (chosen["manufacturer_id"], chosen["manufacturer"]) == (maker["id"], "Sigma-Aldrich")
    cleared = change(crm, f"/nomenclatures/{chosen['id']}", {"version": chosen["version"], "manufacturer": None})
    assert (cleared["manufacturer_id"], cleared["manufacturer"]) == (None, None)
    listed = get(crm, f"/nomenclatures?manufacturer_id={maker['id']}")
    assert {row["name"] for row in listed["items"]} == {"Atropine", "Caffeine", "Codeine"}


def test_registry_keeps_spellings_unique_and_a_rename_keeps_the_old_one(crm):
    start(crm)
    merck = post(crm, "/manufacturers", {"name": "Merck", "aliases": ["Merck KGaA", "MERCK  kgaa"]})
    assert (merck["aliases"], merck["nomenclatures"]) == (["Merck KGaA"], 0)
    taken = post(crm, "/manufacturers", {"name": "merck kgaa"}, status=409)
    assert (taken["code"], taken["message"]) == (
        "MANUFACTURER_EXISTS", f"Написание «Merck KGaA» уже есть у производителя «Merck» (№ {merck['internal_code']})")
    assert post(crm, "/manufacturers", {"name": "123"}, status=422)["code"] == "MANUFACTURER_NAME"
    item = nomenclature(crm, "Atropine", "Merck KGaA")
    assert (item["manufacturer_id"], item["manufacturer"]) == (merck["id"], "Merck")

    renamed = change(crm, f"/manufacturers/{merck['id']}", {"version": merck["version"], "name": "Merck Group"})
    assert (renamed["name"], renamed["aliases"]) == ("Merck Group", ["Merck", "Merck KGaA"])
    with crm["sessions"]() as db:
        assert db.get(Nomenclature, item["id"]).manufacturer == "Merck Group"
    assert nomenclature(crm, "Caffeine", "merck")["manufacturer_id"] == merck["id"]
    trimmed = change(crm, f"/manufacturers/{merck['id']}", {"version": renamed["version"], "aliases": ["Merck"]})
    assert trimmed["aliases"] == ["Merck"]
    stale = change(crm, f"/manufacturers/{merck['id']}", {"version": merck["version"], "name": "Другое"}, 409)
    assert stale["code"] == "VERSION_CONFLICT"
    assert [row["name"] for row in get(crm, "/manufacturers?q=kgaa")["items"]] == []
    assert [row["name"] for row in get(crm, "/manufacturers?q=merck")["items"]] == ["Merck Group"]
    assert [row["name"] for row in get(crm, f"/manufacturers?q={merck['internal_code']}")["items"]] == ["Merck Group"]


def test_umbrella_brands_have_one_level(crm):
    start(crm)
    merck = post(crm, "/manufacturers", {"name": "Merck"})
    sigma = post(crm, "/manufacturers", {"name": "Sigma-Aldrich", "umbrella_id": merck["id"]})
    supelco = post(crm, "/manufacturers", {"name": "Supelco"})
    assert (sigma["umbrella_id"], sigma["umbrella_name"]) == (merck["id"], "Merck")
    assert get(crm, f"/manufacturers/{merck['id']}")["brand_members"] == 1
    nested = change(crm, f"/manufacturers/{supelco['id']}", {"version": supelco["version"], "umbrella_id": sigma["id"]},
                    422)
    assert nested["code"] == "UMBRELLA_INVALID"
    owner = change(crm, f"/manufacturers/{merck['id']}", {"version": merck["version"], "umbrella_id": supelco["id"]},
                   422)
    assert owner["message"] == "«Merck» сам зонтичный бренд для других производителей"
    itself = change(crm, f"/manufacturers/{supelco['id']}", {"version": supelco["version"], "umbrella_id": supelco["id"]},
                    422)
    assert itself["code"] == "UMBRELLA_INVALID"
    joined = change(crm, f"/manufacturers/{supelco['id']}", {"version": supelco["version"], "umbrella_id": merck["id"]})
    assert joined["umbrella_name"] == "Merck"
    assert {row["name"] for row in get(crm, "/manufacturers?umbrellas=true")["items"]} == {"Merck"}
    alone = change(crm, f"/manufacturers/{joined['id']}", {"version": joined["version"], "umbrella_id": None})
    assert alone["umbrella_id"] is None


def test_duplicates_are_suggested_and_merged_into_the_main_entry(crm):
    start(crm)
    for name, maker in (("Atropine", "Sigma-Aldrich"), ("Caffeine", "Sigma Aldrich"), ("Codeine", "Sigma"),
                        ("Morphine", "Chemicea"), ("Heroin", "Chemicea Pharmaceuticals Pvt Ltd,"), ("Salt", "EP"),
                        ("Sugar", "EPA Standards")):
        nomenclature(crm, name, maker)
    groups = get(crm, "/manufacturers/duplicates")["items"]
    assert [sorted(row["name"] for row in group["manufacturers"]) for group in groups] == [
        ["Sigma", "Sigma Aldrich", "Sigma-Aldrich"], ["Chemicea", "Chemicea Pharmaceuticals Pvt Ltd,"]]

    main, *copies = sorted(groups[0]["manufacturers"], key=lambda row: row["name"] != "Sigma-Aldrich")
    umbrella = post(crm, "/manufacturers", {"name": "Merck"})
    copy = change(crm, f"/manufacturers/{copies[0]['id']}", {"version": copies[0]["version"],
                                                              "umbrella_id": umbrella["id"]})
    key = str(uuid4())
    merged = merge(crm, copy, main, key=key)
    assert merged["moved"] == {"aliases": 1, "nomenclatures": 1, "members": 0}
    assert merge(crm, copy, main, key=key) == merged
    merge(crm, copies[1], main)
    assert sorted(merged["aliases"]) == [copy["name"]]
    current = get(crm, f"/manufacturers/{main['id']}")
    assert (current["nomenclatures"], sorted(current["aliases"])) == (3, ["Sigma", "Sigma Aldrich"])
    with crm["sessions"]() as db:
        texts = set(db.scalars(select(Nomenclature.manufacturer).where(Nomenclature.manufacturer_id == main["id"])))
        assert texts == {"Sigma-Aldrich"}
        assert db.get(Manufacturer, copy["id"]).merged_into_id == main["id"]
        assert db.query(AuditEvent).filter_by(entity_id=copy["id"], action="merged").one().reason == (
            "Одно и то же написание")
    assert nomenclature(crm, "Codeine 2", "sigma aldrich")["manufacturer_id"] == main["id"]
    gone = get(crm, f"/manufacturers/{copy['id']}")
    assert (gone["merged_into_id"], gone["merged_into_name"]) == (main["id"], "Sigma-Aldrich")
    assert copy["id"] not in {row["id"] for row in get(crm, "/manufacturers")["items"]}
    assert copy["id"] in {row["id"] for row in get(crm, "/manufacturers?merged=true")["items"]}
    again = merge(crm, gone, main, status=409)
    assert (again["code"], again["message"]) == ("MANUFACTURER_MERGED",
                                                 f"Производитель объединён с «Sigma-Aldrich» (№ {main['internal_code']})")
    assert merge(crm, current, current, status=422)["code"] == "MERGE_TARGET"
    assert [len(group["manufacturers"]) for group in get(crm, "/manufacturers/duplicates")["items"]] == [2]


def test_quote_import_takes_registry_numbers_and_adds_new_spellings(crm):
    login(crm)
    req = request(crm)
    supplier = post(crm, "/counterparties", {"name": "Поставщик", "kind": "supplier"})
    raw = preview(crm, req["id"], ["Name", "Article", "Quantity"],
                  [["Потребность 1", "A-1", "1"], ["Потребность 2", "A-2", "1"], ["Потребность 3", "A-3", "1"]])
    confirm(crm, req["id"], raw["id"])
    known = post(crm, "/manufacturers", {"name": "Sigma-Aldrich"})
    rows = [
        quote_file(supplier["internal_code"], {"source_row": "1", "article": "A-1", "manufacturer": str(known["internal_code"])}),
        quote_file(supplier["internal_code"], {"source_row": "2", "article": "A-2", "manufacturer": "Toronto Research"}),
        quote_file(supplier["internal_code"], {"source_row": "3", "article": "A-3", "manufacturer": "99"}),
    ]
    refused = preview(crm, req["id"], list(QUOTE_COLUMNS.values()), rows, "quotes")
    assert refused["rows"][2]["errors"] == ["Производителя с номером 99 нет в реестре"]
    batch = preview(crm, req["id"], list(QUOTE_COLUMNS.values()), rows[:2], "quotes")
    assert batch["summary"] == {"total": 2, "checked": 0, "create_nomenclature": 2, "new_manufacturers": 1,
                                "errors": 0, "request_positions": 3, "missing_positions": [3]}
    assert [row["data"]["manufacturer"] for row in batch["rows"]] == ["Sigma-Aldrich", "Toronto Research"]
    result = confirm(crm, req["id"], batch["id"])["result"]
    assert (result["nomenclatures_created"], result["manufacturers_created"]) == (2, 1)
    with crm["sessions"]() as db:
        made = {row.article: row for row in db.scalars(select(Nomenclature))}
        assert made["A-1"].manufacturer_id == known["id"]
        added = db.get(Manufacturer, made["A-2"].manufacturer_id)
        assert (added.name, added.internal_code) == ("Toronto Research", known["internal_code"] + 1)

    response = crm["client"].get(f"/api/v1/requests/{req['id']}/quote-import/template.xlsx")
    book = load_workbook(io.BytesIO(response.content))
    makers = [[cell.value for cell in row] for row in book["Производители"].iter_rows()]
    assert makers[0] == ["Номер", "Производитель", "Зонтичный бренд", "Другие написания"]
    assert ["1", "Sigma-Aldrich", None, None] in makers
    column = list(QUOTE_COLUMNS).index("manufacturer") + 1
    assert {book["Квоты"].cell(row, column).value for row in (2, 3)} == {"1", "2"}
    book.close()


def test_requests_and_orders_by_manufacturer_and_by_umbrella_brand(crm):
    req, data, _, _ = structured(crm)
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    _, execution = release(crm, req, calculation, "1")
    other = post(crm, "/requests", {"title": "Таблица клиента", "client_id": req["client_id"]})
    raw = preview(crm, other["id"], ["Name", "Producer", "Qty"],
                  [["Atropine", "aozeal", "2"], ["Caffeine", "Your", "1"]])
    confirm(crm, other["id"], raw["id"])

    report = get(crm, "/analytics/manufacturers")
    rows = {row["name"]: row for row in report["items"]}
    aozeal = rows["Aozeal"]
    assert (aozeal["requests"], aozeal["positions"], aozeal["quoted_positions"], aozeal["ordered_positions"]) == (
        2, 2, 1, 1)
    with crm["sessions"]() as db:
        total = Decimal(db.get(Execution, execution["id"]).snapshot["total"]).quantize(Decimal("0.01"))
    assert (aozeal["ordered_quantity"], aozeal["ordered_sales"]) == ("1", str(total))
    unknown = report["items"][-1]
    assert (unknown["name"], unknown["manufacturer_id"], unknown["positions"]) == ("Производитель не указан", None, 1)

    group = post(crm, "/manufacturers", {"name": "Aozeal Group"})
    member = get(crm, f"/manufacturers/{aozeal['manufacturer_id']}")
    change(crm, f"/manufacturers/{member['id']}", {"version": member["version"], "umbrella_id": group["id"]})
    brands = {row["name"]: row for row in get(crm, "/analytics/manufacturers?by=umbrella")["items"]}
    assert (brands["Aozeal Group"]["positions"], brands["Aozeal Group"]["ordered_positions"]) == (2, 1)
    assert get(crm, "/analytics/manufacturers")["items"][0]["umbrella_name"] == "Aozeal Group"
    assert get(crm, "/analytics/manufacturers?date_from=2000-01-01&date_to=2000-01-02")["items"] == []

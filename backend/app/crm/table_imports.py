"""Preview and atomic application of request tables and supplier quote spreadsheets."""

import csv
import io
import json
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, Depends, File, Form, Header, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.commerce.calculator import convert
from app.commerce.files import put_file
from app.core.config import settings
from app.core.db import get_db
from app.core.errors import DomainError
from app.core.models import User
from app.core.routes import Input
from app.core.security import check_request, current_user, require_permission
from app.core.service import advisory, audit, idem, lock, serialize
from app.crm.imports import string_value, xlsx
from app.crm.models import (
    Counterparty,
    Currency,
    Nomenclature,
    Packing,
    ProductGroup,
    Request,
    RequestItem,
    RequestItemRevision,
    TableImport,
)
from app.crm.routes import create_quote_sheet, packing_name, structured_item_fields
from app.crm.schemas import ItemInput, QuoteItemInput, QuoteSheetInput

router = APIRouter(tags=["Импорт заявок и квот"])
QUOTE_COLUMNS = {
    "source_row": "Позиция заявки",
    "supplier_code": "Код поставщика",
    "article": "Артикул",
    "name": "Наименование",
    "manufacturer": "Производитель",
    "cas": "CAS",
    "linear_formula": "Линейная формула",
    "purity": "Чистота",
    "packing_value": "Фасовка",
    "packing_unit": "Единица фасовки",
    "quantity": "Количество",
    "unit_price": "Цена",
    "currency": "Валюта",
    "delivery_days": "Срок поставки, дней",
    "product_group": "Товарная группа",
}
ALIASES = {
    "description": [
        "наименование",
        "название",
        "товар",
        "product name (english)",
        "product name",
        "name",
        "description",
    ],
    "article": ["артикул", "cat. no.", "cat no", "catalog number", "article", "sku"],
    "quantity": ["количество", "кол-во", "quantity", "qty"],
    "unit": ["единица", "ед.", "unit"],
    "packaging": ["фасовка", "packaging", "packing"],
    "manufacturer": ["производитель", "бренд", "manufacturer", "brand"],
    "source_row": ["позиция заявки", "номер позиции", "source row"],
    "supplier_code": ["код поставщика", "поставщик (код)", "supplier code"],
    "name": ["наименование", "название", "product name", "product name (english)", "name"],
    "cas": ["cas", "cas-номер", "cas no."],
    "linear_formula": ["линейная формула", "formula"],
    "purity": ["чистота", "purity"],
    "packing_value": ["фасовка", "объём фасовки", "packing value"],
    "packing_unit": ["единица фасовки", "packing unit"],
    "unit_price": ["цена", "цена за единицу", "unit price", "price"],
    "currency": ["валюта", "currency"],
    "delivery_days": ["срок поставки, дней", "срок поставки", "delivery days"],
    "product_group": ["товарная группа", "группа товаров", "группа", "код группы", "product group", "product_group"],
}
UNITS = {
    "шт": "pcs",
    "шт.": "pcs",
    "pcs": "pcs",
    "set": "pcs",
    "pack": "pcs",
    "g": "g",
    "kg": "kg",
    "mg": "mg",
    "ml": "ml",
    "l": "l",
    "г": "g",
    "мл": "ml",
}
MAX_ROWS = 10000


def read_table(data: bytes, filename: str, *, sheet_name: str | None = None):
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                files = archive.infolist()
                if (
                    len(files) > 10000
                    or sum(f.file_size for f in files) > 100 * 1024 * 1024
                    or any("vbaProject" in f.filename for f in files)
                ):
                    raise DomainError("IMPORT_UNSAFE", "Файл слишком большой или содержит макросы", 422)
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        except (BadZipFile, ValueError, KeyError, OSError):
            raise DomainError("IMPORT_INVALID_XLSX", "Не удалось прочитать XLSX", 422) from None
        try:
            sheet = book[sheet_name] if sheet_name and sheet_name in book.sheetnames else book.active
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 100:
                raise DomainError("IMPORT_LIMIT", "Максимум 10 000 строк и 100 столбцов", 422)
            matrix = []
            for cells in sheet.iter_rows():
                if len(matrix) > MAX_ROWS or len(cells) > 100:
                    raise DomainError("IMPORT_LIMIT", "Максимум 10 000 строк и 100 столбцов", 422)
                if any(cell.data_type == "f" for cell in cells):
                    raise DomainError("IMPORT_FORMULA", "Вставьте значения вместо формул", 422)
                matrix.append([string_value(cell) for cell in cells])
        finally:
            book.close()
    elif extension in (".txt", ".tsv", ".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise DomainError("IMPORT_ENCODING", "Сохраните текст в UTF-8", 422) from None
        lines = text.strip().splitlines()
        if lines and lines[0].strip().startswith("|") and "\t" not in lines[0]:
            matrix = [
                [value.strip() for value in line.strip().strip("|").split("|")]
                for line in lines
                if line.strip()
            ]
            if len(matrix) > 1 and all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in matrix[1]):
                matrix.pop(1)
        else:
            delimiter = (
                "\t"
                if "\t" in (lines[0] if lines else "")
                else ";"
                if ";" in (lines[0] if lines else "")
                else ","
            )
            matrix = list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
    else:
        raise DomainError("IMPORT_TYPE", "Поддерживаются XLSX, TSV, CSV и TXT", 422)
    if not matrix or len(matrix) < 2:
        raise DomainError("IMPORT_EMPTY", "Нужна строка заголовков и хотя бы одна позиция", 422)
    if len(matrix) > MAX_ROWS + 1 or len(matrix[0]) > 100:
        raise DomainError("IMPORT_LIMIT", "Максимум 10 000 строк и 100 столбцов", 422)
    columns = [str(value).strip() for value in matrix[0]]
    while (
        columns
        and not columns[-1]
        and all(len(row) <= len(columns) - 1 or not row[len(columns) - 1] for row in matrix[1:])
    ):
        columns.pop()
    if not columns or any(not col or len(col) > 250 for col in columns) or len(set(columns)) != len(columns):
        raise DomainError("IMPORT_HEADERS", "Заголовки должны быть заполнены и не повторяться", 422)
    rows = []
    for row in matrix[1:]:
        if not any(row):
            continue
        if any(len(str(cell)) > 10000 for cell in row) or any(row[len(columns) :]):
            raise DomainError("IMPORT_CELL", "Проверьте размер ячеек и число столбцов", 422)
        rows.append([str(row[i]) if i < len(row) else "" for i in range(len(columns))])
    if not rows:
        raise DomainError("IMPORT_EMPTY", "В таблице нет позиций", 422)
    return columns, rows


def infer_mapping(columns, mapping, kind):
    keys = (
        QUOTE_COLUMNS
        if kind == "quotes"
        else {key: key for key in ("description", "article", "quantity", "unit", "packaging", "cas")}
    )
    if not isinstance(mapping, dict) or any(
        key not in keys or not isinstance(value, str) or value and value not in columns
        for key, value in mapping.items()
    ):
        raise DomainError("IMPORT_MAPPING", "Выберите существующие столбцы таблицы", 422)
    result = {}
    for key in keys:
        if key in mapping:
            result[key] = mapping[key]
        else:
            result[key] = next((col for col in columns if col.casefold().strip() in ALIASES.get(key, [])), "")
    if kind == "items" and not result.get("description"):
        result["description"] = columns[0]
    return result


def values_for(columns, row, mapping):
    raw = dict(zip(columns, row, strict=True))
    return {key: raw.get(col, "").strip() for key, col in mapping.items()}


def optional(value):
    return None if value.strip() in ("", "0", "0.0", "0,0") else value.strip()


def group_reference(value):
    return " ".join(value.split()).casefold()


def decimal(value, *, positive=False, whole=False, places=8):
    try:
        number = Decimal(value.replace("\u00a0", "").replace(" ", "").replace(",", "."))
        if (
            not number.is_finite()
            or number < 0
            or positive
            and number <= 0
            or whole
            and number != number.to_integral_value()
        ):
            raise ValueError
        if (
            number.as_tuple().exponent < -places
            or len(number.as_tuple().digits) > 24
            or number >= Decimal("1e16")
        ):
            raise ValueError
        return str(number)
    except (InvalidOperation, ValueError):
        raise ValueError(
            "Укажите корректное " + ("положительное целое число" if whole else "число")
        ) from None


def request_plan(columns, rows, mapping):
    plan = []
    for i, row in enumerate(rows, 2):
        values = values_for(columns, row, mapping)
        errors = []
        fields = {key: value or None for key, value in values.items() if key != "quantity"}
        fields["description"] = fields.get("description") or " / ".join(cell for cell in row if cell)[:10000]
        quantity = values.get("quantity", "")
        if quantity:
            match = re.fullmatch(r"([\d.,\s]+)\s*([^\d.,\s]+)?", quantity)
            if match:
                try:
                    fields["quantity"] = decimal(match[1], positive=True, places=6)
                    fields["unit"] = UNITS.get(
                        (fields.get("unit") or match[2] or "pcs").casefold(), fields.get("unit") or match[2]
                    )
                except ValueError:
                    pass  # Raw client wording remains available even when it is not a number.
        try:
            fields = ItemInput(**fields).model_dump(mode="json")
        except ValueError as exc:
            errors.append(str(exc))
        plan.append(
            {"row_number": i, "action": "error" if errors else "checked", "data": fields, "errors": errors}
        )
    return plan


def quote_plan(db, request_id, columns, rows, mapping):
    items = list(
        db.scalars(
            select(RequestItem)
            .where(RequestItem.request_id == request_id, RequestItem.archived.is_(False), RequestItem.quote_only.is_(False))
            .order_by(RequestItem.created_at, RequestItem.id)
        )
    )
    catalogue = defaultdict(list)
    for n in db.scalars(select(Nomenclature)):
        if n.article:
            catalogue[n.article.strip().casefold()].append(n)
    packings = defaultdict(list)
    for p in db.scalars(select(Packing).where(Packing.active.is_(True))):
        packings[p.nomenclature_id].append(p)
    suppliers = {
        str(s.internal_code): s
        for s in db.scalars(
            select(Counterparty).where(
                Counterparty.kind.in_(["supplier", "both"]),
                Counterparty.archived.is_(False),
                Counterparty.client_base == "working",
            )
        )
    }
    currencies = {c.code: c for c in db.scalars(select(Currency).where(Currency.active.is_(True)))}
    product_groups = {group.id: group for group in db.scalars(select(ProductGroup))}
    group_names = defaultdict(dict)
    for group in product_groups.values():
        for reference in (group.name, group.slug, str(group.internal_code)):
            group_names[group_reference(reference)][group.id] = group

    def resolve_group(reference):
        matches = list(group_names[group_reference(reference)].values())
        active = [group for group in matches if group.active]
        if len(active) > 1:
            raise ValueError("Товарная группа неоднозначна. Укажите её уникальный код из справочника")
        if not active:
            raise ValueError(
                "Товарная группа неактивна" if matches else "Товарная группа не найдена в справочнике"
            )
        return active[0]

    seen = set()
    new_definitions = {}
    plan = []
    for i, raw in enumerate(rows, 2):
        values = values_for(columns, raw, mapping)
        errors = []
        data = {key: optional(value) for key, value in values.items()}
        entry = {"row_number": i, "action": "error", "data": data, "errors": errors}
        try:
            article = values.get("article", "")
            if not article or len(article) > 150:
                raise ValueError("Укажите артикул (до 150 символов), сохраняя начальные нули")
            data["article"] = article
            matches = catalogue[article.casefold()]
            if len(matches) > 1:
                raise ValueError("Артикул неоднозначен: в номенклатуре несколько совпадений")
            n = matches[0] if matches else None
            if n and not n.active:
                raise ValueError("Номенклатура с этим артикулом неактивна")
            s = suppliers.get(values.get("supplier_code", ""))
            if not s:
                raise ValueError("Код действующего поставщика не найден. Создайте поставщика в справочнике")
            c = currencies.get(values.get("currency", "").upper())
            if not c:
                raise ValueError("Валюта не найдена в справочнике")
            data["supplier_id"], data["currency_id"] = s.id, c.id
            data["supplier_name"] = s.name
            data["quantity"] = decimal(values.get("quantity", ""), positive=True, whole=True, places=6)
            data["unit_price"] = decimal(values.get("unit_price", ""))
            days = decimal(values.get("delivery_days", ""), whole=True)
            if Decimal(days) > 2147483647:
                raise ValueError("Срок поставки слишком большой")
            data["delivery_days"] = int(Decimal(days))
            source_row = values.get("source_row", "")
            if source_row:
                if not source_row.isdigit() or not 1 <= int(source_row) <= len(items):
                    raise ValueError("Номер позиции заявки не найден")
                source = items[int(source_row) - 1]
            else:
                candidates = [
                    item
                    for item in items
                    if (item.article or "").strip().casefold() == article.casefold()
                    or n
                    and item.nomenclature_id == n.id
                ]
                if len(candidates) > 1:
                    raise ValueError("Укажите номер позиции заявки: найдено несколько строк")
                source = candidates[0] if candidates else RequestItem(quantity=Decimal(data["quantity"]), unit="pcs")
            identity = (s.id, source.id, article.casefold(), data.get("packing_value"), data.get("packing_unit"))
            if identity in seen:
                raise ValueError("Позиция заявки повторяется для одного поставщика")
            seen.add(identity)
            data["source_request_item_id"], data["source_version"] = source.id, source.version
            data["nomenclature_id"] = n.id if n else None
            selected_group = resolve_group(data["product_group"]) if data.get("product_group") else None
            if n:
                group = product_groups.get(n.product_group_id)
            elif selected_group:
                group = selected_group
            elif source.product_group_id:
                group = product_groups.get(source.product_group_id)
            else:
                group = next((group for group in product_groups.values() if group.slug == "other"), None)
            if (not n or n.product_group_id) and (not group or not group.active):
                raise ValueError("Выберите действующую товарную группу")
            data["product_group_id"] = group.id if group else None
            data["product_group_name"] = group.name if group else None
            data["product_group_slug"] = group.slug if group else None
            if not n:
                if not data.get("name"):
                    raise ValueError("Для новой номенклатуры заполните наименование")
                limits = {"name": 250, "manufacturer": 250, "cas": 30, "linear_formula": 500, "purity": 200}
                if any(len(data.get(key) or "") > limit for key, limit in limits.items()):
                    raise ValueError("Слишком длинное описание новой номенклатуры")
            existing_packings = packings[n.id] if n else []
            packing = None
            if data.get("packing_value"):
                data["packing_value"] = decimal(data["packing_value"], positive=True, places=6)
                data["packing_unit"] = UNITS.get(
                    (data.get("packing_unit") or "").casefold(), data.get("packing_unit")
                )
                if data["packing_unit"] not in ("g", "kg", "mg", "l", "ml", "pcs"):
                    raise ValueError("Укажите единицу фасовки: g, kg, mg, l, ml или pcs")
                packing = next(
                    (
                        p
                        for p in existing_packings
                        if p.value == Decimal(data["packing_value"]) and p.unit == data["packing_unit"]
                    ),
                    None,
                )
            else:
                packing = next((p for p in existing_packings if p.id == source.packing_id), None)
                if not packing and len(existing_packings) > 1:
                    raise ValueError("Укажите фасовку: у номенклатуры несколько вариантов")
                packing = packing or (existing_packings[0] if existing_packings else None)
                data["packing_value"], data["packing_unit"] = (
                    (str(packing.value), packing.unit) if packing else ("1", "pcs")
                )
            data["packing_id"] = packing.id if packing else None
            source_quantity = source.quantity
            if source_quantity is not None and source.unit and source.unit != "pcs":
                try:
                    source_quantity = convert(source_quantity, source.unit, data["packing_unit"]) / Decimal(
                        data["packing_value"]
                    )
                except DomainError:
                    raise ValueError(
                        "Единица потребности не соответствует фасовке. Уточните позицию заявки"
                    ) from None
            data["source_quantity"] = decimal(
                str(source_quantity if source_quantity is not None else data["quantity"]),
                positive=True,
                whole=True,
                places=6,
            )
            if not n:
                definition = {
                    key: data.get(key)
                    for key in ("name", "manufacturer", "cas", "linear_formula", "purity", "product_group_id")
                }
                old = new_definitions.setdefault(article.casefold(), definition)
                if old != definition:
                    raise ValueError("Для одного нового артикула указаны разные характеристики")
            entry["action"] = "checked" if n else "create_nomenclature"
        except ValueError as exc:
            errors.append(str(exc))
        plan.append(entry)
    return plan


def summarize(plan):
    return {
        "total": len(plan),
        "checked": sum(row["action"] == "checked" for row in plan),
        "create_nomenclature": sum(row["action"] == "create_nomenclature" for row in plan),
        "errors": sum(bool(row["errors"]) for row in plan),
    }


def batch_view(batch):
    return {**serialize(batch), "rows": batch.plan[:100], "plan": None, "preview_limit": 100}


@router.post("/requests/{request_id}/table-imports/preview")
def preview(
    request_id: str,
    kind: str = Form("items"),
    file: UploadFile = File(),
    mapping: str = Form("{}"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    if kind not in ("items", "quotes"):
        raise DomainError("IMPORT_KIND", "Неизвестный вид импорта", 422)
    request = check_request(db, user, request_id, "quotes.write" if kind == "quotes" else "requests.write")
    if request.archived:
        raise DomainError("REQUEST_ARCHIVED", "Заявка в архиве", 409)
    if kind == "quotes":
        require_permission(db, user, "finance.purchase.read", request_id)
    content = file.file.read(settings.max_file_size + 1)
    if len(content) > settings.max_file_size:
        raise DomainError("FILE_TOO_LARGE", "Файл превышает допустимый размер", 413)
    if kind == "quotes" and not (file.filename or "").lower().endswith(".xlsx"):
        raise DomainError("IMPORT_TYPE", "Квоты загружаются из XLSX", 422)
    try:
        chosen = json.loads(mapping)
        columns, rows = read_table(
            content, file.filename or "", sheet_name="Квоты" if kind == "quotes" else None
        )
    except (json.JSONDecodeError, csv.Error):
        raise DomainError(
            "IMPORT_INVALID", "Проверьте формат таблицы и сопоставление столбцов", 422
        ) from None
    chosen = infer_mapping(columns, chosen, kind)
    plan = (
        quote_plan(db, request_id, columns, rows, chosen)
        if kind == "quotes"
        else request_plan(columns, rows, chosen)
    )
    batch = TableImport(
        request_id=request_id,
        author_id=user.id,
        kind=kind,
        source_name=Path(file.filename).name[:250],
        file_metadata=put_file(content, Path(file.filename).suffix.lstrip(".")),
        columns=columns,
        rows=rows,
        mapping=chosen,
        plan=plan,
        summary=summarize(plan),
    )
    if kind == 'items' and Path(file.filename or '').suffix.lower() == '.xlsx':
        # Retain row references to the immutable original instead of duplicating style XML per item.
        source_book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            source_sheet = source_book.active
            batch.file_metadata = {**batch.file_metadata, 'source_sheet': source_sheet.title,
                                   'source_rows': [index for index, cells in enumerate(
                                       source_sheet.iter_rows(min_row=2), 2)
                                       if any(string_value(cell) for cell in cells)]}
        finally:
            source_book.close()
    db.add(batch)
    db.flush()
    audit(db, user, "table_import", batch.id, "preview", after=batch.summary)
    return batch_view(batch)


def owned_batch(db, user, request_id, batch_id):
    batch = db.get(TableImport, batch_id)
    if not batch or batch.request_id != request_id or batch.author_id != user.id:
        raise DomainError("NOT_FOUND", "Импорт недоступен", 404)
    check_request(db, user, request_id, "quotes.write" if batch.kind == "quotes" else "requests.write")
    if batch.kind == "quotes":
        require_permission(db, user, "finance.purchase.read", request_id)
    return batch


@router.get("/requests/{request_id}/table-imports/{batch_id}/errors.xlsx")
def error_file(
    request_id: str, batch_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    batch = owned_batch(db, user, request_id, batch_id)
    require_permission(db, user, "exports.download", request_id)
    rows = [
        [entry["row_number"], "; ".join(entry["errors"]), *batch.rows[i]]
        for i, entry in enumerate(batch.plan)
        if entry["errors"]
    ]
    return Response(
        xlsx([["Строка", "Ошибки", *batch.columns], *rows]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/requests/{request_id}/quote-import/template.xlsx")
def quote_template(request_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    check_request(db, user, request_id, "quotes.write")
    require_permission(db, user, "exports.download", request_id)
    rows = []
    items = db.scalars(
        select(RequestItem)
        .where(RequestItem.request_id == request_id, RequestItem.archived.is_(False), RequestItem.quote_only.is_(False))
        .order_by(RequestItem.created_at, RequestItem.id)
    )
    for index, item in enumerate(items, 1):
        n = db.get(Nomenclature, item.nomenclature_id) if item.nomenclature_id else None
        p = db.get(Packing, item.packing_id) if item.packing_id else None
        group_id = n.product_group_id if n else item.product_group_id
        group = db.get(ProductGroup, group_id) if group_id else None
        values = {
            "source_row": str(index),
            "article": n.article if n else item.article,
            "name": n.name if n else item.description,
            "manufacturer": n.manufacturer if n else "",
            "cas": n.cas if n else item.cas,
            "packing_value": p.value if p else "0",
            "packing_unit": p.unit if p else "0",
            "quantity": item.quantity,
            "product_group": group.name if group else "",
        }
        rows.append([values.get(key, "") for key in QUOTE_COLUMNS])
    book = load_workbook(io.BytesIO(xlsx([list(QUOTE_COLUMNS.values()), *rows])))
    sheet = book.active
    sheet.title = "Квоты"
    directory = book.create_sheet("Товарные группы")
    directory.append(["Товарная группа", "Код группы", "Системный код"])
    groups = list(
        db.scalars(select(ProductGroup).where(ProductGroup.active.is_(True)).order_by(ProductGroup.name))
    )
    for group in groups:
        directory.append([group.name, str(group.internal_code), group.slug])
    for row in directory:
        for cell in row:
            cell.data_type = "s"
            cell.number_format = "@"
    directory.column_dimensions["A"].width = 40
    directory.column_dimensions["B"].width = 28
    directory.freeze_panes = "A2"
    if groups:
        book.defined_names.add(
            DefinedName("QuoteProductGroups", attr_text=f"'Товарные группы'!$A$2:$A${len(groups) + 1}")
        )
        validation = DataValidation(type="list", formula1="QuoteProductGroups", allow_blank=True)
        validation.showInputMessage = True
        validation.promptTitle = "Группа новой номенклатуры"
        validation.prompt = "Выберите группу из списка или укажите её код с листа «Товарные группы»."
        # Codes and 0 are valid typed inputs too; the import preview validates every value.
        validation.showErrorMessage = False
        sheet.add_data_validation(validation)
        column = sheet.cell(1, len(QUOTE_COLUMNS)).column_letter
        validation.add(f"{column}2:{column}{MAX_ROWS + 1}")
        sheet.column_dimensions[column].width = 40
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return Response(
        stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.post("/requests/{request_id}/table-imports/{batch_id}/confirm")
def confirm(
    request_id: str,
    batch_id: str,
    body: Input,
    idempotency_key: str | None = Header(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    owned_batch(db, user, request_id, batch_id)

    def operation():
        parent = lock(db, Request, request_id)
        batch = lock(db, TableImport, batch_id)
        if batch.status == "completed":
            return batch_view(batch)
        if parent.archived:
            raise DomainError("REQUEST_ARCHIVED", "Заявка в архиве", 409)
        if batch.summary["errors"]:
            raise DomainError("IMPORT_ERRORS", "Исправьте строки с ошибками и повторите проверку файла", 422)
        previous = db.scalars(
            select(TableImport).where(
                TableImport.request_id == request_id,
                TableImport.kind == batch.kind,
                TableImport.status == "completed",
                TableImport.id != batch.id,
            )
        )
        duplicate = next(
            (
                row
                for row in previous
                if (row.file_metadata["sha256"] == batch.file_metadata["sha256"]
                    or (batch.kind == "items" and row.columns == batch.columns and row.rows == batch.rows))
                and {key: value for key, value in row.mapping.items() if value}
                == {key: value for key, value in batch.mapping.items() if value}
            ),
            None,
        )
        if duplicate:
            batch.status = "completed"
            batch.result = {**duplicate.result, "already_imported": True}
            audit(db, user, "table_import", batch.id, "duplicate_skipped", after=batch.result)
            return batch_view(batch)
        created = 0
        if batch.kind == "items":
            for index, (entry, raw) in enumerate(zip(batch.plan, batch.rows, strict=True)):
                fields = structured_item_fields(db, ItemInput(**entry["data"]).model_dump())
                item = RequestItem(
                    request_id=request_id, **fields, source_columns=batch.columns, source_values=raw,
                    source_format={'import_id': batch.id, 'row_number': batch.file_metadata['source_rows'][index]}
                    if batch.file_metadata.get('source_rows') else {},
                )
                db.add(item)
                db.flush()
                snapshot = serialize(item)
                db.add(RequestItemRevision(item_id=item.id, revision=1, snapshot=snapshot, author_id=user.id))
                audit(
                    db, user, "request_item", item.id, "imported", after=snapshot, reason=f"Импорт {batch.id}"
                )
        else:
            advisory(db, "catalog.article-import")
            fresh = quote_plan(db, request_id, batch.columns, batch.rows, batch.mapping)
            if fresh != batch.plan:
                raise DomainError(
                    "IMPORT_STALE", "Данные изменились после проверки. Загрузите файл ещё раз", 409
                )
            if any(not row["data"]["nomenclature_id"] or not row["data"]["packing_id"] for row in fresh):
                require_permission(db, user, "catalog.write")
            groups = defaultdict(list)
            new_catalogue, new_packings, updated_sources = {}, {}, set()
            for entry in fresh:
                data = entry["data"]
                n_id = data["nomenclature_id"] or new_catalogue.get(data["article"].casefold())
                if not n_id:
                    n = Nomenclature(
                        **{
                            key: data.get(key)
                            for key in ("name", "article", "manufacturer", "cas", "linear_formula", "purity")
                        },
                        product_group_id=data["product_group_id"],
                    )
                    db.add(n)
                    db.flush()
                    n_id = n.id
                    new_catalogue[data["article"].casefold()] = n_id
                    created += 1
                    audit(db, user, "nomenclature", n.id, "imported", after=serialize(n))
                packing_id = data["packing_id"]
                packing_key = (n_id, str(Decimal(data["packing_value"]).normalize()), data["packing_unit"])
                if not packing_id:
                    packing_id = new_packings.get(packing_key)
                if not packing_id:
                    p = Packing(
                        nomenclature_id=n_id,
                        value=Decimal(data["packing_value"]),
                        unit=data["packing_unit"],
                        display_name=packing_name(Decimal(data["packing_value"]), data["packing_unit"]),
                    )
                    db.add(p)
                    db.flush()
                    packing_id = p.id
                    new_packings[packing_key] = packing_id
                    audit(db, user, "packing", p.id, "imported", after=serialize(p))
                source = db.get(RequestItem, data["source_request_item_id"]) if data["source_request_item_id"] else None
                if source is None:
                    source = RequestItem(request_id=request_id, description=data.get("name") or data["article"],
                                         nomenclature_id=n_id, packing_id=packing_id,
                                         product_group_id=data["product_group_id"], quantity=Decimal(data["quantity"]),
                                         unit="pcs", quote_only=True)
                    db.add(source)
                    db.flush()
                if source.id not in updated_sources and (
                    not source.nomenclature_id
                    or not source.packing_id
                    or source.quantity is None
                    or source.unit != "pcs"
                ):
                    fields = structured_item_fields(
                        db,
                        {
                            "nomenclature_id": n_id,
                            "packing_id": packing_id,
                            "quantity": Decimal(data["source_quantity"]),
                        },
                        source,
                    )
                    before = serialize(source)
                    for key, value in fields.items():
                        setattr(source, key, value)
                    source.version += 1
                    source.revision += 1
                    db.add(
                        RequestItemRevision(
                            item_id=source.id,
                            revision=source.revision,
                            snapshot=serialize(source),
                            author_id=user.id,
                            reason=f"Уточнение по квоте {batch.id}",
                        )
                    )
                    audit(db, user, "request_item", source.id, "structured", before, serialize(source))
                    updated_sources.add(source.id)
                groups[data["supplier_id"]].append(
                    QuoteItemInput(
                        source_request_item_id=source.id,
                        nomenclature_id=n_id,
                        packing_id=packing_id,
                        quantity=data["quantity"],
                        unit_price=data["unit_price"],
                        currency_id=data["currency_id"],
                        delivery_days=data["delivery_days"],
                    )
                )
            sheet_ids = []
            previous_batch = db.info.get("quote_import_batch")
            db.info["quote_import_batch"] = batch.id
            try:
                for supplier_id, entries in groups.items():
                    # Use the same invariants and audit trail as the regular quote editor.
                    for offset in range(0, len(entries), 100):
                        result = create_quote_sheet(
                            request_id,
                            QuoteSheetInput(supplier_id=supplier_id, items=entries[offset : offset + 100]),
                            None,
                            user,
                            db,
                        )
                        sheet_ids.append(result["id"])
            finally:
                if previous_batch is None:
                    db.info.pop("quote_import_batch", None)
                else:
                    db.info["quote_import_batch"] = previous_batch
            batch.result = {"quote_sheet_ids": sheet_ids}
        parent.version += 1
        batch.status = "completed"
        batch.result = {**batch.result, "imported": len(batch.rows), "nomenclatures_created": created}
        audit(db, user, "table_import", batch.id, "completed", after=batch.result)
        return batch_view(batch)

    return idem(db, user, idempotency_key, f"table-import:{batch_id}", {}, operation)

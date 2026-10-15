"""October 6 acceptance: quantity, procurement handoff, forecasts and immutable exports."""

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.commerce.files import read_file
from app.commerce.models import CommercialDocument, Execution, WaveAllocation
from app.core.models import Notification
from app.crm.models import QuoteItem, Request, RequestItem
from app.crm.table_imports import QUOTE_COLUMNS
from tests.support.business import (
    cmd,
    confirm,
    handoff,
    peer_request,
    preview,
    quote_file,
    release,
    structured,
    wave_row,
)


@pytest.mark.parametrize("quantity", ["2", "100"])
def test_quantity_is_independent_of_quote_and_initial_demand(crm, quantity):  # noqa: F811
    req, data, seller, _ = structured(crm)
    data["selections"][0]["quantity"] = quantity
    quote_id = data["selections"][0]["quote_item_id"]
    seller = crm["client"].patch(f"/api/v1/sellers/{seller['id']}", json={
        "version": seller["version"], "name": seller["name"], "currency": seller["currency"],
        "details": {**seller["details"], "Подписант": "Иван Иванов", "Должность подписанта": "Директор"},
    }).json()
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    assert Decimal(calculation["snapshot"]["totals"]["purchase_rub"]) == Decimal(quantity) * 1000
    proposal, execution = release(crm, req, calculation, quantity)
    assert proposal["snapshot"]["lines"][0]["description"] == "Стандартный образец Sample (100 mg) Арт. A123"
    assert proposal["snapshot"]["signature"] == {"name": "Гильмутдинов Т.Ф", "position": "Генеральный директор"}
    assert "bank_details" not in proposal["snapshot"]
    invoice = cmd(crm, f"/requests/{req['id']}/invoices", {
        "proposal_id": proposal["id"], "due_date": (date.today() + timedelta(days=30)).isoformat(),
        "terms": "Отсрочка", "lines": [{"execution_id": execution["id"], "quantity": quantity}],
    })
    assert Decimal(invoice["snapshot"]["lines"][0]["quantity"]) == Decimal(quantity)
    assert invoice["snapshot"]["signature"] == proposal["snapshot"]["signature"]
    assert invoice["snapshot"]["bank_details"]["Банк"] == "Банк A"
    with crm["sessions"]() as db:
        quote = db.get(QuoteItem, quote_id)
        assert quote.quantity == 1
        assert db.get(RequestItem, quote.source_request_item_id).quantity == 1
        proposal_book = load_workbook(io.BytesIO(read_file(db.get(CommercialDocument, proposal["id"]).files["xlsx"])))
        invoice_book = load_workbook(io.BytesIO(read_file(db.get(CommercialDocument, invoice["id"]).files["xlsx"])))
        proposal_values = [cell.value for row in proposal_book.active for cell in row]
        invoice_values = [cell.value for row in invoice_book.active for cell in row]
        assert "Банк A" not in proposal_values
        assert "Банк A" in invoice_values
        assert "Генеральный директор" in proposal_values and "Гильмутдинов Т.Ф" in proposal_values
        # Business remark of 15 October 2026: the invoice is laid out and signed like the proposal.
        assert "Генеральный директор" in invoice_values and "Гильмутдинов Т.Ф" in invoice_values
        assert "Директор" not in invoice_values and "Иван Иванов" not in invoice_values
        assert "Заявка" not in invoice_values and "Цена за единицу с НДС" in invoice_values
        assert "Наименование" in proposal_values and "Цена за единицу с НДС" in proposal_values
        assert "Без налога" not in proposal_values
    assert crm["client"].get("/api/v1/supplier-orders/positions").json()["items"] == []
    handoff(crm, req, execution)
    positions = crm["client"].get("/api/v1/supplier-orders/positions").json()["items"]
    assert Decimal(positions[0]["quantity"]) == Decimal(quantity)
    cmd(crm, "/supplier-orders", {"execution_ids": [execution["id"]], "seller_id": seller["id"],
        "expected_date": (date.today() + timedelta(days=30)).isoformat()})
    # The released line still cannot be accepted twice beyond its own quantity.
    response = cmd(crm, f"/proposals/{proposal['id']}/accept", {
        "version": proposal["version"] + 1, "reason": "Повторное принятие",
        "lines": [{"line_id": calculation["snapshot"]["lines"][0]["line_id"], "quantity": "1"}],
    }, status=422)
    assert response["code"] == "QUANTITY_EXCEEDED"


def test_planned_wave_calculates_on_its_volume_and_confirmed_orders_go_to_its_working_wave(crm):  # noqa: F811
    req, data, seller, supplier = structured(crm, mode="IMPORT", wave_kind="plan")
    with crm["sessions"]() as db:
        quote = db.get(QuoteItem, data["selections"][0]["quote_item_id"])
        item = db.get(RequestItem, quote.source_request_item_id)
        group_id = item.product_group_id
        item_data = {"nomenclature_id": quote.nomenclature_id, "packing_id": quote.packing_id,
            "quantity": "1", "unit_price": "1000", "currency_id": quote.currency_id, "delivery_days": 7}
    plan = wave_row(crm, req["wave_id"])
    forecast_data = {"product_group_id": group_id, "target_quantity": "10", "unit_price_rub": "1000",
        "weight_per_unit": "1", "reason": "Минимальный объём запуска"}
    current = cmd(crm, f"/waves/{plan['id']}/forecasts", {**forecast_data, "version": plan["version"]}, method="put")
    assert Decimal(current["financial_summary"]["real_quantity"]) == 0
    assert Decimal(current["financial_summary"]["forecast_quantity"]) == 10
    data["selections"][0]["quantity"] = "2"
    data["expenses"] = [{"name": "Общая логистика", "amount": "1000", "currency": "RUB", "scope": "WAVE",
        "stage": "INTERNATIONAL_LOGISTICS", "method": "BY_QUANTITY"}]
    current = cmd(crm, f"/waves/{plan['id']}/budget", {"version": current["version"], "profile_id": data["profile_id"],
        "expenses": data["expenses"], "reason": "Бюджет запуска"}, method="put")
    first = cmd(crm, f"/requests/{req['id']}/calculations", data)
    assert Decimal(first["snapshot"]["wave_distribution"]["total_quantity"]) == 10
    assert Decimal(first["snapshot"]["wave_distribution"]["forecast_quantity"]) == 8
    assert Decimal(first["snapshot"]["totals"]["international_logistics"]) == 200
    proposal, execution = release(crm, req, first, "2")
    assert wave_row(crm, plan["id"])["working_waves"] == []
    assert crm["client"].get("/api/v1/supplier-orders/positions").json()["items"] == []
    handoff(crm, req, execution)
    # The plan keeps planning; the confirmed order went to a working wave copied from it.
    plan = wave_row(crm, plan["id"])
    assert plan["allocations"] == []
    assert Decimal(plan["financial_summary"]["real_quantity"]) == 0
    assert Decimal(plan["financial_summary"]["forecast_quantity"]) == 10
    assert Decimal(plan["forecasts"][0]["target_quantity"]) == 10
    assert Decimal(plan["forecasts"][0]["filled_quantity"]) == 2
    assert Decimal(plan["forecasts"][0]["remaining_quantity"]) == 8
    [working] = plan["working_waves"]
    work = wave_row(crm, working["id"])
    assert (work["kind"], work["plan_id"], work["plan_number"]) == ("work", plan["id"], plan["number"])
    assert work["number"] == f"{plan['number']} — рабочая 1" and work["status"] == "assembling"
    assert (work["route"], work["departure_week"], work["owner_id"]) == (plan["route"], plan["departure_week"], plan["owner_id"])
    assert work["budget_expenses"] == current["budget_expenses"]
    assert work["forecasts"] == [] and len(work["allocations"]) == 1
    assert Decimal(work["financial_summary"]["real_quantity"]) == 2
    assert Decimal(work["financial_summary"]["forecast_quantity"]) == 0
    assert work["financial_summary"]["actual"]["status"] == "current"
    assert crm["client"].get(f"/api/v1/requests/{req['id']}").json()["wave_id"] == plan["id"]
    # Another client's calculation on the plan still divides the costs by the planned volume...
    peer, peer_quote = peer_request(crm, seller, supplier, item_data, plan["id"])
    peer_payload = {"request_version": peer["version"], "profile_id": data["profile_id"], "delivery_days": 14,
        "selections": [{"quote_item_id": peer_quote["id"], "quantity": "2"}], "delivery_required": False}
    previewed = cmd(crm, f"/requests/{peer['id']}/calculations/preview", peer_payload)["snapshot"]
    distribution = previewed["wave_distribution"]
    assert Decimal(distribution["existing_quantity"]) == 0
    assert Decimal(distribution["selected_quantity"]) == 2
    assert Decimal(distribution["forecast_quantity"]) == 8
    assert Decimal(distribution["total_quantity"]) == 10
    assert Decimal(previewed["totals"]["international_logistics"]) == 200
    assert sum(Decimal(x) for x in previewed["expense_allocations"][0]["wave_parts"].values()) == 1000
    # ...and on the working wave by its real orders only.
    other, other_quote = peer_request(crm, seller, supplier, item_data, work["id"])
    real = cmd(crm, f"/requests/{other['id']}/calculations/preview", {**peer_payload,
        "request_version": other["version"], "selections": [{"quote_item_id": other_quote["id"], "quantity": "2"}]})["snapshot"]
    assert Decimal(real["wave_distribution"]["existing_quantity"]) == 2
    assert Decimal(real["wave_distribution"]["forecast_quantity"]) == 0
    assert Decimal(real["wave_distribution"]["total_quantity"]) == 4
    assert Decimal(real["totals"]["international_logistics"]) == 500
    # A wave version/forecast change invalidates a preview without mutating released documents.
    cmd(crm, f"/waves/{plan['id']}/forecasts", {**forecast_data, "version": plan["version"], "active": False}, method="put")
    conflict = cmd(crm, f"/requests/{peer['id']}/calculations", {**peer_payload,
        "expected_wave_digest": distribution["digest"]}, status=409)
    assert conflict["code"] == "WAVE_CHANGED"
    fresh = cmd(crm, f"/requests/{peer['id']}/calculations/preview", peer_payload)["snapshot"]
    assert Decimal(fresh["wave_distribution"]["total_quantity"]) == 2
    assert Decimal(fresh["totals"]["international_logistics"]) == 1000
    with crm["sessions"]() as db:
        assert db.get(CommercialDocument, proposal["id"]).snapshot == proposal["snapshot"]
        assert db.get(Execution, execution["id"]).procurement_at is not None
        assert db.get(Execution, execution["id"]).financing_deficit is True
        assert db.query(WaveAllocation).filter_by(wave_id=plan["id"]).count() == 0
        assert db.query(WaveAllocation).filter_by(wave_id=work["id"], active=True).count() == 1


def test_large_quote_import_notifies_manager_once_and_accepts_extra_lines(crm):  # noqa: F811
    req, data, _, supplier = structured(crm)
    with crm["sessions"].begin() as db:
        db.get(Request, req["id"]).owner_id = crm["manager"].id
    rows = [quote_file(supplier["internal_code"], {"source_row": "", "article": f"EXTRA-{i:03}",
        "name": f"Дополнительный стандарт {i}", "packing_value": "50", "packing_unit": "mg",
        "product_group": "Стандартные образцы"}) for i in range(125)]
    batch = preview(crm, req["id"], list(QUOTE_COLUMNS.values()), rows, kind="quotes")
    assert batch["summary"]["errors"] == 0, batch["rows"][0]
    confirmed = confirm(crm, req["id"], batch["id"], key="same-batch")
    assert confirmed["result"]["imported"] == 125
    confirm(crm, req["id"], batch["id"], key="same-batch")
    with crm["sessions"]() as db:
        notifications = db.scalars(select(Notification).where(Notification.event_key == f"quote-import:{batch['id']}")).all()
        assert len(notifications) == 1 and notifications[0].user_id == crm["manager"].id
        assert db.query(RequestItem).filter_by(request_id=req["id"], quote_only=True).count() == 126
    all_quotes = crm["client"].get(f"/api/v1/requests/{req['id']}/quote-items?all=true").json()
    assert len(all_quotes["items"]) == 126
    latest = crm["client"].get(f"/api/v1/requests/{req['id']}").json()
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", {**data, "request_version": latest["version"],
        "selections": [{"quote_item_id": row["id"], "quantity": "2"} for row in all_quotes["items"]]})
    assert len(calculation["snapshot"]["lines"]) == 126


@pytest.mark.parametrize("method", ["BY_QUANTITY", "BY_PURCHASE_VALUE", "BY_WEIGHT"])
def test_forecast_budget_bases_and_overfilled_plan(crm, method):  # noqa: F811
    req, data, _, _ = structured(crm, mode="IMPORT", wave_kind="plan")
    with crm["sessions"]() as db:
        quote = db.get(QuoteItem, data["selections"][0]["quote_item_id"])
        group_id = db.get(RequestItem, quote.source_request_item_id).product_group_id
    initial = wave_row(crm, req["wave_id"])
    payload = {"version": initial["version"], "product_group_id": group_id,
        "target_quantity": "10", "unit_price_rub": "1000", "weight_per_unit": "1", "reason": "Запуск поставщика"}
    current = cmd(crm, f"/waves/{initial['id']}/forecasts", payload, method="put", key="forecast-retry")
    assert cmd(crm, f"/waves/{initial['id']}/forecasts", payload, method="put", key="forecast-retry") == current
    assert len(wave_row(crm, initial["id"])["forecasts"]) == 1
    assert cmd(crm, f"/waves/{initial['id']}/forecasts", payload, method="put", status=409)["code"] == "VERSION_CONFLICT"
    data["expenses"] = [{"name": "Транспорт", "amount": "1000", "scope": "WAVE", "stage": "GENERAL",
        "method": method, "currency": "RUB"}]
    data["selections"][0].update(quantity="2", weight="2")
    calculated = cmd(crm, f"/requests/{req['id']}/calculations/preview", data)["snapshot"]
    assert Decimal(calculated["totals"]["general_expenses"]) == 200
    assert Decimal(calculated["wave_distribution"]["forecast_quantity"]) == 8
    data["selections"][0].update(quantity="12", weight="12")
    overfilled = cmd(crm, f"/requests/{req['id']}/calculations/preview", data)["snapshot"]
    assert Decimal(overfilled["wave_distribution"]["forecast_quantity"]) == 0
    assert Decimal(overfilled["wave_distribution"]["total_quantity"]) == 12
    assert Decimal(overfilled["totals"]["general_expenses"]) == 1000

    current = cmd(crm, f"/waves/{initial['id']}/budget", {"version": current["version"],
        "profile_id": data["profile_id"], "expenses": data["expenses"], "reason": "Бюджет по выбранной базе"}, method="put")
    data["selections"][0].update(quantity="2", weight="2")
    saved = cmd(crm, f"/requests/{req['id']}/calculations", data)
    _, execution = release(crm, req, saved, "2")
    handoff(crm, req, execution)
    plan = wave_row(crm, initial["id"])
    assert Decimal(plan["financial_summary"]["real_quantity"]) == 0
    assert Decimal(plan["financial_summary"]["forecast_quantity"]) == 10
    summary = wave_row(crm, plan["working_waves"][0]["id"])["financial_summary"]
    assert summary["actual"]["status"] == "current"
    assert Decimal(summary["real_quantity"]) == 2
    assert Decimal(summary["forecast_quantity"]) == 0

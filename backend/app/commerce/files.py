"""Immutable generated documents in protected content-addressed storage."""

import hashlib
import io
import os
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import settings
from app.core.errors import error


def put_file(content: bytes, extension: str) -> dict:
    checksum = hashlib.sha256(content).hexdigest()
    key = f"generated/{checksum[:2]}/{checksum}.{extension}"
    target = settings.storage_dir / key
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        # Atomic publication; an interrupted write never replaces an existing document.
        temporary = target.with_suffix(f".{os.getpid()}.tmp")
        temporary.write_bytes(content)
        os.replace(temporary, target)
    return {"key": key, "sha256": checksum, "size": len(content), "format": extension}


def read_file(metadata: dict) -> bytes:
    target = (settings.storage_dir / metadata["key"]).resolve()
    if not target.is_relative_to(settings.storage_dir.resolve()):
        error("FILE_NOT_FOUND", "Файл не найден", 404)
    try:
        content = target.read_bytes()
    except OSError:
        error("FILE_UNAVAILABLE", "Файл временно недоступен. Обратитесь к администратору", 503)
    if hashlib.sha256(content).hexdigest() != metadata["sha256"]:
        error(
            "FILE_INTEGRITY",
            "Контрольная сумма документа не совпадает. Восстановите файл из резервной копии",
            503,
        )
    return content


def safe_cell(value):
    # Prevent spreadsheet formula injection from product names and user remarks.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def format_details(value: object) -> str:
    if isinstance(value, dict):
        parts = [
            f"{key}: {item}"
            for key, item in value.items()
            if item is not None and str(item).strip()
        ]
        return " · ".join(parts) or "—"
    return str(value) if value else "—"


def workbook(headers: list[str], rows: list[list], title: str) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.append(headers)
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="213D37")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            45, max(13, max(len(str(cell.value or "")) for cell in column) + 2)
        )
        for cell in column[1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.number_format = "@"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def document_files(snapshot: dict) -> dict:
    rows = [
        [
            str(index),
            row["description"],
            row.get("cas", ""),
            row["quantity"],
            row["unit"],
            row["unit_price"],
            row["net"],
            row["tax"],
            row["total"],
        ]
        for index, row in enumerate(snapshot["lines"], 1)
    ]
    headers = [
        "№",
        "Наименование и характеристики",
        "CAS",
        "Кол-во",
        "Ед.",
        "Цена без налога",
        "Без налога",
        "Налог",
        "Итого",
    ]
    xlsx_rows = [
        [snapshot["title"], snapshot["number"], snapshot["date"]],
        ["Продавец", snapshot["seller"]["name"], format_details(snapshot["seller"].get("details"))],
        ["Клиент", snapshot["client"]["name"], format_details(snapshot["client"].get("details"))],
        ["Валюта", snapshot["currency"]],
        [],
        headers,
        *rows,
        [],
        [
            "Итого",
            "",
            "",
            "",
            "",
            "",
            snapshot["totals"]["net"],
            snapshot["totals"]["tax"],
            snapshot["totals"]["total"],
        ],
        ["Условия", snapshot["terms"]],
        ["Действует / оплатить до", snapshot["valid_until"]],
    ]
    xlsx = workbook([snapshot["title"]], xlsx_rows, "Документ")
    font_name = "CRMUnicode"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        candidates = [
            os.getenv("CRM_PDF_FONT", ""),
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/Library/Fonts/Arial.ttf",
        ]
        font = next((path for path in candidates if path and Path(path).is_file()), None)
        if not font:
            error("PDF_FONT_MISSING", "Не установлен шрифт Unicode для печатных форм", 503)
        pdfmetrics.registerFont(TTFont(font_name, font))
    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=32,
        bottomMargin=32,
        title=f"{snapshot['title']} {snapshot['number']}",
    )
    style = ParagraphStyle("body", fontName=font_name, fontSize=8, leading=11, spaceAfter=8)
    title_style = ParagraphStyle("title", parent=style, fontSize=16, leading=20, spaceAfter=16)

    def p(text):
        return Paragraph(escape(str(text)), style)

    content = [
        Paragraph(escape(f"{snapshot['title']} № {snapshot['number']}"), title_style),
        p(f"Дата: {snapshot['date']} · Валюта: {snapshot['currency']}"),
        p(f"Продавец: {snapshot['seller']['name']}"),
        p(format_details(snapshot["seller"].get("details"))),
        p(f"Клиент: {snapshot['client']['name']}"),
        p(format_details(snapshot["client"].get("details"))),
        Spacer(1, 8),
    ]
    pdf_headers = ["№", "Наименование", "Кол-во / ед.", "Без налога", "Налог", "Итого"]
    table_rows = [[p(value) for value in pdf_headers]]
    for index, row in enumerate(snapshot["lines"], 1):
        title = row["description"] + (f" · CAS {row['cas']}" if row.get("cas") else "")
        table_rows.append(
            [
                p(index),
                p(title),
                p(f"{row['quantity']} {row['unit']}"),
                p(row["net"]),
                p(row["tax"]),
                p(row["total"]),
            ]
        )
    table = Table(table_rows, colWidths=[22, 220, 67, 75, 60, 91], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F0EB")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C6D3CC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    # Available A4 width is 535pt; explicit widths match exactly.
    content.extend(
        [
            table,
            Spacer(1, 14),
            p(
                f"Без налога: {snapshot['totals']['net']} · Налог: {snapshot['totals']['tax']} · Итого: {snapshot['totals']['total']} {snapshot['currency']}"
            ),
            p(f"Условия: {snapshot['terms']}"),
            p(f"Действует / оплатить до: {snapshot['valid_until']}"),
        ]
    )
    pdf.build(content)
    return {"pdf": put_file(buffer.getvalue(), "pdf"), "xlsx": put_file(xlsx, "xlsx")}

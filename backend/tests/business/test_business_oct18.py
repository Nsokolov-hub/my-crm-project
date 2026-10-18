"""Business remarks of 18 October 2026: supplier orders carry logistics, ordered positions move to
«Прибыло в РФ» and «Доставлено клиенту» or are cancelled with a reason, proposal files are named for clients."""

import io
from datetime import date, timedelta
from decimal import Decimal
from urllib.parse import quote

from openpyxl import load_workbook

from app.core.models import PermissionGrant, User
from tests.support.business import cmd, handoff, issue, release, structured, wave_row
from tests.support.crm import login

LOGISTICS = [{"name": "Международная логистика", "amount": "300", "currency": "RUB", "scope": "WAVE",
              "stage": "INTERNATIONAL_LOGISTICS", "method": "BY_QUANTITY"}]


def ordered_position(crm, mode="IMPORT", logistics=True):
    """An approved position handed to procurement; an import wave budgets its international logistics."""
    req, data, seller, supplier = structured(crm, mode=mode)
    if mode == "IMPORT" and logistics:
        wave = wave_row(crm, req["wave_id"])
        cmd(crm, f"/waves/{wave['id']}/budget", {"version": wave["version"], "profile_id": data["profile_id"],
                                                 "expenses": LOGISTICS, "reason": "Бюджет волны"}, method="put")
        data["expenses"] = LOGISTICS
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    _, execution = release(crm, req, calculation, "1")
    handoff(crm, req, execution)
    return req, seller, execution


def create_order(crm, seller, execution, **extra):
    return cmd(crm, "/supplier-orders", {
        "execution_ids": [execution["id"]], "seller_id": seller["id"],
        "expected_date": (date.today() + timedelta(days=21)).isoformat(), "contract": "C-123",
        "payment_terms": "Prepayment", "delivery_terms": "CPT Moscow", **extra,
    })


def logistics_hint(crm, execution):
    response = crm["client"].get("/api/v1/supplier-orders/logistics", params={"execution_ids": [execution["id"]]})
    assert response.status_code == 200, response.text
    return response.json()


def ordered_lines(crm, **params):
    response = crm["client"].get("/api/v1/supplier-orders/positions", params={"ordered": "true", **params})
    assert response.status_code == 200, response.text
    return response.json()["items"]


def orders(crm):
    return crm["client"].get("/api/v1/supplier-orders").json()["items"]


def refs(lines):
    return [{"id": line["id"], "version": line["version"]} for line in lines]


def test_order_takes_wave_logistics_into_total_and_purchase_order(crm):
    req, seller, execution = ordered_position(crm)
    hint = logistics_hint(crm, execution)
    assert Decimal(hint["amount"]) == 300
    assert hint["currency"] == "RUB" and "бюджет волны" in hint["basis"]
    assert hint["wave_ids"] == [req["wave_id"]]
    order = create_order(crm, seller, execution)
    assert Decimal(order["logistics_amount"]) == 300
    assert Decimal(order["total"]) == 1000 + 300
    assert orders(crm)[0]["logistics_amount"] == order["logistics_amount"]

    book = load_workbook(io.BytesIO(crm["client"].get(f"/api/v1/supplier-orders/{order['id']}/po.xlsx").content))
    values = [cell.value for row in book.active for cell in row]
    assert "International logistics" in values and Decimal(str(max(v for v in values if isinstance(v, (int, float, Decimal))))) == 1300
    pdf = crm["client"].get(f"/api/v1/supplier-orders/{order['id']}/po.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-type"] == "application/pdf"
    assert f"PO_{order['number']}.pdf" in pdf.headers["content-disposition"]

    # The wave logistics is in a live order: another order of the wave is not charged with it again.
    again = logistics_hint(crm, execution)
    assert Decimal(again["amount"]) == 0 and order["number"] in again["notes"][0]


def test_logistics_can_be_corrected_by_hand_and_comes_from_the_supplier_card_without_wave(crm):
    req, seller, execution = ordered_position(crm, mode="RUSSIA")
    assert Decimal(logistics_hint(crm, execution)["amount"]) == 0
    order = create_order(crm, seller, execution, logistics_amount="125.50")
    assert Decimal(order["total"]) == Decimal("1125.50")
    assert order["logistics_basis"] == "Указано вручную"

    supplier_id = order["supplier_id"]
    party = crm["client"].get(f"/api/v1/counterparties/{supplier_id}").json()
    response = crm["client"].patch(f"/api/v1/counterparties/{supplier_id}", json={
        "version": party["version"], "name": party["name"], "kind": party["kind"],
        "details": {**party["details"], "Логистика по умолчанию": "40", "Валюта логистики": "RUB"},
    })
    assert response.status_code == 200, response.text
    cancel = cmd(crm, "/supplier-order-lines/cancel", {"lines": refs(ordered_lines(crm)),
                                                       "reason": "Поставщик не отгрузит"})
    assert cancel["count"] == 1
    hint = logistics_hint(crm, execution)
    assert Decimal(hint["amount"]) == 40 and hint["basis"] == "карточка поставщика"


def test_logistician_moves_ordered_positions_but_cannot_order_or_cancel(crm):
    req, seller, execution = ordered_position(crm)
    order = create_order(crm, seller, execution)
    with crm["sessions"].begin() as db:
        logistician = User(email="logist@example.com", name="Логист", password_hash=crm["admin"].password_hash)
        db.add(logistician)
        db.flush()
        db.add_all(PermissionGrant(user_id=logistician.id, code=code, scope="all")
                   for code in ("requests.read", "supplier_orders.delivery"))
    login(crm, "logist@example.com")
    assert crm["client"].get("/api/v1/supplier-orders/positions").status_code == 403
    assert crm["client"].get("/api/v1/supplier-orders").status_code == 403
    assert crm["client"].get(f"/api/v1/supplier-orders/{order['id']}/po.pdf").status_code == 403
    lines = ordered_lines(crm)
    assert "unit_price" not in lines[0]
    assert lines[0]["request_number"] == req["number"] and lines[0]["client_name"] == "Рабочий клиент"
    arrived = cmd(crm, "/supplier-order-lines/status", {"lines": refs(lines), "status": "arrived_ru"})
    assert arrived["count"] == 1 and arrived["items"][0]["status"] == "arrived_ru"
    assert [line["id"] for line in ordered_lines(crm, status="arrived_ru")] == [lines[0]["id"]]
    denied = cmd(crm, "/supplier-order-lines/cancel", {"lines": refs(ordered_lines(crm)), "reason": "Ошибка"},
                 status=403)
    assert denied["code"] == "FORBIDDEN"
    cmd(crm, "/supplier-order-lines/status", {"lines": refs(ordered_lines(crm)), "status": "delivered"})
    stale = cmd(crm, "/supplier-order-lines/status", {"lines": refs(lines), "status": "delivery"}, status=409)
    assert stale["code"] == "VERSION_CONFLICT"

    login(crm)
    assert orders(crm)[0]["status"] == "delivered"
    assert ordered_lines(crm)[0]["unit_price"]


def test_cancelled_position_keeps_its_reason_and_can_be_ordered_again(crm):
    req, seller, execution = ordered_position(crm, logistics=False)
    first = create_order(crm, seller, execution)
    lines = ordered_lines(crm)
    blank = cmd(crm, "/supplier-order-lines/cancel", {"lines": refs(lines), "reason": "   "}, status=422)
    assert blank["code"] == "CANCEL_REASON_REQUIRED"
    cmd(crm, "/supplier-order-lines/cancel", {"lines": refs(lines), "reason": "Поставщик  снял позицию"})
    cancelled = ordered_lines(crm)[0]
    assert cancelled["status"] == "cancelled" and cancelled["cancel_reason"] == "Поставщик снял позицию"
    assert cancelled["cancelled_at"]
    assert orders(crm)[0]["status"] == "cancelled"
    moved = cmd(crm, "/supplier-order-lines/status", {"lines": refs([cancelled]), "status": "delivery"},
                status=409)
    assert moved["code"] == "ORDER_LINE_CANCELLED"

    positions = crm["client"].get("/api/v1/supplier-orders/positions").json()["items"]
    assert [row["id"] for row in positions] == [execution["id"]]
    second = create_order(crm, seller, execution)
    assert second["number"] != first["number"]
    states = sorted(line["status"] for line in ordered_lines(crm))
    assert states == ["cancelled", "waiting"]
    live = [line for line in ordered_lines(crm) if line["status"] == "waiting"]
    cmd(crm, "/supplier-order-lines/status", {"lines": refs(live), "status": "delivered"})
    refused = cmd(crm, "/supplier-order-lines/cancel",
                  {"lines": refs([line for line in ordered_lines(crm) if line["status"] == "delivered"]),
                   "reason": "Поздно"}, status=409)
    assert refused["code"] == "ORDER_LINE_DELIVERED"
    export = load_workbook(io.BytesIO(
        crm["client"].get("/api/v1/supplier-orders/positions/export.xlsx").content)).active
    exported = [cell.value for row in export for cell in row]
    assert "Отменено" in exported and "Поставщик снял позицию" in exported and "Доставлено клиенту" in exported


def test_proposal_file_is_named_by_buyer_number_and_date(crm):
    req, data, _, _ = structured(crm)
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    proposal = issue(crm, req, calculation)
    issued = date.fromisoformat(proposal["snapshot"]["date"]).strftime("%d.%m.%Y")
    expected = f"Рабочий клиент, КП №{proposal['display_number']}, {issued}"
    listed = crm["client"].get(f"/api/v1/requests/{req['id']}/proposals").json()["items"][0]
    assert listed["file_name"] == expected
    response = crm["client"].get(f"/api/v1/documents/{proposal['id']}/file", params={"format": "pdf"})
    assert response.status_code == 200
    assert f"filename*=UTF-8''{quote(expected + '.pdf')}" in response.headers["content-disposition"]

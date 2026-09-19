import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, Depends, File, Form, Header, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import Field
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.db import get_db
from app.core.errors import DomainError
from app.core.models import OutboxEvent, User
from app.core.routes import Input
from app.core.security import check_client, client_predicate, current_user, require_permission
from app.core.service import advisory, audit, idem, lock, serialize
from app.core.service import page as paginate
from app.crm.models import Contact, Counterparty, ImportBatch, ImportRow

router = APIRouter(tags=["Импорт клиентской базы"])
COLUMNS = {
    "external_id": "Внешний ID",
    "name": "Организация",
    "country": "Страна",
    "tax_id": "ИНН",
    "contact": "Контактное лицо",
    "phone": "Телефон",
    "email": "Электронная почта",
    "source": "Источник",
    "comment": "Комментарий",
    "owner_id": "Ответственный",
}
FIELD_LIMITS = {
    "external_id": 250,
    "name": 250,
    "country": 100,
    "tax_id": 100,
    "contact": 250,
    "phone": 100,
    "email": 254,
    "source": 250,
    "comment": 2000,
    "owner_id": 36,
}


def xlsx(rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Клиентская база"
    for row in rows:
        sheet.append([str(v) if v is not None else "" for v in row])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.number_format = "@"
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 25
    sheet.freeze_panes = "A2"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def string_value(cell: Any) -> str:
    if cell.value is None:
        return ""
    if isinstance(cell.value, int) and re.fullmatch(r"0+", cell.number_format or ""):
        return str(cell.value).zfill(len(cell.number_format))
    return str(cell.value).strip()


def parse_rows(data: bytes, mapping: dict[str, str]) -> list[dict[str, Any]]:
    try:
        archive = ZipFile(io.BytesIO(data))
        infos = archive.infolist()
        if (
            len(infos) > 10000
            or sum(i.file_size for i in infos) > 100 * 1024 * 1024
            or any("vbaProject" in i.filename for i in infos)
        ):
            raise DomainError("IMPORT_UNSAFE", "Архив слишком большой или содержит макросы")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError):
        raise DomainError(
            "IMPORT_INVALID_XLSX", "Не удалось прочитать XLSX. Используйте шаблон импорта."
        ) from None
    try:
        sheet = workbook.active
        if sheet.max_row and sheet.max_row > 10001:
            raise DomainError("IMPORT_ROW_LIMIT", "В одной партии допускается до 10 000 строк")
        raw = iter(sheet.iter_rows())
        headers = [string_value(c) for c in next(raw, [])]
        if len(headers) > 100:
            raise DomainError("IMPORT_COLUMN_LIMIT", "В файле слишком много колонок")
        index = {
            key: headers.index(mapping.get(key, label))
            for key, label in COLUMNS.items()
            if mapping.get(key, label) in headers
        }
        if "name" not in index and "contact" not in index:
            raise DomainError(
                "IMPORT_MAPPING_REQUIRED", "Сопоставьте колонку организации или контакта", field="mapping"
            )
        output = []
        for number, cells in enumerate(raw, 2):
            if number > 10001:
                raise DomainError("IMPORT_ROW_LIMIT", "В одной партии допускается до 10 000 строк")
            if not any(c.value is not None for c in cells):
                continue
            row = {key: string_value(cells[i]) if i < len(cells) else "" for key, i in index.items()}
            errors = []
            if any(c.data_type == "f" for c in cells):
                errors.append({"field": "row", "message": "Формулы не допускаются: вставьте значения"})
            if not row.get("name") and not row.get("contact"):
                errors.append({"field": "name", "message": "Укажите организацию или контакт"})
            row["email"] = row.get("email", "").strip().lower()
            row["phone"] = re.sub(r"[^\d+]", "", row.get("phone", ""))
            if not row["email"] and not row["phone"]:
                errors.append({"field": "phone", "message": "Нужен телефон или электронная почта"})
            if row["email"] and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", row["email"]):
                errors.append({"field": "email", "message": "Некорректная электронная почта"})
            for field, value in row.items():
                if len(value) > FIELD_LIMITS[field]:
                    errors.append(
                        {"field": field, "message": f"Допускается не более {FIELD_LIMITS[field]} символов"}
                    )
            output.append({"row_number": number, "data": row, "errors": errors})
        return output
    finally:
        workbook.close()


@router.get("/imports/template.xlsx")
def template(user: User = Depends(current_user), db: Session = Depends(get_db)) -> Response:
    require_permission(db, user, "imports.write")
    return Response(
        xlsx([list(COLUMNS.values())]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="clients-template-v1.xlsx"'},
    )


@router.post("/imports/preview")
async def preview(
    file: UploadFile = File(),
    mapping: str = Form("{}"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_permission(db, user, "imports.write")
    require_permission(db, user, "clients.write")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise DomainError("IMPORT_TYPE_INVALID", "Загрузите файл XLSX без макросов")
    data = await file.read(settings.max_file_size + 1)
    if len(data) > settings.max_file_size:
        raise DomainError("FILE_TOO_LARGE", "Размер файла превышает разрешённый предел", 413)
    try:
        columns = json.loads(mapping)
        if not isinstance(columns, dict) or any(
            k not in COLUMNS or not isinstance(v, str) for k, v in columns.items()
        ):
            raise ValueError
    except (ValueError, TypeError):
        raise DomainError(
            "MAPPING_INVALID", "Сопоставление колонок должно быть JSON-объектом", field="mapping"
        ) from None
    parsed = parse_rows(data, columns)
    key = f"imports/{uuid4().hex}.xlsx"
    path = settings.storage_dir / key
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    batch = ImportBatch(
        author_id=user.id,
        source_name=Path(file.filename or "import.xlsx").name[:250],
        file_key=key,
        file_hash=hashlib.sha256(data).hexdigest(),
        mapping=columns,
    )
    db.add(batch)
    db.flush()
    seen_external: set[str] = set()
    for raw in parsed:
        values = raw["data"]
        row = ImportRow(batch_id=batch.id, **raw)
        external = values.get("external_id")
        match = (
            db.scalar(select(Counterparty).where(Counterparty.external_id == external)) if external else None
        )
        if external and external in seen_external:
            row.errors = [
                *row.errors,
                {"field": "external_id", "message": "Внешний ID повторяется внутри файла"},
            ]
        if external:
            seen_external.add(external)
        if match:
            try:
                check_client(db, user, match.id, "clients.write")
                row.match_id = match.id
                row.action = "update"
            except DomainError:
                row.errors = [
                    *row.errors,
                    {"field": "external_id", "message": "Идентификатор недоступен для обновления"},
                ]
        else:
            conditions = []
            for attr in ("email", "phone", "tax_id"):
                if values.get(attr):
                    conditions.append(getattr(Counterparty, attr) == values[attr])
            if conditions:
                candidates = list(
                    db.scalars(
                        select(Counterparty.id)
                        .where(client_predicate(db, user, "clients.write"), or_(*conditions))
                        .limit(20)
                    )
                )
                row.candidate_ids = candidates
                if candidates:
                    row.action = "conflict"
        owner_id = values.get("owner_id")
        if owner_id:
            target = db.get(User, owner_id)
            if not target or not target.active:
                row.errors = [*row.errors, {"field": "owner_id", "message": "Ответственный не найден"}]
            elif owner_id != user.id:
                require_permission(db, user, "requests.assign")
        if row.errors:
            row.action = "error"
        db.add(row)
    db.flush()
    batch.summary = summary(db, batch.id)
    audit(db, user, "import", batch.id, "preview", after=batch.summary)
    return {
        **serialize(batch),
        "rows": paginate(
            db, select(ImportRow).where(ImportRow.batch_id == batch.id).order_by(ImportRow.row_number), 1, 100
        )["items"],
    }


def summary(db: Session, batch_id: str) -> dict[str, int]:
    rows = db.scalars(select(ImportRow).where(ImportRow.batch_id == batch_id)).all()
    return {
        **{
            k: sum(r.action == k for r in rows)
            for k in ("create", "update", "error", "conflict", "skip", "created", "updated")
        },
        "total": len(rows),
    }


def owned_batch(db: Session, user: User, entity_id: str) -> ImportBatch:
    require_permission(db, user, "imports.write")
    row = db.get(ImportBatch, entity_id)
    if row is None or row.author_id != user.id:
        raise DomainError("NOT_FOUND", "Партия импорта недоступна", 404)
    return row


@router.get("/imports")
def batches(
    page: int = 1, page_size: int = 25, user: User = Depends(current_user), db: Session = Depends(get_db)
) -> dict[str, Any]:
    require_permission(db, user, "imports.write")
    return paginate(
        db,
        select(ImportBatch).where(ImportBatch.author_id == user.id).order_by(ImportBatch.created_at.desc()),
        page,
        page_size,
    )


@router.get("/imports/{entity_id}")
def batch_detail(
    entity_id: str,
    page: int = 1,
    page_size: int = 100,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    row = owned_batch(db, user, entity_id)
    return {
        **serialize(row),
        "rows": paginate(
            db,
            select(ImportRow).where(ImportRow.batch_id == entity_id).order_by(ImportRow.row_number),
            page,
            page_size,
        )["items"],
    }


class Resolution(Input):
    row_id: str
    action: Literal["create", "update", "skip"]
    match_id: str | None = None


class Confirmation(Input):
    decisions: list[Resolution] = Field(default_factory=list)


@router.post("/imports/{entity_id}/confirm")
def confirm(
    entity_id: str,
    body: Confirmation,
    idempotency_key: str | None = Header(default=None),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    owned_batch(db, user, entity_id)
    require_permission(db, user, "clients.write")

    def operation() -> dict[str, Any]:
        batch = lock(db, ImportBatch, entity_id)
        if batch.status != "preview":
            return serialize(batch)
        for decision in body.decisions:
            row = db.get(ImportRow, decision.row_id)
            if row is None or row.batch_id != batch.id or row.errors:
                raise DomainError(
                    "IMPORT_RESOLUTION_INVALID", "Решение не относится к корректной строке партии"
                )
            if decision.action == "update":
                if not decision.match_id:
                    raise DomainError("MATCH_REQUIRED", "Для обновления выберите существующую карточку")
                check_client(db, user, decision.match_id, "clients.write")
            row.action = decision.action
            row.match_id = decision.match_id if decision.action == "update" else None
        db.flush()
        if db.scalar(
            select(ImportRow.id).where(ImportRow.batch_id == batch.id, ImportRow.action == "conflict")
        ):
            raise DomainError(
                "IMPORT_CONFLICT_UNRESOLVED", "Сопоставьте конфликтующие строки или пропустите их", 409
            )
        batch.status = "queued"
        db.add(
            OutboxEvent(
                event_key=f"import:{batch.id}",
                kind="import.confirm",
                payload={"batch_id": batch.id, "user_id": user.id},
            )
        )
        audit(db, user, "import", batch.id, "queued", after=summary(db, batch.id))
        return serialize(batch)

    return idem(db, user, idempotency_key, f"imports.confirm:{entity_id}", body.model_dump(), operation)


def process_import(db: Session, payload: dict[str, Any]) -> dict[str, Any]:
    user = db.get(User, payload["user_id"])
    if user is None or not user.active:
        raise DomainError("IMPORT_ACCESS_REVOKED", "Доступ инициатора импорта отозван")
    require_permission(db, user, "imports.write")
    require_permission(db, user, "clients.write")
    batch = lock(db, ImportBatch, payload["batch_id"])
    if batch.status == "completed":
        return serialize(batch)
    advisory(db, "import.counterparties")
    for row in db.scalars(
        select(ImportRow).where(ImportRow.batch_id == batch.id).order_by(ImportRow.row_number)
    ):
        if row.action not in ("create", "update"):
            continue
        data = row.data
        external = data.get("external_id") or None
        matched = (
            db.scalar(select(Counterparty).where(Counterparty.external_id == external)) if external else None
        )
        target_id = row.match_id or (matched.id if matched else None)
        client = check_client(db, user, target_id, "clients.write") if target_id else None
        before = serialize(client) if client else None
        if not client:
            client = Counterparty(
                name=data.get("name") or data.get("contact"),
                owner_id=data.get("owner_id") or user.id,
                external_id=external,
                details={},
            )
            db.add(client)
        elif external and client.external_id not in (None, external):
            raise DomainError("EXTERNAL_ID_CONFLICT", "Карточка уже связана с другим внешним ID", 409)
        for field in ("name", "country", "tax_id", "email", "phone"):
            if data.get(field):
                setattr(client, field, data[field])
        # The original acquisition source survives later imports, like the call history.
        if data.get("source") and not client.source:
            client.source = data["source"]
        if external:
            client.external_id = external
        if data.get("comment"):
            client.details = {**(client.details or {}), "comment": data["comment"]}
        if before:
            client.version += 1
        db.flush()
        if data.get("contact"):
            existing = db.scalar(
                select(Contact).where(Contact.client_id == client.id, Contact.name == data["contact"])
            )
            if not existing:
                db.add(
                    Contact(
                        client_id=client.id,
                        name=data["contact"],
                        email=data.get("email") or None,
                        phone=data.get("phone") or None,
                    )
                )
        row.match_id = client.id
        row.action = "updated" if before else "created"
        audit(
            db,
            user,
            "counterparty",
            client.id,
            "imported",
            before,
            serialize(client),
            f"Партия {batch.id}, строка {row.row_number}",
        )
    db.flush()
    batch.status = "completed"
    batch.summary = summary(db, batch.id)
    audit(db, user, "import", batch.id, "completed", after=batch.summary)
    return serialize(batch)


@router.get("/imports/{entity_id}/errors.xlsx")
def errors_file(
    entity_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
) -> Response:
    owned_batch(db, user, entity_id)
    require_permission(db, user, "exports.download")
    rows = [["Строка", "Поле", "Ошибка"]]
    for row in db.scalars(select(ImportRow).where(ImportRow.batch_id == entity_id)):
        rows += [[row.row_number, e["field"], e["message"]] for e in row.errors]
    audit(db, user, "import", entity_id, "errors_exported")
    return Response(
        xlsx(rows),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="import-errors.xlsx"'},
    )

"""Regression coverage for mistaken quotes, single bonuses and customer documents."""

import io
from datetime import date, timedelta
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from test_business_imports import confirm, preview, request
from test_business_oct5 import cmd, configured_profile, structured
from test_crm import crm, login, post  # noqa: F401
from test_itemized_wave_calculation import _selection

from app.business.models import SupplierMail
from app.business.worker import send_mail
from app.commerce.files import PROPOSAL_NOTICE, PROPOSAL_SIGNATURE, read_file
from app.commerce.itemized import calculate_itemized
from app.commerce.models import Calculation, CommercialDocument
from app.core.config import settings
from app.core.models import OutboxEvent
from app.crm.models import QuoteItem, QuoteSheet


@pytest.mark.parametrize("kind,value", [("PERCENTAGE", "15"), ("MULTIPLIER", "1.15"), ("FIXED", "225")])
def test_line_bonus_is_not_added_again_by_legacy_global_adjustment(kind, value):
    row = _selection("line", "reference_standards", "1", "1000", "1.15")
    row.update(currency_code="RUB", calculation_type="RUSSIA")
    baseline = calculate_itemized(configured_profile(), [row], [], [])
    result = calculate_itemized(configured_profile(), [row], [], [], {
        "enabled": True, "type": kind, "value": value, "service_fee_percent": "0",
    })
    assert result["totals"] == baseline["totals"]
    assert Decimal(result["totals"]["internal_bonus"]) == 225
    assert Decimal(result["lines"][0]["detail"]["bonus_withdrawal_fee"]) == 36


def test_legacy_global_bonus_only_and_separate_service_fee_remain_supported():
    row = _selection("line", "reference_standards", "1", "1000", "1")
    row.update(currency_code="RUB", calculation_type="RUSSIA")
    result = calculate_itemized(configured_profile(), [row], [], [], {
        "enabled": True, "type": "PERCENTAGE", "value": "15", "service_fee_percent": "2",
    })
    assert Decimal(result["totals"]["internal_bonus"]) == 225
    assert Decimal(result["lines"][0]["detail"]["additional_service_fee"]) == 20


def release(env, req, calculation, days=30):
    return cmd(env, f"/requests/{req['id']}/proposals", {
        "calculation_id": calculation["id"],
        "valid_until": (date.today() + timedelta(days=days)).isoformat(),
        "terms": "Оплата согласно договору.",
    })


def test_documents_have_five_day_lifetime_short_number_and_fixed_proposal_signature(crm):  # noqa: F811
    req, data, _, supplier = structured(crm)
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    assert calculation["snapshot"]["lines"][0]["delivery_days"] == 7
    assert calculation["snapshot"]["lines"][0]["supplier_name"] == supplier["name"]
    proposal = release(crm, req, calculation)
    until = (date.today() + timedelta(days=5)).isoformat()
    assert proposal["valid_until"] == proposal["snapshot"]["valid_until"] == until
    assert proposal["display_number"] == f"{req['number']}-1"
    assert proposal["snapshot"]["signature"] == PROPOSAL_SIGNATURE
    assert proposal["snapshot"]["validity_notice"] == PROPOSAL_NOTICE
    assert release(crm, req, calculation, days=2)["valid_until"] == (date.today() + timedelta(days=2)).isoformat()
    with crm["sessions"]() as db:
        document = db.get(CommercialDocument, proposal["id"])
        book = load_workbook(io.BytesIO(read_file(document.files["xlsx"])))
        values = [cell.value for row in book.active for cell in row]
        assert "Заявка" not in values
        assert f"Коммерческое предложение №{req['number']}-1 от {date.today():%d.%m.%Y}" in values
        assert "Генеральный директор" in values and "Гильмутдинов Т.Ф" in values
        signature_row = next(row[0].row for row in book.active if row[0].value == "Генеральный директор")
        assert book.active.cell(signature_row, 5).value == "Гильмутдинов Т.Ф"
        assert book.active.cell(signature_row, 5).alignment.horizontal == "right"
        assert PROPOSAL_NOTICE in values
        assert "https://ogk-chem.ru · info@ogk-chem.ru" in values
        header = next(row[0].row for row in book.active if row[0].value == "№")
        assert book.active.cell(header, 1).fill.fgColor.rgb == "00173F35"
        assert book.active.print_title_rows == f"${header}:${header}"
        book.close()
    accepted = cmd(crm, f"/proposals/{proposal['id']}/accept", {
        "version": proposal["version"], "reason": "Клиент согласовал",
        "lines": [{"line_id": calculation["snapshot"]["lines"][0]["line_id"], "quantity": "1"}],
    })
    invoice_data = {"proposal_id": proposal["id"], "due_date": (date.today() + timedelta(days=30)).isoformat(),
        "terms": "По договору", "lines": [{"execution_id": accepted["executions"][0]["id"], "quantity": "1"}]}
    expired = cmd(crm, f"/requests/{req['id']}/invoices", {
        **invoice_data, "due_date": (date.today() - timedelta(days=1)).isoformat(),
    }, status=422)
    assert expired["code"] == "DOCUMENT_EXPIRED"
    invoice = cmd(crm, f"/requests/{req['id']}/invoices", invoice_data)
    assert invoice["valid_until"] == invoice["snapshot"]["valid_until"] == until


def test_delete_quote_item_preserves_saved_documents_but_prevents_new_use(crm):  # noqa: F811
    req, data, _, _ = structured(crm)
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    proposal = release(crm, req, calculation)
    before = crm["client"].get(f"/api/v1/documents/{proposal['id']}/file?format=pdf").content
    row = crm["client"].get(f"/api/v1/requests/{req['id']}/quote-items").json()["items"][0]
    key = str(uuid4())
    deleted = cmd(crm, f"/quote-items/{row['id']}/delete", {"version": row["version"]}, key=key)
    assert deleted == {"deleted_ids": [row["id"]]}
    assert cmd(crm, f"/quote-items/{row['id']}/delete", {"version": row["version"]}, key=key) == deleted
    assert crm["client"].get(f"/api/v1/requests/{req['id']}/quote-items").json()["items"] == []
    assert crm["client"].get(f"/api/v1/requests/{req['id']}/quote-sheets").json()["items"] == []
    data["request_version"] = crm["client"].get(f"/api/v1/requests/{req['id']}").json()["version"]
    assert cmd(crm, f"/requests/{req['id']}/calculations/preview", data, status=404)["code"] == "QUOTE_ITEM_NOT_FOUND"
    assert cmd(crm, f"/requests/{req['id']}/proposals", {
        "calculation_id": calculation["id"], "valid_until": date.today().isoformat(), "terms": "По договору",
    }, status=422)["code"] == "QUOTE_DELETED"
    assert crm["client"].get(f"/api/v1/documents/{proposal['id']}/file?format=pdf").content == before
    with crm["sessions"]() as db:
        assert db.get(QuoteItem, row["id"]).archived
        assert db.get(QuoteSheet, row["quote_sheet_id"]).archived
        assert db.get(Calculation, calculation["id"]).snapshot == calculation["snapshot"]
        assert db.get(CommercialDocument, proposal["id"]).snapshot == proposal["snapshot"]


def test_delete_entire_quote_checks_permission_and_version_and_excludes_price_reuse(crm):  # noqa: F811
    req, data, _, supplier = structured(crm, owner=crm["manager"].id)
    with crm["sessions"]() as db:
        q = db.get(QuoteItem, data["selections"][0]["quote_item_id"])
        item = {"nomenclature_id": q.nomenclature_id, "packing_id": q.packing_id,
                "quantity": "1", "unit_price": "500", "currency_id": q.currency_id, "delivery_days": 9}
        source_item_id = q.source_request_item_id
    sheet = post(crm, f"/requests/{req['id']}/quote-sheets", {"supplier_id": supplier["id"], "items": [item, {**item, "source_request_item_id": source_item_id}]})
    login(crm, "manager@example.com")
    assert cmd(crm, f"/quote-sheets/{sheet['id']}/delete", {"version": sheet["version"]}, status=404)["code"] == "NOT_FOUND"
    login(crm)
    cmd(crm, f"/quote-sheets/{sheet['id']}/delete", {"version": sheet["version"] + 1}, status=409)
    result = cmd(crm, f"/quote-sheets/{sheet['id']}/delete", {"version": sheet["version"]})
    assert set(result["deleted_ids"]) == {q["id"] for q in sheet["items"]}
    active = crm["client"].get(f"/api/v1/requests/{req['id']}/quote-items?all=true").json()["items"]
    assert [q["id"] for q in active] == [data["selections"][0]["quote_item_id"]]
    history = crm["client"].get("/api/v1/quote-items/history").json()["items"]
    assert [q["id"] for q in history] == [data["selections"][0]["quote_item_id"]]
    latest = crm["client"].get("/api/v1/quote-items/latest", params={"supplier_id": supplier["id"],
        "nomenclature_id": item["nomenclature_id"], "packing_id": item["packing_id"]}).json()["item"]
    assert latest["id"] == data["selections"][0]["quote_item_id"]


def test_mail_matches_imported_rfq_columns_and_sends_cc_in_smtp_envelope(crm, monkeypatch):  # noqa: F811
    login(crm)
    req = request(crm)
    columns = ["Name", "Produce", "Art / CAS", "Packing", "Qty", "Comment <extra>"]
    rows = [["Lornoxicam <sample>", "Your", "70374-39-9", "250 mg", "1", "Keep & test"],
            ["Second", "Maker", "0042", "100 mg", "0", ""]]
    batch = preview(crm, req["id"], columns, rows)
    confirm(crm, req["id"], batch["id"])
    items = crm["client"].get(f"/api/v1/requests/{req['id']}/items").json()["items"]
    supplier = post(crm, "/counterparties", {"name": "Supplier", "kind": "supplier", "email": "sales@example.com"})
    payload = {"item_ids": [item["id"] for item in reversed(items)], "supplier_ids": [supplier["id"]],
               "cc": ["info@ogk-chem.ru", "sales@example.com", "info@ogk-chem.ru"]}
    draft = cmd(crm, f"/requests/{req['id']}/supplier-mail/preview", payload)["items"][0]
    assert draft["cc"] == ["info@ogk-chem.ru"]
    assert "\t".join(columns) in draft["body"]
    assert "Your\t70374-39-9\t250 mg\t1" in draft["body"]
    assert "Maker\t0042\t100 mg\t0" in draft["body"]
    assert "Lornoxicam &lt;sample&gt;" in draft["table_html"]
    assert "Comment &lt;extra&gt;" in draft["table_html"]
    rfq = cmd(crm, f"/requests/{req['id']}/rfqs", {"supplier_id": supplier["id"],
        "item_ids": payload["item_ids"], "response_due": date.today().isoformat()})
    book = load_workbook(io.BytesIO(crm["client"].get(f"/api/v1/rfqs/{rfq['id']}/file").content))
    assert list(book.active.values)[0] == tuple(columns)
    assert [line.split("\t")[0] for line in draft["body"].splitlines()[5:]] == [row[0].value for row in list(book.active)[1:]]
    book.close()
    cmd(crm, f"/requests/{req['id']}/supplier-mail/preview", {**payload, "cc": ["bad-address"]}, status=422)
    monkeypatch.setattr(settings, "smtp_host", "smtp.example")
    monkeypatch.setattr(settings, "smtp_from", "crm@example.com")
    messages, envelopes = [], []

    class SMTP:
        def __init__(self, *args, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def starttls(self):
            pass

        def send_message(self, message):
            messages.append(message)
            # smtplib.send_message derives its recipient envelope from To and Cc.
            envelopes.append([message["To"], message["Cc"]])

    monkeypatch.setattr("app.business.worker.smtplib.SMTP", SMTP)
    sent = cmd(crm, f"/requests/{req['id']}/supplier-mail/send", {**payload, "subject": "Request"})["items"][0]
    assert sent["cc"] == ["info@ogk-chem.ru"]
    with crm["sessions"].begin() as db:
        mail = db.get(SupplierMail, sent["id"])
        assert mail.body == draft["body"]
        event = db.scalar(select(OutboxEvent).where(OutboxEvent.kind == "supplier.mail"))
        send_mail(db, event)
        send_mail(db, event)
    assert len(messages) == 1
    assert envelopes == [["sales@example.com", "info@ogk-chem.ru"]]
    assert draft["table_html"] in messages[0].get_body(preferencelist=("html",)).get_content()

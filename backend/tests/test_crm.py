"""Real HTTP authentication, import A01 and CRM A02/A10 regression checks."""

import io
import json
from datetime import datetime, timedelta, timezone
from uuid import uuid4

import pyotp
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.db import Base, get_db
from app.core.models import AuthSession, PermissionGrant, User
from app.core.security import PERMISSIONS, password_hasher
from app.crm.imports import COLUMNS, parse_rows, process_import, xlsx
from app.crm.models import Call, Contact, Counterparty, ImportBatch, Request, Task
from app.main import app

PASSWORD = "Only-for-automated-tests-1234"


@pytest.fixture
def crm(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)

    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _record):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    sessions = sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(settings, "storage_dir", tmp_path / "files")
    monkeypatch.setattr(settings, "require_mfa", False)
    with sessions.begin() as db:
        admin = User(
            email="owner@example.com", name="Владелец теста", password_hash=password_hasher.hash(PASSWORD)
        )
        manager = User(email="manager@example.com", name="Менеджер теста", password_hash=admin.password_hash)
        db.add_all([admin, manager])
        db.flush()
        db.add_all(PermissionGrant(user_id=admin.id, code=code, scope="all") for code in PERMISSIONS)
        db.add_all(
            PermissionGrant(user_id=manager.id, code=code, scope="own")
            for code in (
                "clients.read",
                "clients.write",
                "requests.read",
                "requests.write",
                "calls.write",
                "tasks.read",
                "tasks.write",
            )
        )

    def get_session():
        with sessions() as db:
            try:
                yield db
                db.commit()
            except Exception:
                db.rollback()
                raise

    app.dependency_overrides[get_db] = get_session
    try:
        with TestClient(app) as client:
            res1 = client.post("/api/v1/auth/login", json={"email": "owner@example.com", "password": PASSWORD})
            assert res1.status_code == 200, res1.text
            client.headers["X-CSRF-Token"] = res1.json()["csrf_token"]
            res2 = client.post("/api/v1/setup/wizard", json={"company_name": "Test Company"})
            assert res2.status_code == 200, res2.text
            client.headers.pop("X-CSRF-Token", None)
            client.cookies.clear()
            client.cookies.clear()
            client.cookies.clear()
            yield {"client": client, "sessions": sessions, "admin": admin, "manager": manager}
    finally:
        app.dependency_overrides.pop(get_db, None)
        engine.dispose()


def login(env, email="owner@example.com", **extra):
    response = env["client"].post("/api/v1/auth/login", json={"email": email, "password": PASSWORD, **extra})
    assert response.status_code == 200, response.text
    env["client"].headers["X-CSRF-Token"] = response.json()["csrf_token"]
    return response.json()


def post(env, path, data, status=201, key=None):
    response = env["client"].post(
        "/api/v1" + path, json=data, headers={"Idempotency-Key": key or str(uuid4())}
    )
    assert response.status_code == status, response.text
    return response.json()


def test_a01_import_90_new_5_updates_5_errors_repeat_and_history(crm):
    login(crm)
    with crm["sessions"].begin() as db:
        for i in range(5):
            row = Counterparty(
                name=f"Старый {i}",
                owner_id=crm["admin"].id,
                external_id=f"existing-{i}",
                source="Исходная выставка",
                phone=f"+7999111111{i}",
            )
            db.add(row)
            db.flush()
            db.add(
                Call(
                    client_id=row.id,
                    author_id=crm["admin"].id,
                    result="no_answer",
                    occurred_at=datetime.now(timezone.utc),
                )
            )
    rows = [list(COLUMNS.values())]
    for i in range(100):
        values = {
            "external_id": f"new-{i}" if i < 90 else f"existing-{i - 90}",
            "name": f"Клиент {i}",
            "contact": f"Контакт {i}",
            "email": f" Person{i}@Example.COM " if i < 95 else "не-почта",
            "phone": "+7 (900) 123-45-67",
            "source": "Поздний импорт",
            "tax_id": f"0000{i}",
        }
        rows.append([values.get(key, "") for key in COLUMNS])
    upload = crm["client"].post(
        "/api/v1/imports/preview", files={"file": ("clients.xlsx", xlsx(rows))}, data={"mapping": "{}"}
    )
    assert upload.status_code == 200, upload.text
    batch = upload.json()
    assert batch["summary"] == {
        "create": 90,
        "update": 5,
        "error": 5,
        "conflict": 0,
        "skip": 0,
        "created": 0,
        "updated": 0,
        "total": 100,
    }
    assert batch["rows"][0]["data"]["email"] == "person0@example.com"
    assert batch["rows"][0]["data"]["phone"] == "+79001234567"
    assert batch["rows"][0]["data"]["tax_id"] == "00000"
    key = str(uuid4())
    post(crm, f"/imports/{batch['id']}/confirm", {"decisions": []}, 200, key)
    post(crm, f"/imports/{batch['id']}/confirm", {"decisions": []}, 200, key)
    payload = {"batch_id": batch["id"], "user_id": crm["admin"].id}
    with crm["sessions"].begin() as db:
        first = process_import(db, payload)
    with crm["sessions"].begin() as db:
        second = process_import(db, payload)
        assert first["summary"] == second["summary"]
        assert second["summary"]["created"] == 90
        assert second["summary"]["updated"] == 5
        assert second["summary"]["error"] == 5
        assert db.scalar(select(func.count()).select_from(Counterparty)) == 95
        assert db.scalar(select(func.count()).select_from(Contact)) == 95
        assert db.scalar(select(func.count()).select_from(Call)) == 5
        old = db.scalars(select(Counterparty).where(Counterparty.external_id.like("existing-%"))).all()
        assert all(c.source == "Исходная выставка" for c in old)
        assert all(c.name.startswith("Клиент") for c in old)
    errors = crm["client"].get(f"/api/v1/imports/{batch['id']}/errors.xlsx")
    book = load_workbook(io.BytesIO(errors.content))
    assert book.active.max_row == 6
    book.close()


def test_import_conflict_resolution_pages_and_ownership(crm):
    login(crm)
    client = post(crm, "/counterparties", {"name": "Существующий", "email": "shared@example.com"})
    data = xlsx(
        [["Организация", "Электронная почта"], *[[f"Строка {i}", "shared@example.com"] for i in range(101)]]
    )
    response = crm["client"].post("/api/v1/imports/preview", files={"file": ("clients.xlsx", data)})
    batch = response.json()
    assert len(batch["rows"]) == 100
    page2 = crm["client"].get(f"/api/v1/imports/{batch['id']}?page=2&page_size=100").json()
    assert len(page2["rows"]) == 1
    assert page2["rows"][0]["candidate_ids"] == [client["id"]]
    rejected = post(crm, f"/imports/{batch['id']}/confirm", {"decisions": []}, 409)
    assert rejected["code"] == "IMPORT_CONFLICT_UNRESOLVED"
    decisions = [{"row_id": row["id"], "action": "skip"} for row in batch["rows"]]
    decisions.append({"row_id": page2["rows"][0]["id"], "action": "update", "match_id": client["id"]})
    post(crm, f"/imports/{batch['id']}/confirm", {"decisions": decisions}, 200)
    with crm["sessions"].begin() as db:
        process_import(db, {"batch_id": batch["id"], "user_id": crm["admin"].id})
        assert db.scalar(select(func.count()).select_from(Counterparty)) == 1
        assert db.get(Counterparty, client["id"]).name == "Строка 100"
        db.add(PermissionGrant(user_id=crm["manager"].id, code="imports.write", scope="own"))
    login(crm, "manager@example.com")
    assert crm["client"].get(f"/api/v1/imports/{batch['id']}").status_code == 404


def test_import_errors_before_writes_formula_lengths_and_numeric_identifiers(crm):
    login(crm)
    book = Workbook()
    sheet = book.active
    sheet.append(["Организация", "Телефон", "ИНН"])
    sheet.append(["Допустимо", "+79991112233", 123])
    sheet["C2"].number_format = "0000000000"
    sheet.append(["x" * 251, "+79991112233", ""])
    sheet.append(["=1+1", "+79991112233", ""])
    stream = io.BytesIO()
    book.save(stream)
    rows = parse_rows(stream.getvalue(), {})
    assert rows[0]["data"]["tax_id"] == "0000000123"
    assert rows[1]["errors"][0]["field"] == "name"
    assert "Формулы" in rows[2]["errors"][0]["message"]
    exported = load_workbook(io.BytesIO(xlsx([["=1+1", "+cmd"]])))
    assert exported.active["A1"].data_type == "s"
    exported.close()
    response = crm["client"].post(
        "/api/v1/imports/preview",
        files={"file": ("file.xlsx", stream.getvalue())},
        data={"mapping": json.dumps({"unknown": "Организация"})},
    )
    assert response.status_code == 422
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(ImportBatch)) == 0


def test_a02_callback_then_request_preserves_client_and_history(crm):
    login(crm, "manager@example.com")
    client = post(crm, "/counterparties", {"name": "Клиент обзвона"})
    next_at = (datetime.now(timezone.utc) + timedelta(days=2)).isoformat()
    post(crm, "/calls", {"client_id": client["id"], "result": "callback"}, 422)
    key = str(uuid4())
    body = {"client_id": client["id"], "result": "callback", "next_at": next_at}
    first = post(crm, "/calls", body, key=key)
    assert post(crm, "/calls", body, key=key)["id"] == first["id"]
    successful = post(crm, "/calls", {"client_id": client["id"], "result": "request_received"})
    created = post(
        crm,
        "/requests",
        {"client_id": client["id"], "title": "После звонка", "source_call_id": successful["id"]},
    )
    assert created["source_call_id"] == successful["id"]
    with crm["sessions"]() as db:
        task = db.get(Task, first["task_id"])
        assert task.due_at.replace(tzinfo=timezone.utc).isoformat() == next_at
        assert task.assignee_id == crm["manager"].id
        assert db.scalar(select(func.count()).select_from(Call)) == 2
        assert db.scalar(select(func.count()).select_from(Counterparty)) == 1
        assert db.scalar(select(func.count()).select_from(Request)) == 1


def test_auth_csrf_live_revocation_idor_and_version_conflict(crm):
    assert crm["client"].get("/api/v1/requests").status_code == 401
    session = login(crm)
    token = crm["client"].cookies.get("crm_session")
    with crm["sessions"]() as db:
        assert db.scalar(select(AuthSession)).token_hash != token
    rejected = crm["client"].post(
        "/api/v1/counterparties", json={"name": "Rejected"}, headers={"X-CSRF-Token": "wrong"}
    )
    assert rejected.status_code == 403
    assert rejected.json()["code"] == "CSRF_INVALID"
    denied_origin = crm["client"].post(
        "/api/v1/counterparties", json={"name": "Rejected"}, headers={"Origin": "https://untrusted.example"}
    )
    assert denied_origin.status_code == 403
    client = post(crm, "/counterparties", {"name": "Приватный клиент"})
    request = post(crm, "/requests", {"client_id": client["id"], "title": "Приватная заявка"})
    update = crm["client"].patch(
        f"/api/v1/requests/{request['id']}", json={"version": 1, "title": "Первая правка"}
    )
    assert update.status_code == 200
    conflict = crm["client"].patch(
        f"/api/v1/requests/{request['id']}", json={"version": 1, "title": "Потеря правки"}
    )
    assert conflict.status_code == 409
    assert crm["client"].get(f"/api/v1/requests/{request['id']}").json()["title"] == "Первая правка"
    login(crm, "manager@example.com")
    assert crm["client"].get(f"/api/v1/requests/{request['id']}").status_code == 404
    assert crm["client"].get("/api/v1/requests").json()["total"] == 0
    with crm["sessions"].begin() as db:
        db.get(User, crm["manager"].id).active = False
    assert crm["client"].get("/api/v1/auth/me").status_code == 401
    assert session["user"]["email"] == "owner@example.com"


def test_mfa_enrollment_gate_and_login_replay(crm, monkeypatch):
    monkeypatch.setattr(settings, "require_mfa", True)
    login(crm)
    assert crm["client"].get("/api/v1/counterparties").json()["code"] == "MFA_REQUIRED"
    setup = post(crm, "/auth/mfa/setup", {}, 200)
    code = pyotp.TOTP(setup["secret"]).now()
    post(crm, "/auth/mfa/enable", {"otp": code}, 200)
    assert crm["client"].get("/api/v1/counterparties").status_code == 200
    post(crm, "/auth/logout", {}, 200)
    invalid = crm["client"].post(
        "/api/v1/auth/login", json={"email": "owner@example.com", "password": PASSWORD}
    )
    assert invalid.status_code == 401
    login(crm, otp=code)
    replay = crm["client"].post(
        "/api/v1/auth/login", json={"email": "owner@example.com", "password": PASSWORD, "otp": code}
    )
    assert replay.status_code == 401

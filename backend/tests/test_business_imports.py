"""Business workflows: preserve demand, resolve catalogue and keep orders separate from money."""

import io
from datetime import date, datetime, timezone
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from test_commerce import approve, command, get, prepare, wave
from test_commerce import commerce as commerce
from test_crm import crm as crm
from test_crm import login, post

from app.commerce.schemas import PaymentIn
from app.crm.imports import xlsx
from app.crm.models import Counterparty, Currency, Nomenclature, ProductGroup, QuoteItem, RequestItem
from app.crm.table_imports import QUOTE_COLUMNS, read_table


def request(crm):
    with crm["sessions"].begin() as db:
        if not db.scalar(select(ProductGroup.id).where(ProductGroup.slug == "other")):
            db.add(ProductGroup(name="Прочее", slug="other"))
        if not db.scalar(select(Currency.id).where(Currency.code == "RUB")):
            db.add(Currency(name="Рубль", code="RUB"))
    client = post(crm, "/counterparties", {"name": "Рабочий клиент"})
    return post(crm, "/requests", {"title": "Исходная заявка", "client_id": client["id"]})


def preview(crm, request_id, columns, rows, kind="items"):
    response = crm["client"].post(
        f"/api/v1/requests/{request_id}/table-imports/preview",
        files={
            "file": (
                "table.xlsx",
                xlsx([columns, *rows]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"kind": kind},
    )
    assert response.status_code == 200, response.text
    return response.json()


def confirm(crm, request_id, batch_id, status=200, key=None):
    return post(crm, f"/requests/{request_id}/table-imports/{batch_id}/confirm", {}, status, key=key)


def quote_file(supplier_code, overrides=None):
    row = {
        "source_row": "1",
        "supplier_code": str(supplier_code),
        "article": "0341",
        "name": "Новый товар",
        "manufacturer": "0",
        "cas": "0",
        "linear_formula": "0",
        "purity": "0",
        "packing_value": "0",
        "packing_unit": "0",
        "quantity": "1",
        "unit_price": "100,25",
        "currency": "rub",
        "delivery_days": "14",
    }
    row.update(overrides or {})
    return [row.get(key, "") for key in QUOTE_COLUMNS]


def test_client_bases_promotion_contacts_codes_and_scope(crm):
    login(crm)
    cold = post(crm, "/counterparties", {"name": "Холодный", "client_base": "cold"})
    supplier = post(crm, "/counterparties", {"name": "Поставщик", "kind": "supplier"})
    assert isinstance(cold["internal_code"], int) and supplier["internal_code"] > cold["internal_code"]
    call = post(crm, "/calls", {"client_id": cold["id"], "result": "request_received"})
    assert cold["id"] in {
        r["id"] for r in crm["client"].get("/api/v1/counterparties?client_base=cold").json()["items"]
    }
    assert cold["id"] not in {
        r["id"] for r in crm["client"].get("/api/v1/counterparties?client_base=working").json()["items"]
    }
    post(crm, "/requests", {"title": "Рано", "client_id": cold["id"]}, 422)
    post(crm, f"/counterparties/{cold['id']}/contacts", {"name": "Закупщик"}, 422)
    promoted = post(crm, f"/counterparties/{cold['id']}/promote", {"version": cold["version"]}, 200)
    assert promoted["id"] == cold["id"] and promoted["internal_code"] == cold["internal_code"]
    assert crm["client"].get(f"/api/v1/calls?client_id={cold['id']}").json()["items"][0]["id"] == call["id"]
    first = post(
        crm,
        "/contacts",
        {"client_id": cold["id"], "name": "Иван", "department": "Стандарты", "purchase_area": "USP"},
    )
    second = post(crm, f"/counterparties/{cold['id']}/contacts", {"name": "Анна", "department": "Колонки"})
    assert {first["id"], second["id"]} <= {
        r["id"] for r in crm["client"].get("/api/v1/contacts").json()["items"]
    }
    assert crm["client"].get(f"/api/v1/contacts/{first['id']}").json()["department"] == "Стандарты"
    login(crm, "manager@example.com")
    assert crm["client"].get(f"/api/v1/contacts/{first['id']}").status_code == 404
    assert not crm["client"].get("/api/v1/contacts").json()["items"]


def test_range_of_codes_is_unique_for_batched_orm_inserts(crm):
    with crm["sessions"].begin() as db:
        rows = [
            Counterparty(name=f"Поставщик {i}", kind="supplier", owner_id=crm["admin"].id) for i in range(40)
        ]
        db.add_all(rows)
        db.flush()
        assert len({row.internal_code for row in rows}) == len(rows)


def test_raw_table_roundtrip_rfq_and_repeat_import(crm):
    login(crm)
    req = request(crm)
    columns = ["Product Name (English)", "Manufacturer", "Cat. No.", "Packaging", "Quantity", "Любая колонка"]
    rows = [
        ["Сыворотка\nPanel I", "Chromsystems", "0341", "5 × 3 ml", "1 pack", "=literal text"],
        ["Контроль", "Chromsystems", "0342", "7 × 3 ml", "По согласованию", "Значение"],
    ]
    batch = preview(crm, req["id"], columns, rows)
    assert batch["summary"] == {"total": 2, "checked": 2, "create_nomenclature": 0, "errors": 0}
    confirm(crm, req["id"], batch["id"])
    assert confirm(crm, req["id"], batch["id"])["result"]["imported"] == 2
    again = preview(crm, req["id"], columns, rows)
    assert confirm(crm, req["id"], again["id"])["result"]["already_imported"] is True
    items = crm["client"].get(f"/api/v1/requests/{req['id']}/items").json()["items"]
    assert len(items) == 2 and items[0]["source_values"] == rows[0]
    assert not items[0]["nomenclature_id"]
    supplier = post(crm, "/counterparties", {"name": "Ответчик", "kind": "supplier"})
    rfq = post(
        crm,
        f"/requests/{req['id']}/rfqs",
        {
            "idempotency_key": str(uuid4()),
            "supplier_id": supplier["id"],
            "item_ids": [row["id"] for row in items],
            "response_due": date.today().isoformat(),
        },
        200,
    )
    response = crm["client"].get(f"/api/v1/rfqs/{rfq['id']}/file")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), data_only=False)
    exported = list(book.active.values)
    assert exported[0] == tuple(columns)
    assert {tuple(row) for row in exported[1:]} == {tuple(row) for row in rows}
    assert all(cell.data_type != "f" for row in book.active for cell in row)
    book.close()
    started = post(crm, f"/requests/{req['id']}/items/start", {"item_ids": [row["id"] for row in items]}, 200)
    assert started["count"] == 2
    assert all(
        row["work_status"] == "in_progress"
        for row in crm["client"].get(f"/api/v1/requests/{req['id']}/items").json()["items"]
    )


def test_quote_import_matches_articles_creates_only_catalogue_and_keeps_zero_fields_empty(crm):
    login(crm)
    req = request(crm)
    supplier = post(crm, "/counterparties", {"name": "Поставщик", "kind": "supplier"})
    raw = preview(
        crm,
        req["id"],
        ["Name", "Article", "Quantity"],
        [["Заявка A", "0341", "1"], ["Заявка B", "002-new", "3"]],
    )
    confirm(crm, req["id"], raw["id"])
    known = post(
        crm,
        "/nomenclatures",
        {"name": "Известный", "article": "0341", "packings": [{"value": "100", "unit": "mg"}]},
    )
    rows = [
        quote_file(supplier["internal_code"]),
        quote_file(supplier["internal_code"], {"source_row": "2", "article": "002-new", "quantity": "3"}),
    ]
    batch = preview(crm, req["id"], list(QUOTE_COLUMNS.values()), rows, "quotes")
    assert batch["summary"] == {"total": 2, "checked": 1, "create_nomenclature": 1, "errors": 0}
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(Nomenclature)) == 1
    result = confirm(crm, req["id"], batch["id"])
    assert result["result"]["imported"] == 2 and result["result"]["nomenclatures_created"] == 1
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(Counterparty)) == 2
        new = db.scalar(select(Nomenclature).where(Nomenclature.article == "002-new"))
        assert new.cas is None and new.manufacturer is None and new.linear_formula is None
        quotes = list(db.scalars(select(QuoteItem)))
        assert len(quotes) == 2 and {q.nomenclature_id for q in quotes} == {known["id"], new.id}
        assert all(q.unit_price == Decimal("100.25") and q.delivery_days == 14 for q in quotes)
    assert confirm(crm, req["id"], batch["id"])["result"]["imported"] == 2


def test_bulk_quotes_cross_regular_sheet_limit_and_are_atomic(crm):
    login(crm)
    req = request(crm)
    supplier = post(crm, "/counterparties", {"name": "Большой поставщик", "kind": "supplier"})
    raw = preview(
        crm,
        req["id"],
        ["Name", "Article", "Quantity"],
        [[f"Потребность {i}", f"BULK-{i:04}", "1"] for i in range(150)],
    )
    confirm(crm, req["id"], raw["id"])
    batch = preview(
        crm,
        req["id"],
        list(QUOTE_COLUMNS.values()),
        [
            quote_file(
                supplier["internal_code"],
                {
                    "source_row": str(i + 1),
                    "article": f"BULK-{i:04}",
                    "name": f"Товар {i}",
                },
            )
            for i in range(150)
        ],
        "quotes",
    )
    assert batch["summary"]["total"] == 150 and len(batch["rows"]) == 100
    result = confirm(crm, req["id"], batch["id"])["result"]
    assert result["imported"] == result["nomenclatures_created"] == 150
    assert len(result["quote_sheet_ids"]) == 2
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(QuoteItem)) == 150
        assert db.scalar(select(func.count()).select_from(Nomenclature)) == 150
        assert db.scalar(select(func.count()).select_from(Counterparty)) == 2


@pytest.mark.parametrize(
    "overrides",
    [
        {"supplier_code": "999999"},
        {"currency": "ZZZ"},
        {"name": "0"},
        {"quantity": "1.5"},
        {"unit_price": "-1"},
        {"delivery_days": "x"},
    ],
)
def test_invalid_quote_file_never_writes_suppliers_or_catalogue(crm, overrides):
    login(crm)
    req = request(crm)
    supplier = post(crm, "/counterparties", {"name": "Поставщик", "kind": "supplier"})
    raw = preview(crm, req["id"], ["Name", "Article", "Quantity"], [["Заявка", "0341", "1"]])
    confirm(crm, req["id"], raw["id"])
    batch = preview(
        crm,
        req["id"],
        list(QUOTE_COLUMNS.values()),
        [quote_file(supplier["internal_code"], overrides)],
        "quotes",
    )
    assert batch["summary"]["errors"] == 1
    assert confirm(crm, req["id"], batch["id"], 422)["code"] == "IMPORT_ERRORS"
    errors = crm["client"].get(f"/api/v1/requests/{req['id']}/table-imports/{batch['id']}/errors.xlsx")
    assert errors.status_code == 200 and load_workbook(io.BytesIO(errors.content)).active.max_row == 2
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(Counterparty)) == 2
        assert db.scalar(select(func.count()).select_from(Nomenclature)) == 0
        assert db.scalar(select(func.count()).select_from(QuoteItem)) == 0


def test_stale_and_ambiguous_quote_imports_are_blocked(crm):
    login(crm)
    req = request(crm)
    supplier = post(crm, "/counterparties", {"name": "Поставщик", "kind": "supplier"})
    raw = preview(crm, req["id"], ["Name", "Article", "Quantity"], [["Заявка", "0341", "1"]])
    confirm(crm, req["id"], raw["id"])
    batch = preview(
        crm, req["id"], list(QUOTE_COLUMNS.values()), [quote_file(supplier["internal_code"])], "quotes"
    )
    with crm["sessions"].begin() as db:
        item = db.scalar(select(RequestItem).where(RequestItem.request_id == req["id"]))
        item.version += 1
    assert confirm(crm, req["id"], batch["id"], 409)["code"] == "IMPORT_STALE"
    for name in ("Brand A", "Brand B"):
        post(
            crm,
            "/nomenclatures",
            {"name": name, "article": "0341", "packings": [{"value": "1", "unit": "pcs"}]},
        )
    batch = preview(
        crm, req["id"], list(QUOTE_COLUMNS.values()), [quote_file(supplier["internal_code"])], "quotes"
    )
    assert batch["summary"]["errors"] == 1 and "неоднозначен" in batch["rows"][0]["errors"][0]


def test_different_supplier_packings_for_one_source_are_allowed(crm):
    login(crm)
    req = request(crm)
    suppliers = [
        post(crm, "/counterparties", {"name": f"Поставщик {i}", "kind": "supplier"}) for i in range(2)
    ]
    raw = preview(crm, req["id"], ["Name", "Article", "Quantity"], [["Потребность", "0341", "1"]])
    confirm(crm, req["id"], raw["id"])
    batch = preview(
        crm,
        req["id"],
        list(QUOTE_COLUMNS.values()),
        [
            quote_file(suppliers[0]["internal_code"], {"packing_value": "100", "packing_unit": "mg"}),
            quote_file(suppliers[1]["internal_code"], {"packing_value": "200", "packing_unit": "mg"}),
        ],
        "quotes",
    )
    assert batch["summary"]["errors"] == 0
    confirm(crm, req["id"], batch["id"])
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(Nomenclature)) == 1
        assert db.scalar(select(func.count()).select_from(QuoteItem)) == 2


def test_text_and_xlsx_readers_preserve_identifiers_and_reject_formulas():
    assert read_table("Name\tCat. No.\nA\t0341\n".encode(), "input.txt")[1] == [["A", "0341"]]
    assert read_table(b"| Name | Article |\n| --- | --- |\n| A | 0341 |", "input.txt")[1] == [["A", "0341"]]
    book = Workbook()
    sheet = book.active
    sheet.append(["Name", "Article"])
    sheet.append(["A", 341])
    sheet["B2"].number_format = "0000"
    stream = io.BytesIO()
    book.save(stream)
    assert read_table(stream.getvalue(), "input.xlsx")[1][0][1] == "0341"
    sheet["B2"] = "=1+1"
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(Exception, match="формул"):
        read_table(stream.getvalue(), "input.xlsx")


def test_payment_currency_is_normalized_and_order_does_not_require_payment(commerce):
    env = commerce
    state = prepare(env, with_invoice=True)
    approval = approve(env, state)
    shipment = wave(env, "unpaid-order")
    allocation = command(
        env,
        f"/waves/{shipment['id']}/allocations",
        {"execution_id": state["execution"]["id"], "approval_id": approval["id"], "quantity": "1"},
    )
    command(
        env,
        f"/allocations/{allocation['id']}/events",
        {
            "kind": "ordered",
            "quantity": "1",
            "occurred_at": datetime.now(timezone.utc).isoformat(),
            "reason": "Заказ без платежа",
        },
    )
    assert not get(env, f"/requests/{env['request_id']}/payments")["items"]
    payment = command(
        env,
        f"/requests/{env['request_id']}/payments",
        {
            "invoice_id": state["invoice"]["id"],
            "amount": "10",
            "currency": " Rub ",
            "payment_date": date.today().isoformat(),
            "number": "135",
        },
    )
    assert payment["currency"] == "RUB"
    assert (
        PaymentIn(
            idempotency_key=str(uuid4()), amount="10", currency="eur", payment_date=date.today(), number="P"
        ).currency
        == "EUR"
    )

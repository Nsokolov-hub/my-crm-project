"""Purchase order to a supplier in the proposal layout: PDF to send and XLSX to edit, both in English."""

import io
from datetime import date
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.commerce.files import LOGO, PROPOSAL_INK, PROPOSAL_LIME, pdf_decimal, pdf_fonts, safe_cell

PO_SIGNATURE = {"position": "General Director", "name": "T.F. Gilmutdinov"}
PO_UNITS = {"pcs": "pcs", "kg": "kg", "g": "g", "mg": "mg", "l": "l", "ml": "ml"}


def po_date(value) -> str:
    return (value if isinstance(value, date) else date.fromisoformat(str(value))).strftime("%d.%m.%Y")


def party_lines(party: dict) -> list[str]:
    """Address, tax number and e-mail of a party, whichever of their spellings the card uses."""
    saved = party.get("details")
    details: dict = saved if isinstance(saved, dict) else {}

    def first(*keys):
        return next((str(details[key]).strip() for key in keys if str(details.get(key) or "").strip()), "")

    address = first("legal_address", "Юридический адрес", "address", "Адрес", "Фактический адрес")
    tax_id = party.get("tax_id") or first("ИНН", "inn", "tax_id")
    email = party.get("email") or first("email", "E-mail", "Email")
    return [line for line in (address, f"Tax ID: {tax_id}" if tax_id else "", email) if line]


def po_lines(order) -> list[dict]:
    rows = []
    for row in order.snapshot["lines"]:
        quantity, price = Decimal(row["quantity"]), Decimal(row["unit_price"])
        rows.append({**row, "quantity": quantity, "unit_price": price, "total": quantity * price})
    return rows


def po_logistics(order, with_logistics: bool = True) -> Decimal:
    if not with_logistics:
        return Decimal("0")
    return Decimal(str((order.snapshot.get("logistics") or {}).get("amount") or "0"))


def po_total(order, with_logistics: bool = True) -> Decimal:
    return sum((row["total"] for row in po_lines(order)), Decimal("0")) + po_logistics(order, with_logistics)


def logistics_note(order, with_logistics: bool = True) -> str:
    """How the PO speaks of freight it does not carry: paid separately, or simply not included."""
    separately = not with_logistics or (order.snapshot.get("logistics") or {}).get("separately")
    return ("International logistics: paid separately for the consolidated shipment" if separately
            else "International logistics: not included")


def po_terms(order) -> list[tuple[str, str]]:
    return [
        ("Contract", order.contract or "—"),
        ("Payment conditions", order.payment_terms or "—"),
        ("Delivery terms", order.delivery_terms or "—"),
        ("Expected delivery date", po_date(order.expected_date)),
    ]


def purchase_order_pdf(order, with_logistics: bool = True) -> bytes:
    font_name, bold_font = pdf_fonts()
    ink = colors.HexColor(f"#{PROPOSAL_INK}")
    muted = colors.HexColor("#63716B")
    lime = colors.HexColor(f"#{PROPOSAL_LIME}")
    pale = colors.HexColor("#F5F8F1")
    border = colors.HexColor("#DDE6DA")
    customer, supplier = order.snapshot["customer"], order.snapshot["supplier"]
    currency = order.currency
    issued = po_date(order.created_at.date())
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=34, rightMargin=34, topMargin=104, bottomMargin=50,
        title=f"Purchase order {order.number}", author=customer["name"],
    )
    width = A4[0] - 80
    body = ParagraphStyle("po-body", fontName=font_name, fontSize=8.5, leading=12.5, textColor=ink)
    small = ParagraphStyle("po-small", parent=body, fontSize=7.3, leading=10.5, textColor=muted)
    label = ParagraphStyle("po-label", parent=small, fontName=bold_font, fontSize=7, leading=10)
    heading = ParagraphStyle("po-heading", parent=body, fontName=bold_font, fontSize=22, leading=26)
    name = ParagraphStyle("po-name", parent=body, fontName=bold_font, fontSize=10, leading=14)
    strong = ParagraphStyle("po-strong", parent=body, fontName=bold_font)
    right = ParagraphStyle("po-right", parent=body, alignment=TA_RIGHT, splitLongWords=0)
    header = ParagraphStyle("po-column", parent=label, textColor=colors.white, leading=10)
    header_right = ParagraphStyle("po-column-right", parent=header, alignment=TA_RIGHT)
    counter = ParagraphStyle("po-counter", parent=small, splitLongWords=0)
    white_label = ParagraphStyle("po-total-label", parent=label, textColor=colors.HexColor("#D3E5CB"))
    amount_text = f"{pdf_decimal(po_total(order, with_logistics), money=True)} {currency}"
    amount_size = min(21, 21 * (width * .44 - 30) / max(pdfmetrics.stringWidth(amount_text, bold_font, 21), 1))
    amount = ParagraphStyle("po-total", parent=body, fontName=bold_font, fontSize=amount_size,
                            leading=amount_size + 6, textColor=colors.white)

    def p(text, style=body):
        return Paragraph(escape(str(text)), style)

    def party_block(title, party):
        block = [p(title, label), Spacer(1, 6), p(party["name"], name), Spacer(1, 4)]
        for line in party_lines(party):
            block.extend([p(line, small), Spacer(1, 2)])
        return block

    content = [p("Purchase order", heading), Spacer(1, 8),
               p("Please supply the items below on the terms of this order and confirm it in reply.", body),
               Spacer(1, 16)]
    parties = Table([[party_block("CUSTOMER", customer), "", party_block("SUPPLIER", supplier)]],
                    colWidths=[(width - 16) / 2, 16, (width - 16) / 2])
    parties.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (0, 0), pale),
        ("BACKGROUND", (2, 0), (2, 0), pale), ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12), ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    terms = po_terms(order)
    conditions = Table(
        [[[p(key.upper(), label), Spacer(1, 3), p(value)] for key, value in terms[index:index + 2]]
         for index in (0, 2)],
        colWidths=[width / 2, width / 2],
    )
    conditions.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), .5, border),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    content.extend([parties, Spacer(1, 12), conditions, Spacer(1, 14)])
    lines = po_lines(order)
    section = Table([[p("ORDER ITEMS", label),
                      p(f"Items: {len(lines)} · Currency: {currency}",
                        ParagraphStyle("po-section-meta", parent=small, alignment=TA_RIGHT))]],
                    colWidths=[width / 2, width / 2])
    section.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    widths = [26, width - 26 - 52 - 86 - 96, 52, 86, 96]
    rows = [[p(value, header_right if index > 2 else header)
             for index, value in enumerate(("№", "Item", "Qty", f"Price, {currency}", f"Total, {currency}"))]]

    def money(value, column_width):
        text = pdf_decimal(value, money=True)
        size = min(8.5, 8.5 * (column_width - 18) / max(pdfmetrics.stringWidth(text, font_name, 8.5), 1))
        return p(text, ParagraphStyle("po-price", parent=right, fontSize=size, leading=size + 4))

    for index, row in enumerate(lines, 1):
        facts = " · ".join(
            f"{caption}: {row[key]}" for caption, key in
            (("Manufacturer", "manufacturer"), ("Cat.", "article"), ("Packaging", "packaging")) if row.get(key)
        )
        item = [p(row["name"], strong)] + ([Spacer(1, 2), p(facts, small)] if facts else [])
        rows.append([p(index, counter), item, p(pdf_decimal(row["quantity"])),
                     money(row["unit_price"], widths[3]), money(row["total"], widths[4])])
    logistics = po_logistics(order, with_logistics)
    if logistics:
        rows.append([p(len(lines) + 1, counter), [p("International logistics", strong)], p("1"),
                     money(logistics, widths[3]), money(logistics, widths[4])])
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ink), ("LINEABOVE", (0, 0), (-1, 0), 3, lime),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pale]),
        ("LINEBELOW", (0, 1), (-1, -1), .5, border), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("LEFTPADDING", (0, 0), (0, -1), 6), ("RIGHTPADDING", (0, 0), (0, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    content.extend([section, Spacer(1, 6), table, Spacer(1, 14)])
    goods = sum((row["total"] for row in lines), Decimal("0"))
    summary = [p("ORDER SUMMARY", label), Spacer(1, 6),
               p(f"Goods: {pdf_decimal(goods, money=True)} {currency}"), Spacer(1, 4),
               p(f"International logistics: {pdf_decimal(logistics, money=True)} {currency}" if logistics
                 else logistics_note(order, with_logistics)), Spacer(1, 4),
               p(f"Delivery terms: {order.delivery_terms or '—'}", small)]
    total = [p("ORDER TOTAL", white_label), Spacer(1, 5), p(amount_text, amount), Spacer(1, 4),
             p(f"Payment: {order.payment_terms or 'as per contract'}", white_label)]
    decision = Table([[summary, total]], colWidths=[width * .56, width * .44])
    decision.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (1, 0), (1, 0), ink),
        ("LINEABOVE", (1, 0), (1, 0), 3, lime), ("LEFTPADDING", (0, 0), (0, 0), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 20), ("LEFTPADDING", (1, 0), (1, 0), 15),
        ("RIGHTPADDING", (1, 0), (1, 0), 15), ("TOPPADDING", (0, 0), (-1, -1), 11),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
    ]))
    signature = Table([[p(PO_SIGNATURE["position"]), p(PO_SIGNATURE["name"], right)]],
                      colWidths=[width * .6, width * .4])
    signature.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    notice = p("Please confirm the order, the prices and the delivery date in reply. "
               "Ship with the order number on the invoice and the packing list.", small)
    content.append(KeepTogether([decision, Spacer(1, 12), notice, Spacer(1, 14), signature]))

    def page_brand(canvas, document):
        canvas.saveState()
        page_width, page_height = A4
        canvas.drawImage(str(LOGO), 40, page_height - 87, width=145, height=145 * 425 / 1280,
                         preserveAspectRatio=True, mask="auto")
        canvas.setFillColor(muted)
        canvas.setFont(font_name, 7)
        canvas.drawRightString(page_width - 40, page_height - 48, "PURCHASE ORDER")
        canvas.setFillColor(ink)
        canvas.setFont(bold_font, 11)
        canvas.drawRightString(page_width - 40, page_height - 68, f"№{order.number} dated {issued}")
        canvas.setStrokeColor(border)
        canvas.setLineWidth(.6)
        canvas.line(40, page_height - 98, page_width - 40, page_height - 98)
        canvas.line(40, 42, page_width - 40, 42)
        canvas.setFont(bold_font, 8)
        canvas.drawString(40, 27, "ogk-chem.ru")
        canvas.linkURL("https://ogk-chem.ru", (40, 24, 40 + pdfmetrics.stringWidth("ogk-chem.ru", bold_font, 8), 36),
                       relative=0)
        canvas.setFont(font_name, 8)
        canvas.setFillColor(muted)
        canvas.drawString(160, 27, "info@ogk-chem.ru")
        canvas.drawRightString(page_width - 40, 27, f"{document.page}")
        canvas.restoreState()

    doc.build(content, onFirstPage=page_brand, onLaterPages=page_brand)
    return buffer.getvalue()


def purchase_order(order, with_logistics: bool = True) -> bytes:
    """The same order as an editable workbook, laid out like the proposal workbook."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Purchase order"
    customer, supplier = order.snapshot["customer"], order.snapshot["supplier"]
    currency = order.currency
    headers = ["№", "Name", "Manufacturer", "Cat.", "Packaging", "Qty", f"Price ({currency})",
               f"Total ({currency})"]
    columns = len(headers)
    ink_font = Font(name="Arial", size=10, color=PROPOSAL_INK)
    thin = Side(style="thin", color="DDE6DA")
    sheet.append(["ОГК-ХИМ"])
    sheet.append([f"Purchase order №{order.number} dated {po_date(order.created_at.date())}"])
    sheet.append([])
    info = [
        ("Customer", customer["name"], "\n".join(party_lines(customer))),
        ("Supplier", supplier["name"], "\n".join(party_lines(supplier))),
        *((key, value, None) for key, value in po_terms(order)),
        ("Currency", currency, None),
    ]
    # As in the proposal workbook: the label spans A:B, the value the rest; a party adds its details.
    for key, value, extra in info:
        sheet.append([key, None, safe_cell(value), None, safe_cell(extra) if extra else None])
        row = sheet.max_row
        sheet.cell(row, 1).font = Font(name="Arial", bold=True, color=PROPOSAL_INK)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        if extra:
            sheet.cell(row, 3).font = Font(name="Arial", bold=True, color=PROPOSAL_INK)
            sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
            sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=columns)
            sheet.row_dimensions[row].height = max(30, 15 * (extra.count("\n") + 1))
        else:
            sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=columns)
        for cell in sheet[row][:columns]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
    sheet.append([])
    sheet.append(headers)
    header_row = sheet.max_row
    lines = po_lines(order)
    for index, row in enumerate(lines, 1):
        sheet.append([index, *(safe_cell(row.get(key)) for key in ("name", "manufacturer", "article", "packaging")),
                      row["quantity"], row["unit_price"], row["total"]])
    logistics = po_logistics(order, with_logistics)
    if logistics:
        sheet.append([len(lines) + 1, "International logistics", None, None, None, Decimal("1"), logistics,
                      logistics])
    last_line = sheet.max_row
    goods = sum((row["total"] for row in lines), Decimal("0"))
    sheet.append([])
    summary = [("Goods", goods),
               ("International logistics", logistics) if logistics else (logistics_note(order, with_logistics), None),
               (f"Total ({currency})", po_total(order, with_logistics))]
    for key, value in summary:
        sheet.append([key, None, None, None, None, None, None, value])
        row = sheet.max_row
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns - 1)
        sheet.cell(row, columns).number_format = "#,##0.00"
        for cell in sheet[row][:columns]:
            cell.font = Font(name="Arial", bold=True, color=PROPOSAL_INK)
    total_row = sheet.max_row
    for cell in sheet[total_row][:columns]:
        cell.fill = PatternFill("solid", fgColor=PROPOSAL_LIME)
        cell.font = Font(name="Arial", bold=True, color=PROPOSAL_INK, size=12)
    signature_row = sheet.max_row + 3
    sheet.cell(signature_row, 1, PO_SIGNATURE["position"]).font = ink_font
    sheet.merge_cells(start_row=signature_row, start_column=1, end_row=signature_row, end_column=4)
    sheet.cell(signature_row, 6, PO_SIGNATURE["name"]).font = ink_font
    sheet.cell(signature_row, 6).alignment = Alignment(horizontal="right")
    sheet.merge_cells(start_row=signature_row, start_column=6, end_row=signature_row, end_column=columns)
    sheet.cell(signature_row + 1, 1, "ogk-chem.ru · info@ogk-chem.ru").font = ink_font

    logo = ExcelImage(str(LOGO))
    logo.width, logo.height = 210, 70
    sheet.add_image(logo, "F1")
    sheet.row_dimensions[1].height = 60
    sheet.cell(1, 1).font = Font(name="Arial", bold=True, color=PROPOSAL_INK, size=14)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sheet.cell(2, 1).font = Font(name="Arial", bold=True, color=PROPOSAL_INK, size=18)
    sheet.row_dimensions[2].height = 40
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PROPOSAL_INK)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[header_row].height = 32
    for row in range(header_row + 1, last_line + 1):
        for cell in sheet[row][:columns]:
            cell.fill = PatternFill("solid", fgColor="F5F8F1" if row % 2 else "FFFFFF")
            cell.border = Border(bottom=Side(style="thin", color="B6D49B"))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in sheet[row][6:columns]:
            cell.number_format = "#,##0.00"
        sheet.row_dimensions[row].height = max(28, 16 * ((len(str(sheet.cell(row, 2).value or "")) + 44) // 45))
    for column, width in zip("ABCDEFGH", [7, 46, 22, 16, 14, 10, 16, 18], strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()

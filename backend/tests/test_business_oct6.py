"""October 6 acceptance: quantity, procurement handoff, forecasts and immutable exports."""

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from test_business_imports import confirm, preview, quote_file, request
from test_business_oct5 import cmd, structured
from test_crm import crm, post  # noqa: F401

from app.commerce.files import read_file
from app.commerce.models import CommercialDocument, Execution, WaveAllocation
from app.core.models import Notification
from app.crm.models import QuoteItem, Request, RequestItem
from app.crm.table_imports import QUOTE_COLUMNS


def wave(env, wave_id):
    return next(row for row in env["client"].get("/api/v1/waves").json()["items"] if row["id"] == wave_id)


def release(env, req, calculation, quantity):
    proposal = cmd(env, f"/requests/{req['id']}/proposals", {
        "calculation_id": calculation["id"], "valid_until": (date.today() + timedelta(days=15)).isoformat(),
        "terms": "Отсрочка, оплата после поставки",
    })
    accepted = cmd(env, f"/proposals/{proposal['id']}/accept", {
        "version": proposal["version"], "reason": "Клиент подтвердил новый объём",
        "lines": [{"line_id": calculation["snapshot"]["lines"][0]["line_id"], "quantity": quantity}],
    })
    return proposal, accepted["executions"][0]


def handoff(env, req, execution):
    approval = cmd(env, f"/requests/{req['id']}/approvals", {
        "execution_ids": [execution["id"]], "reviewer_id": env["admin"].id,
    })
    return cmd(env, f"/approvals/{approval['id']}/decision", {
        "version": approval["version"], "decision": "approved", "reason": "Передать с отсрочкой",
    })


@pytest.mark.parametrize("quantity", ["2", "100"])
def test_quantity_is_independent_of_quote_and_initial_demand(crm, quantity):  # noqa: F811
    req, data, seller, _ = structured(crm)
    data["selections"][0]["quantity"] = quantity
    quote_id = data["selections"][0]["quote_item_id"]
    seller = crm["client"].patch(f"/api/v1/sellers/{seller['id']}", json={
        "version": seller["version"], "name": seller["name"], "currency": seller["currency"],
        "details": {**seller["details"], "Подписант": "Иван Иванов", "Должность подписанта": "Директор"},
    }).json()
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", data)
    assert Decimal(calculation["snapshot"]["totals"]["purchase_rub"]) == Decimal(quantity) * 1000
    proposal, execution = release(crm, req, calculation, quantity)
    assert proposal["snapshot"]["lines"][0]["description"] == "Стандартный образец Sample (100 mg) Арт. A123"
    assert proposal["snapshot"]["signature"] == {"name": "Гильмутдинов Т.Ф", "position": "Генеральный директор"}
    assert "bank_details" not in proposal["snapshot"]
    invoice = cmd(crm, f"/requests/{req['id']}/invoices", {
        "proposal_id": proposal["id"], "due_date": (date.today() + timedelta(days=30)).isoformat(),
        "terms": "Отсрочка", "lines": [{"execution_id": execution["id"], "quantity": quantity}],
    })
    assert Decimal(invoice["snapshot"]["lines"][0]["quantity"]) == Decimal(quantity)
    assert invoice["snapshot"]["bank_details"]["Банк"] == "Банк A"
    with crm["sessions"]() as db:
        quote = db.get(QuoteItem, quote_id)
        assert quote.quantity == 1
        assert db.get(RequestItem, quote.source_request_item_id).quantity == 1
        proposal_book = load_workbook(io.BytesIO(read_file(db.get(CommercialDocument, proposal["id"]).files["xlsx"])))
        invoice_book = load_workbook(io.BytesIO(read_file(db.get(CommercialDocument, invoice["id"]).files["xlsx"])))
        proposal_values = [cell.value for row in proposal_book.active for cell in row]
        invoice_values = [cell.value for row in invoice_book.active for cell in row]
        assert "Банк A" not in proposal_values
        assert "Банк A" in invoice_values
        assert "Генеральный директор" in proposal_values and "Гильмутдинов Т.Ф" in proposal_values
        assert "Директор" in invoice_values
        assert "Наименование" in proposal_values and "Цена за единицу с НДС" in proposal_values
        assert "Без налога" not in proposal_values
    assert crm["client"].get("/api/v1/supplier-orders/positions").json()["items"] == []
    handoff(crm, req, execution)
    positions = crm["client"].get("/api/v1/supplier-orders/positions").json()["items"]
    assert Decimal(positions[0]["quantity"]) == Decimal(quantity)
    cmd(crm, "/supplier-orders", {"execution_ids": [execution["id"]], "seller_id": seller["id"],
        "expected_date": (date.today() + timedelta(days=30)).isoformat()})
    # The released line still cannot be accepted twice beyond its own quantity.
    response = cmd(crm, f"/proposals/{proposal['id']}/accept", {
        "version": proposal["version"] + 1, "reason": "Повторное принятие",
        "lines": [{"line_id": calculation["snapshot"]["lines"][0]["line_id"], "quantity": "1"}],
    }, status=422)
    assert response["code"] == "QUANTITY_EXCEEDED"


def test_forecast_replaces_real_orders_and_calculations_do_not_fill_wave(crm):  # noqa: F811
    req, data, seller, supplier = structured(crm, mode="IMPORT")
    with crm["sessions"]() as db:
        quote = db.get(QuoteItem, data["selections"][0]["quote_item_id"])
        item = db.get(RequestItem, quote.source_request_item_id)
        group_id = item.product_group_id
        item_data = {"nomenclature_id": quote.nomenclature_id, "packing_id": quote.packing_id,
            "quantity": "1", "unit_price": "1000", "currency_id": quote.currency_id, "delivery_days": 7}
    initial = wave(crm, req["wave_id"])
    forecast_data = {"product_group_id": group_id, "target_quantity": "10", "unit_price_rub": "1000",
        "weight_per_unit": "1", "reason": "Минимальный объём запуска"}
    current = cmd(crm, f"/waves/{initial['id']}/forecasts", {**forecast_data, "version": initial["version"]}, method="put")
    assert Decimal(current["financial_summary"]["real_quantity"]) == 0
    assert Decimal(current["financial_summary"]["forecast_quantity"]) == 10
    data["selections"][0]["quantity"] = "2"
    data["expenses"] = [{"name": "Общая логистика", "amount": "1000", "currency": "RUB", "scope": "WAVE",
        "stage": "INTERNATIONAL_LOGISTICS", "method": "BY_QUANTITY"}]
    current = cmd(crm, f"/waves/{initial['id']}/budget", {"version": current["version"], "profile_id": data["profile_id"],
        "expenses": data["expenses"], "reason": "Бюджет запуска"}, method="put")
    first = cmd(crm, f"/requests/{req['id']}/calculations", data)
    assert Decimal(first["snapshot"]["wave_distribution"]["total_quantity"]) == 10
    assert Decimal(first["snapshot"]["wave_distribution"]["forecast_quantity"]) == 8
    assert Decimal(first["snapshot"]["totals"]["international_logistics"]) == 200
    assert wave(crm, initial["id"])["allocations"] == []
    proposal, execution = release(crm, req, first, "2")
    assert wave(crm, initial["id"])["allocations"] == []
    assert crm["client"].get("/api/v1/supplier-orders/positions").json()["items"] == []
    handoff(crm, req, execution)
    actual = wave(crm, initial["id"])
    assert len(actual["allocations"]) == 1
    assert Decimal(actual["financial_summary"]["real_quantity"]) == 2
    assert Decimal(actual["financial_summary"]["forecast_quantity"]) == 8
    assert Decimal(actual["forecasts"][0]["remaining_quantity"]) == 8
    assert Decimal(actual["forecasts"][0]["filled_quantity"]) == 2
    assert Decimal(actual["forecasts"][0]["target_quantity"]) == 10
    assert actual["financial_summary"]["actual"]["status"] == "current"
    # A different customer's current selection consumes only its scenario forecast.
    peer = request(crm)
    with crm["sessions"].begin() as db:
        db.get(Request, peer["id"]).seller_id = seller["id"]
    peer_quote = post(crm, f"/requests/{peer['id']}/quote-sheets", {"supplier_id": supplier["id"], "items": [item_data]})
    peer = crm["client"].get(f"/api/v1/requests/{peer['id']}").json()
    assigned = cmd(crm, f"/requests/{peer['id']}/wave", {"request_version": peer["version"], "wave_id": initial["id"]}, method="put")
    peer_payload = {"request_version": assigned["version"], "profile_id": data["profile_id"],
        "selections": [{"quote_item_id": peer_quote["items"][0]["id"], "quantity": "2"}], "delivery_required": False}
    previewed = cmd(crm, f"/requests/{peer['id']}/calculations/preview", peer_payload)["snapshot"]
    distribution = previewed["wave_distribution"]
    assert Decimal(distribution["existing_quantity"]) == 2
    assert Decimal(distribution["selected_quantity"]) == 2
    assert Decimal(distribution["forecast_quantity"]) == 6
    assert Decimal(distribution["total_quantity"]) == 10
    assert Decimal(previewed["totals"]["international_logistics"]) == 200
    assert sum(Decimal(x) for x in previewed["expense_allocations"][0]["wave_parts"].values()) == 1000
    # A wave version/forecast change invalidates a preview without mutating released documents.
    cmd(crm, f"/waves/{initial['id']}/forecasts", {**forecast_data, "version": actual["version"], "active": False}, method="put")
    conflict = cmd(crm, f"/requests/{peer['id']}/calculations", {**peer_payload,
        "expected_wave_digest": distribution["digest"]}, status=409)
    assert conflict["code"] == "WAVE_CHANGED"
    fresh = cmd(crm, f"/requests/{peer['id']}/calculations/preview", peer_payload)["snapshot"]
    assert Decimal(fresh["wave_distribution"]["total_quantity"]) == 4
    assert Decimal(fresh["totals"]["international_logistics"]) == 500
    with crm["sessions"]() as db:
        assert db.get(CommercialDocument, proposal["id"]).snapshot == proposal["snapshot"]
        assert db.get(Execution, execution["id"]).procurement_at is not None
        assert db.get(Execution, execution["id"]).financing_deficit is True
        assert db.query(WaveAllocation).filter_by(wave_id=initial["id"], active=True).count() == 1


def test_large_quote_import_notifies_manager_once_and_accepts_extra_lines(crm):  # noqa: F811
    req, data, _, supplier = structured(crm)
    with crm["sessions"].begin() as db:
        db.get(Request, req["id"]).owner_id = crm["manager"].id
    rows = [quote_file(supplier["internal_code"], {"source_row": "", "article": f"EXTRA-{i:03}",
        "name": f"Дополнительный стандарт {i}", "packing_value": "50", "packing_unit": "mg",
        "product_group": "Стандартные образцы"}) for i in range(125)]
    batch = preview(crm, req["id"], list(QUOTE_COLUMNS.values()), rows, kind="quotes")
    assert batch["summary"]["errors"] == 0, batch["rows"][0]
    confirmed = confirm(crm, req["id"], batch["id"], key="same-batch")
    assert confirmed["result"]["imported"] == 125
    confirm(crm, req["id"], batch["id"], key="same-batch")
    with crm["sessions"]() as db:
        notifications = db.scalars(select(Notification).where(Notification.event_key == f"quote-import:{batch['id']}")).all()
        assert len(notifications) == 1 and notifications[0].user_id == crm["manager"].id
        assert db.query(RequestItem).filter_by(request_id=req["id"], quote_only=True).count() == 126
    all_quotes = crm["client"].get(f"/api/v1/requests/{req['id']}/quote-items?all=true").json()
    assert len(all_quotes["items"]) == 126
    latest = crm["client"].get(f"/api/v1/requests/{req['id']}").json()
    calculation = cmd(crm, f"/requests/{req['id']}/calculations", {**data, "request_version": latest["version"],
        "selections": [{"quote_item_id": row["id"], "quantity": "2"} for row in all_quotes["items"]]})
    assert len(calculation["snapshot"]["lines"]) == 126


@pytest.mark.parametrize("method", ["BY_QUANTITY", "BY_PURCHASE_VALUE", "BY_WEIGHT"])
def test_forecast_budget_bases_and_overfilled_plan(crm, method):  # noqa: F811
    req, data, _, _ = structured(crm, mode="IMPORT")
    with crm["sessions"]() as db:
        quote = db.get(QuoteItem, data["selections"][0]["quote_item_id"])
        group_id = db.get(RequestItem, quote.source_request_item_id).product_group_id
    initial = wave(crm, req["wave_id"])
    payload = {"version": initial["version"], "product_group_id": group_id,
        "target_quantity": "10", "unit_price_rub": "1000", "weight_per_unit": "1", "reason": "Запуск поставщика"}
    current = cmd(crm, f"/waves/{initial['id']}/forecasts", payload, method="put", key="forecast-retry")
    assert cmd(crm, f"/waves/{initial['id']}/forecasts", payload, method="put", key="forecast-retry") == current
    assert len(wave(crm, initial["id"])["forecasts"]) == 1
    assert cmd(crm, f"/waves/{initial['id']}/forecasts", payload, method="put", status=409)["code"] == "VERSION_CONFLICT"
    data["expenses"] = [{"name": "Транспорт", "amount": "1000", "scope": "WAVE", "stage": "GENERAL",
        "method": method, "currency": "RUB"}]
    data["selections"][0].update(quantity="2", weight="2")
    calculated = cmd(crm, f"/requests/{req['id']}/calculations/preview", data)["snapshot"]
    assert Decimal(calculated["totals"]["general_expenses"]) == 200
    assert Decimal(calculated["wave_distribution"]["forecast_quantity"]) == 8
    data["selections"][0].update(quantity="12", weight="12")
    overfilled = cmd(crm, f"/requests/{req['id']}/calculations/preview", data)["snapshot"]
    assert Decimal(overfilled["wave_distribution"]["forecast_quantity"]) == 0
    assert Decimal(overfilled["wave_distribution"]["total_quantity"]) == 12
    assert Decimal(overfilled["totals"]["general_expenses"]) == 1000

    current = cmd(crm, f"/waves/{initial['id']}/budget", {"version": current["version"],
        "profile_id": data["profile_id"], "expenses": data["expenses"], "reason": "Бюджет по выбранной базе"}, method="put")
    data["selections"][0].update(quantity="2", weight="2")
    saved = cmd(crm, f"/requests/{req['id']}/calculations", data)
    _, execution = release(crm, req, saved, "2")
    handoff(crm, req, execution)
    summary = wave(crm, initial["id"])["financial_summary"]
    assert summary["actual"]["status"] == "current"
    assert Decimal(summary["real_quantity"]) == 2
    assert Decimal(summary["forecast_quantity"]) == 8

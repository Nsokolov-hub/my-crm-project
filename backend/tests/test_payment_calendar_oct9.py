"""Plans, actual cash dates, manager decisions, invoice ACLs and dated balances."""

import io
from datetime import date
from decimal import Decimal
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import select
from test_business_oct5 import cmd
from test_crm import crm, login, post  # noqa: F401

from app.communication.models import FileRecord
from app.core.models import AppSetting, AuditEvent, PermissionGrant
from app.crm.models import Request


def invoice(env, entry, status="clean"):
    with env["sessions"].begin() as db:
        db.add(FileRecord(name="Счёт.pdf", media_type="application/pdf", size=20, sha256=uuid4().hex*2,
                          storage_key="test-calendar/"+str(uuid4()), author_id=env["admin"].id,
                          classification="general", status=status, calendar_entry_id=entry["id"]))


def planned(env, direction="expense", amount="100", **extra):
    return cmd(env, "/payment-calendar", {"direction": direction, "planned_date": "2026-10-06",
        "amount": amount, "purpose": "Оплата счёта", **extra})


def submit(env, entry):
    return cmd(env, f"/payment-calendar/{entry['id']}/submit", {"version": entry["version"]})


def decide(env, entry, decision="approved", actual_date="2026-10-08", **extra):
    status = extra.pop("status", 200)
    return cmd(env, f"/workflow-approvals/{entry['review_id']}/decision", {
        "version": entry["review_version"], "decision": decision, "reason": "Решение руководителя",
        "actual_date": actual_date, **extra}, status=status)


def calendar(env, **filters):
    response = env["client"].get("/api/v1/payment-calendar", params={
        "from_date": "2026-10-05", "to_date": "2026-10-09", **filters})
    assert response.status_code == 200, response.text
    return response.json()


def balance(env, day="2026-10-05", amount="1000", **extra):
    status = extra.pop("status", 200)
    return cmd(env, "/payment-calendar/balances", {"balance_date": day, "currency": "RUB",
        "amount": amount, "reason": "Выписка банка на начало дня", **extra}, method="put", status=status)


def test_plan_never_changes_balance_and_confirmation_uses_actual_date_once(crm):  # noqa: F811
    login(crm)
    balance(crm)
    row = planned(crm)
    invoice(crm, row)
    assert Decimal(calendar(crm)["summary"][0]["current_balance"]) == 1000
    row = submit(crm, row)
    summary = calendar(crm)["summary"][0]
    assert Decimal(summary["planned_expense"]) == 100
    assert Decimal(summary["expense"]) == 0
    assert Decimal(summary["current_balance"]) == 1000
    key = str(uuid4())
    payload = {"version": row["review_version"], "decision": "approved", "reason": "Оплата произведена",
               "actual_date": "2026-10-08"}
    route = f"/workflow-approvals/{row['review_id']}/decision"
    result = cmd(crm, route, payload, key=key)
    assert cmd(crm, route, payload, key=key) == result
    data = calendar(crm)
    paid = data["items"][0]
    assert paid["status"] == "confirmed" and paid["actual_date"] == "2026-10-08"
    assert paid["planned_date"] == "2026-10-06"
    daily = {day["date"]: day for day in data["daily"]}
    assert Decimal(daily["2026-10-06"]["balance"]) == 1000
    assert Decimal(daily["2026-10-08"]["balance"]) == 900
    assert Decimal(daily["2026-10-08"]["expense"]) == 100
    assert Decimal(data["summary"][0]["planned_expense"]) == 0
    assert Decimal(calendar(crm, status="rejected")["summary"][0]["current_balance"]) == 900
    with crm["sessions"]() as db:
        assert db.query(AuditEvent).filter_by(entity_id=row["id"], action="confirmed").count() == 1
    assert decide(crm, row, status=409)["code"] in ("VERSION_CONFLICT", "REVIEW_STATE")


def test_rejection_requires_no_actual_date_and_never_changes_cash(crm):  # noqa: F811
    login(crm)
    balance(crm)
    row = planned(crm)
    invoice(crm, row)
    row = submit(crm, row)
    decide(crm, row, decision="rejected", actual_date=None)
    data = calendar(crm)
    assert data["items"][0]["status"] == "rejected"
    assert data["items"][0]["actual_date"] is None
    assert Decimal(data["summary"][0]["current_balance"]) == 1000
    assert Decimal(data["summary"][0]["planned_expense"]) == 0


def test_currencies_have_independent_balances_and_no_invented_opening_saldo(crm):  # noqa: F811
    login(crm)
    balance(crm)
    income = submit(crm, planned(crm, "income", "25", currency="USD"))
    decide(crm, income)
    by_currency = {row["currency"]: row for row in calendar(crm)["summary"]}
    assert Decimal(by_currency["RUB"]["current_balance"]) == 1000
    assert Decimal(by_currency["USD"]["income"]) == 25
    assert by_currency["USD"]["current_balance"] is None
    balance(crm, amount="100", currency="USD")
    usd = calendar(crm, currency="USD")["summary"]
    assert len(usd) == 1 and Decimal(usd[0]["current_balance"]) == 125
    assert Decimal(calendar(crm, currency="RUB")["summary"][0]["current_balance"]) == 1000


def test_payment_days_apply_to_expenses_with_explicit_exception_and_versioned_rules(crm):  # noqa: F811
    login(crm)
    assert crm["client"].get("/api/v1/payment-calendar/rules").json()["weekdays"] == [1, 3]
    error = cmd(crm, "/payment-calendar", {"direction": "expense", "planned_date": "2026-10-09",
                "amount": "100", "purpose": "Срочный платёж"}, status=422)
    assert error["field"] == "planned_date"
    row = planned(crm, planned_date="2026-10-09", outside_payment_days=True)
    invoice(crm, row)
    row = submit(crm, row)
    decide(crm, row, actual_date="2026-10-09", outside_payment_days=True)
    ordinary = planned(crm)
    invoice(crm, ordinary)
    ordinary = submit(crm, ordinary)
    assert decide(crm, ordinary, actual_date="2026-10-09", status=422)["field"] == "actual_date"
    rules = cmd(crm, "/payment-calendar/rules", {"version": 1, "weekdays": [4]}, method="put")
    assert rules == {"version": 2, "weekdays": [4]}
    decide(crm, ordinary, actual_date="2026-10-09")
    assert cmd(crm, "/payment-calendar/rules", {"version": 1, "weekdays": [1]}, method="put", status=409)["code"]
    next_rules = cmd(crm, "/payment-calendar/rules", {"version": 2, "weekdays": [1, 3]}, method="put")
    assert next_rules["version"] == 3
    with crm["sessions"]() as db:
        rows = db.scalars(select(AppSetting).where(AppSetting.key == "payment_days")).all()
        assert len(rows) == 2 and sum(r.status == "published" for r in rows) == 1
    assert planned(crm, "income", planned_date="2026-10-09")["status"] == "draft"


def test_actual_date_and_clean_invoice_are_required_for_approval(crm):  # noqa: F811
    login(crm)
    row = planned(crm)
    assert cmd(crm, f"/payment-calendar/{row['id']}/submit", {"version": row["version"]}, status=422)["code"] == "PAYMENT_INVOICE_REQUIRED"
    invoice(crm, row, status="quarantined")
    row = submit(crm, row)
    assert decide(crm, row, actual_date=None, status=422)["field"] == "actual_date"
    assert decide(crm, row, actual_date="2099-01-01", status=422)["code"] == "ACTUAL_DATE_FUTURE"
    assert decide(crm, row, status=422)["code"] == "PAYMENT_INVOICE_NOT_READY"
    with crm["sessions"].begin() as db:
        db.query(FileRecord).filter_by(calendar_entry_id=row["id"]).update({"status": "clean"})
    decide(crm, row)


def test_dated_balances_reset_at_start_of_day_without_recounting_old_payments(crm):  # noqa: F811
    login(crm)
    balance(crm)
    income = submit(crm, planned(crm, "income", "200", planned_date="2026-10-06"))
    decide(crm, income, actual_date="2026-10-06")
    expense = planned(crm, amount="50")
    invoice(crm, expense)
    expense = submit(crm, expense)
    decide(crm, expense, actual_date="2026-10-08")
    assert Decimal(calendar(crm)["summary"][0]["current_balance"]) == 1150
    anchor = balance(crm, "2026-10-08", "700")
    daily = {day["date"]: day for day in calendar(crm)["daily"]}
    assert Decimal(daily["2026-10-07"]["balance"]) == 1200
    assert Decimal(daily["2026-10-08"]["balance"]) == 650
    assert Decimal(calendar(crm, from_date="2026-10-08")["summary"][0]["opening_balance"]) == 700
    assert balance(crm, "2026-10-08", "800", status=409)["code"] == "VERSION_REQUIRED"
    balance(crm, "2026-10-08", "-100", version=anchor["version"])
    assert Decimal(calendar(crm)["summary"][0]["current_balance"]) == -150
    assert cmd(crm, "/payment-calendar/balances", {"balance_date": "2099-01-01", "amount": "1", "reason": "Будущее"}, method="put", status=422)["code"]


def test_invoice_upload_inherits_payment_access_and_employee_cannot_decide_or_read_company_balance(crm):  # noqa: F811
    with crm["sessions"].begin() as db:
        db.add(PermissionGrant(user_id=crm["manager"].id, code="files.upload", scope="all"))
    login(crm, "manager@example.com")
    row = planned(crm)
    import base64
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jY9sAAAAASUVORK5CYII=")
    uploaded = crm["client"].post("/api/v1/files", data={"entity_type": "calendar_entry", "entity_id": row["id"]},
                                  headers={"Idempotency-Key": str(uuid4())},
                                  files={"file": ("Счёт.png", png, "image/png")})
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["calendar_entry_id"] == row["id"]
    assert crm["client"].get(f"/api/v1/files?entity_type=calendar_entry&entity_id={row['id']}").json()["items"]
    row = submit(crm, row)
    assert decide(crm, row, status=403)["code"]
    assert balance(crm, status=403)["code"]
    assert crm["client"].get("/api/v1/payment-calendar/balances").status_code == 403
    login(crm)
    private = planned(crm, "income")
    invoice(crm, private)
    balance(crm)
    login(crm, "manager@example.com")
    assert crm["client"].get(f"/api/v1/files?entity_type=calendar_entry&entity_id={private['id']}").status_code == 404
    assert calendar(crm)["global_balance"] is False
    assert all(summary["current_balance"] is None for summary in calendar(crm)["summary"])


def test_daily_excel_uses_actual_or_planned_date_and_native_money_dates(crm):  # noqa: F811
    login(crm)
    party = post(crm, "/counterparties", {"name": "=HYPERLINK(\"example\")"})
    row = planned(crm, counterparty_id=party["id"])
    invoice(crm, row)
    row = submit(crm, row)
    decide(crm, row)
    response = crm["client"].get("/api/v1/payment-calendar/export?on_date=2026-10-08&date_basis=actual")
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.active
    assert [sheet.cell(1, c).value for c in range(1, 5)] == ["Контрагент", "Сумма", "Плановая дата", "Фактическая дата"]
    assert sheet.cell(2, 1).value.startswith("'=")
    assert sheet.cell(2, 2).value == 100 and sheet.cell(2, 2).data_type == "n"
    assert sheet.cell(2, 3).value.date() == date(2026, 10, 6)
    assert sheet.cell(2, 4).value.date() == date(2026, 10, 8)
    book.close()
    planned_file = crm["client"].get("/api/v1/payment-calendar/export?on_date=2026-10-06&date_basis=planned")
    assert load_workbook(io.BytesIO(planned_file.content)).active.max_row == 2
    assert crm["client"].get("/api/v1/payment-calendar?from_date=2026-10-08&to_date=2026-10-07").status_code == 422


def test_selecting_loss_reason_closes_quote_stage_and_reopening_clears_it(crm):  # noqa: F811
    login(crm, "manager@example.com")
    party = post(crm, "/counterparties", {"name": "Отказ клиента"})
    req = post(crm, "/requests", {"title": "Запрос с отказом", "client_id": party["id"]})
    with crm["sessions"].begin() as db:
        db.get(Request, req["id"]).commercial_stage = "quote_given"
    saved = crm["client"].patch(f"/api/v1/requests/{req['id']}", json={"version": req["version"],
        "commercial_stage": "quote_given", "loss_reason": "no_budget", "reason": "Отказ от заказа"})
    assert saved.status_code == 200, saved.text
    closed = saved.json()
    assert closed["commercial_stage"] == "closed_lost" and closed["closed_at"]
    listed = crm["client"].get("/api/v1/requests?stage=closed_lost").json()["items"]
    assert listed[0]["id"] == req["id"]
    reopened = crm["client"].patch(f"/api/v1/requests/{req['id']}", json={"version": closed["version"],
        "commercial_stage": "clarification", "loss_reason": "no_budget", "reason": "Клиент возобновил закупку"})
    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["loss_reason"] is None and reopened.json()["closed_at"] is None

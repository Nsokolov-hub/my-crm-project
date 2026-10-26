"""Site requests in Excel, the report by channels and joining a site request to an existing CRM request
(site TZ 9.4.1, 15.3).

- One site request or the list with its filters as an xlsx workbook (step 3 «Просмотр и выгрузка»).
- The report by channels: site requests → qualified leads → proposals sent → confirmed sales → revenue, by the
  channel of the first visit, for a period (the end-to-end funnel of the site's KPI).
- «Добавить в существующую заявку»: several site requests of one client become one CRM request (P1).
"""

import io
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Header, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import Field
from sqlalchemy import func, or_, select

from app.core.config import settings
from app.core.db import DB, utcnow
from app.core.errors import DomainError
from app.core.models import User
from app.core.security import Actor, check_request, require_permission
from app.core.service import audit, check_version, idem, notify, prefetch
from app.crm.common import save
from app.crm.models import Request, RequestItem, RequestItemRevision, Task
from app.crm.requests import structured_item_fields
from app.integrations.models import SiteRequest, SiteRequestAttribution, SiteRequestItem
from app.integrations.site_contract import ConvertItemInput, Input
from app.integrations.site_conversion import ANALOGUE_LABELS, SOURCE_COLUMNS, copy_files, desired_moment
from app.integrations.site_requests import OPEN_TASK_STATUSES, check_site_request, detail_view, sees_all

router = APIRouter(tags=['Заявки с сайта'])
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
KIND_NAMES = {'product_rfq': 'Корзина', 'one_click': 'В 1 клик', 'service': 'Услуга', 'general': 'Общая форма'}
STATUS_NAMES = {'new': 'Новая', 'in_work': 'В работе', 'converted': 'Преобразована', 'rejected': 'Отклонена'}
QUALIFIED = ('collecting_quotes', 'quote_given', 'calculation', 'proposal_sent', 'composition_agreed',
             'awaiting_payment', 'sale_confirmed')
PROPOSED = ('proposal_sent', 'composition_agreed', 'awaiting_payment', 'sale_confirmed')
CHANNELS = ('cpc', 'organic', 'referral', 'direct', 'email', 'social')
CHANNEL_NAMES = {'cpc': 'Реклама', 'organic': 'Поиск', 'referral': 'Переходы с сайтов', 'direct': 'Прямые заходы',
                 'email': 'Рассылки', 'social': 'Соцсети', None: 'Не определён'}


def workbook_response(sheets: list[tuple[str, list[str], list[list[Any]]]], name: str) -> Response:
    book = Workbook()
    book.remove(book.active)
    for title, head, rows in sheets:
        sheet = book.create_sheet(title[:31])
        sheet.append(head)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for index, column in enumerate(head, start=1):
            width = max([len(str(column))] + [len(str(r[index - 1] or '')) for r in rows[:200]])
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(60, max(10, width + 2))
    out = io.BytesIO()
    book.save(out)
    safe = re.sub(r'[^A-Za-z0-9_.-]', '_', name)
    return Response(out.getvalue(), media_type=XLSX, headers={'Content-Disposition': f'attachment; filename="{safe}"'})


def local(value: datetime | None) -> str:
    return value.astimezone(ZoneInfo(settings.company_timezone)).strftime('%d.%m.%Y %H:%M') if value else ''


@router.get('/site-requests/export.xlsx')
def export_list(user: Actor, db: DB, status: str = '', q: str = '') -> Response:
    """Список заявок с сайта с текущими фильтрами в Excel."""
    require_permission(db, user, 'site_requests.read')
    require_permission(db, user, 'exports.download')
    statement = select(SiteRequest)
    if not sees_all(db, user):
        statement = statement.where(SiteRequest.assignee_id == user.id)
    if status:
        statement = statement.where(SiteRequest.status == status)
    if q.strip():
        like = f'%{q.strip()}%'
        statement = statement.where(or_(SiteRequest.number.ilike(like), SiteRequest.company.ilike(like),
                                        SiteRequest.email.ilike(like), SiteRequest.phone.ilike(like),
                                        SiteRequest.inn == q.strip()))
    rows = list(db.scalars(statement.order_by(SiteRequest.submitted_at.desc()).limit(10000)))
    ids = [r.id for r in rows]
    counts: dict[str, int] = dict(db.execute(select(SiteRequestItem.site_request_id, func.count())
                             .where(SiteRequestItem.site_request_id.in_(ids))
                             .group_by(SiteRequestItem.site_request_id)).tuples().all()) if ids else {}
    channels: dict[str, str | None] = dict(db.execute(select(SiteRequestAttribution.site_request_id, SiteRequestAttribution.channel)
                               .where(SiteRequestAttribution.site_request_id.in_(ids))).tuples().all()) if ids else {}
    users = prefetch(db, User, (r.assignee_id for r in rows))
    requests = prefetch(db, Request, (r.request_id for r in rows))
    data = [[r.number, local(r.submitted_at), KIND_NAMES.get(r.kind, r.kind), STATUS_NAMES.get(r.status, r.status),
             r.company, r.inn or '', r.customer_name, r.phone, r.email, r.city or '', counts.get(r.id, 0),
             CHANNEL_NAMES.get(channels.get(r.id), channels.get(r.id)) or '',
             users[r.assignee_id].name if r.assignee_id in users else '',
             requests[r.request_id].number if r.request_id in requests else '', r.reject_reason or '']
            for r in rows]
    audit(db, user, 'site_request', 'export', 'exported', after={'count': len(rows), 'status': status, 'q': q})
    return workbook_response([('Заявки с сайта', ['Номер', 'Получена', 'Тип', 'Статус', 'Компания', 'ИНН', 'Контакт',
                                                  'Телефон', 'E-mail', 'Город', 'Позиций', 'Канал', 'Ответственный',
                                                  'Заявка CRM', 'Причина отказа'], data)],
                             f'site-requests-{date.today().isoformat()}.xlsx')


@router.get('/site-requests/{entity_id}/export.xlsx')
def export_one(entity_id: str, user: Actor, db: DB) -> Response:
    """Одна заявка с сайта в Excel: контакты, позиции, источник."""
    require_permission(db, user, 'exports.download')
    row = check_site_request(db, user, entity_id)
    items = list(db.scalars(select(SiteRequestItem).where(SiteRequestItem.site_request_id == row.id)
                            .order_by(SiteRequestItem.position)))
    attribution = db.scalar(select(SiteRequestAttribution).where(SiteRequestAttribution.site_request_id == row.id))
    card = [['Номер', row.number], ['Номер на сайте', row.site_number], ['Получена', local(row.submitted_at)],
            ['Тип', KIND_NAMES.get(row.kind, row.kind)], ['Статус', STATUS_NAMES.get(row.status, row.status)],
            ['Компания', row.company], ['ИНН', row.inn or ''], ['Контакт', row.customer_name], ['Телефон', row.phone],
            ['E-mail', row.email], ['Город', row.city or ''], ['Комментарий', row.comment or ''],
            ['Желаемый срок', row.desired_date.strftime('%d.%m.%Y') if row.desired_date else ''],
            ['Страница', row.page_url or ''],
            ['Канал', CHANNEL_NAMES.get(attribution.channel if attribution else None, '')],
            ['Источник', ' / '.join(v for v in (attribution.last_utm_source, attribution.last_utm_medium,
                                                 attribution.last_utm_campaign) if v) if attribution else '']]
    lines = [[i.position, i.name, i.brand or '', i.catalog_number or '', i.cas or '', i.packaging or '',
              float(i.quantity), i.unit, ANALOGUE_LABELS.get(i.analogue, i.analogue),
              'да' if i.custom_packaging else '', i.comment or ''] for i in items]
    audit(db, user, 'site_request', row.id, 'exported')
    return workbook_response([('Заявка', ['Поле', 'Значение'], card),
                              ('Позиции', ['№', 'Наименование', 'Бренд', 'Кат. №', 'CAS', 'Фасовка', 'Количество',
                                           'Единица', 'Аналоги', 'Своя фасовка', 'Комментарий'], lines)],
                             f'{row.number}.xlsx')


@router.get('/site-requests/report')
def channel_report(user: Actor, db: DB, date_from: date | None = None, date_to: date | None = None) -> dict[str, Any]:
    """Сквозной отчёт по каналам (ТЗ сайта 15.3): заявки с сайта → квалифицированные → КП → продажи → выручка."""
    require_permission(db, user, 'site_requests.read')
    if not sees_all(db, user):
        raise DomainError('FORBIDDEN', 'Отчёт по каналам доступен сотрудникам с доступом ко всем заявкам', 403)
    zone = ZoneInfo(settings.company_timezone)
    end = date_to or date.today()
    start = date_from or end - timedelta(days=30)
    since = datetime.combine(start, time.min, zone)
    until = datetime.combine(end + timedelta(days=1), time.min, zone)
    rows = db.execute(
        select(SiteRequest, SiteRequestAttribution.channel, Request)
        .outerjoin(SiteRequestAttribution, SiteRequestAttribution.site_request_id == SiteRequest.id)
        .outerjoin(Request, Request.id == SiteRequest.request_id)
        .where(SiteRequest.submitted_at >= since, SiteRequest.submitted_at < until)).all()
    from app.integrations.site_events import revenue

    table: dict[str | None, dict[str, Any]] = {}
    counted_requests: set[str] = set()
    for site, channel, request in rows:
        line = table.setdefault(channel, {'channel': channel, 'name': CHANNEL_NAMES.get(channel, channel),
                                          'leads': 0, 'rejected': 0, 'converted': 0, 'qualified': 0, 'proposals': 0,
                                          'sales': 0, 'revenue': Decimal(0)})
        line['leads'] += 1
        line['rejected'] += site.status == 'rejected'
        line['converted'] += site.status == 'converted'
        if request is None or request.id in counted_requests:
            continue
        # Several site requests joined into one CRM request count it once.
        counted_requests.add(request.id)
        sold = request.sale_confirmed_at is not None
        line['qualified'] += request.commercial_stage in QUALIFIED or sold
        line['proposals'] += request.commercial_stage in PROPOSED or sold
        if sold:
            line['sales'] += 1
            amount = revenue(db, request.id)
            if amount and amount['currency'] in ('RUB', 'RUR'):
                line['revenue'] += Decimal(amount['amount'])
    order = {c: i for i, c in enumerate(CHANNELS)}
    lines = sorted(table.values(), key=lambda line: order.get(line['channel'], 99))
    totals = {key: sum(line[key] for line in lines) for key in
              ('leads', 'rejected', 'converted', 'qualified', 'proposals', 'sales')}
    totals['revenue'] = str(sum((line['revenue'] for line in lines), Decimal(0)))
    for line in lines:
        line['revenue'] = str(line['revenue'])
    return {'date_from': start.isoformat(), 'date_to': end.isoformat(), 'channels': lines, 'totals': totals,
            'currency': 'RUB'}


class AttachInput(Input):
    version: int = Field(ge=1)
    request_id: str = Field(max_length=36)
    items: list[ConvertItemInput] = Field(default_factory=list, max_length=500)


@router.post('/site-requests/{entity_id}/attach')
def attach(entity_id: str, body: AttachInput, user: Actor, db: DB,
           idempotency_key: str | None = Header(default=None)) -> dict[str, Any]:
    """Добавить заявку с сайта в существующую заявку CRM того же клиента (ТЗ сайта 9.4.1, P1): позиции и вложения
    переходят в неё, заявка с сайта — «Преобразована» со ссылкой на неё."""
    require_permission(db, user, 'site_requests.write')

    def operation() -> dict[str, Any]:
        row = check_site_request(db, user, entity_id, for_update=True)
        check_version(row, body.version)
        if row.status in ('converted', 'rejected'):
            raise DomainError('SITE_REQUEST_CLOSED', 'Заявка с сайта уже обработана', 409)
        request: Request = check_request(db, user, body.request_id, 'requests.write')
        if request.closed_at is not None or request.archived:
            raise DomainError('REQUEST_CLOSED', 'Заявка CRM закрыта или в архиве', 409, 'request_id')
        site_items = {item.id: item for item in db.scalars(
            select(SiteRequestItem).where(SiteRequestItem.site_request_id == row.id))}
        desired = desired_moment(row)
        for line in body.items:
            site_item = site_items.get(line.site_item_id or '')
            data = structured_item_fields(db, {**line.model_dump(exclude={'site_item_id'}), 'desired_at': desired})
            item = RequestItem(request_id=request.id, **data, source_columns=SOURCE_COLUMNS if site_item else [],
                               source_values=[site_item.brand or '', ANALOGUE_LABELS[site_item.analogue]]
                               if site_item else [])
            snapshot = save(db, user, item, 'request_item')
            db.add(RequestItemRevision(item_id=item.id, revision=1, snapshot=snapshot, author_id=user.id))
        files = copy_files(db, user, row, request)
        now = utcnow()
        before = {'status': row.status}
        row.status, row.request_id, row.converted_at = 'converted', request.id, now
        if not row.assignee_id:
            row.assignee_id, row.assigned_at = request.owner_id, now
        row.version += 1
        request.version += 1
        for task in db.scalars(select(Task).where(Task.entity_type == 'site_request', Task.entity_id == row.id,
                                                  Task.status.in_(OPEN_TASK_STATUSES))):
            task.status, task.result, task.completed_at = 'completed', f'Добавлена в заявку {request.number}', now
            task.version += 1
        audit(db, user, 'site_request', row.id, 'attached', before,
              {'status': row.status, 'request_number': request.number, 'items': len(body.items), 'files': files})
        if request.owner_id != user.id:
            notify(db, request.owner_id, f'site-request-attached:{row.id}',
                   f'В заявку {request.number} добавлена заявка с сайта {row.number}', 'request', request.id)
        db.flush()
        return {**detail_view(db, user, row), 'created': {'request_id': request.id, 'request_number': request.number}}

    return idem(db, user, idempotency_key, f'site-request.attach:{entity_id}', body.model_dump(mode='json'),
                operation)

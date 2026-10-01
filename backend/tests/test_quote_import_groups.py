"""Quote imports resolve existing product groups without changing existing catalogue entries."""

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from test_business_imports import confirm, preview, quote_file, request
from test_crm import crm as crm
from test_crm import login, post

from app.crm.imports import xlsx
from app.crm.models import Nomenclature, ProductGroup, QuoteItem, QuoteSheet, RequestItem, TableImport
from app.crm.table_imports import QUOTE_COLUMNS


def group(crm, name="Стандарты", slug="standards"):
    return post(crm, "/product-groups", {"name": name, "slug": slug})


def demand(crm, count=1):
    login(crm)
    req = request(crm)
    supplier = post(crm, "/counterparties", {"name": "Поставщик", "kind": "supplier"})
    raw = preview(
        crm,
        req["id"],
        ["Name", "Article", "Quantity"],
        [[f"Потребность {i + 1}", "0341", "1"] for i in range(count)],
    )
    confirm(crm, req["id"], raw["id"])
    return req, supplier


def quote_preview(crm, req, supplier, group_value="", overrides=None):
    return preview(
        crm,
        req["id"],
        list(QUOTE_COLUMNS.values()),
        [quote_file(supplier["internal_code"], {"product_group": group_value, **(overrides or {})})],
        "quotes",
    )


def assert_no_quote_writes(crm, catalogue_count=0):
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(Nomenclature)) == catalogue_count
        assert db.scalar(select(func.count()).select_from(QuoteItem)) == 0


@pytest.mark.parametrize("reference", ["  ФАРМАЦЕВТИЧЕСКИЕ   СТАНДАРТЫ  ", "  STANDARDS  "])
def test_quote_import_resolves_group_by_normalized_name_or_slug(crm, reference):
    req, supplier = demand(crm)
    selected = group(crm, "Фармацевтические стандарты")
    batch = quote_preview(crm, req, supplier, reference)
    assert batch["summary"]["errors"] == 0
    assert batch["rows"][0]["data"]["product_group_name"] == selected["name"]
    assert_no_quote_writes(crm)
    result = confirm(crm, req["id"], batch["id"])["result"]
    assert result["imported"] == result["nomenclatures_created"] == 1
    with crm["sessions"]() as db:
        assert db.scalar(select(Nomenclature)).product_group_id == selected["id"]
        assert db.scalar(select(RequestItem)).product_group_id == selected["id"]
        assert db.scalar(select(func.count()).select_from(ProductGroup)) == 2


@pytest.mark.parametrize("reference", ["", "0"])
@pytest.mark.parametrize("source_has_group", [False, True])
def test_blank_quote_group_uses_source_group_or_other(crm, reference, source_has_group):
    req, supplier = demand(crm)
    selected = group(crm)
    with crm["sessions"].begin() as db:
        if source_has_group:
            db.scalar(select(RequestItem)).product_group_id = selected["id"]
        expected = (
            selected["id"]
            if source_has_group
            else db.scalar(select(ProductGroup.id).where(ProductGroup.slug == "other"))
        )
    batch = quote_preview(crm, req, supplier, reference)
    assert batch["summary"]["errors"] == 0
    confirm(crm, req["id"], batch["id"])
    with crm["sessions"]() as db:
        assert db.scalar(select(Nomenclature)).product_group_id == expected
        assert db.scalar(select(RequestItem)).product_group_id == expected


def test_old_quote_file_without_group_column_still_imports(crm):
    req, supplier = demand(crm)
    selected = group(crm)
    with crm["sessions"].begin() as db:
        db.scalar(select(RequestItem)).product_group_id = selected["id"]
    columns = list(QUOTE_COLUMNS)
    group_index = columns.index("product_group")
    headers = list(QUOTE_COLUMNS.values())
    values = quote_file(supplier["internal_code"])
    headers.pop(group_index)
    values.pop(group_index)
    batch = preview(crm, req["id"], headers, [values], "quotes")
    assert batch["summary"]["errors"] == 0
    confirm(crm, req["id"], batch["id"])
    with crm["sessions"]() as db:
        assert db.scalar(select(Nomenclature)).product_group_id == selected["id"]


def test_default_other_group_uses_its_slug_even_if_another_group_is_named_other(crm):
    req, supplier = demand(crm)
    group(crm, "other", "special")
    with crm["sessions"]() as db:
        expected = db.scalar(select(ProductGroup.id).where(ProductGroup.slug == "other"))
    batch = quote_preview(crm, req, supplier)
    assert batch["summary"]["errors"] == 0
    confirm(crm, req["id"], batch["id"])
    with crm["sessions"]() as db:
        assert db.scalar(select(Nomenclature)).product_group_id == expected


def test_reimport_of_completed_legacy_quote_file_does_not_duplicate_quotes(crm):
    req, supplier = demand(crm)
    group_index = list(QUOTE_COLUMNS).index("product_group")
    headers = list(QUOTE_COLUMNS.values())
    values = quote_file(supplier["internal_code"])
    headers.pop(group_index)
    values.pop(group_index)
    content = xlsx([headers, values])

    def upload_same_file():
        response = crm["client"].post(
            f"/api/v1/requests/{req['id']}/table-imports/preview",
            files={
                "file": (
                    "legacy-quotes.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"kind": "quotes"},
        )
        assert response.status_code == 200, response.text
        return response.json()

    first = upload_same_file()
    imported = confirm(crm, req["id"], first["id"])["result"]
    assert imported["imported"] == 1
    with crm["sessions"].begin() as db:
        legacy_batch = db.get(TableImport, first["id"])
        # Before the group column existed, completed imports did not have this optional mapping key.
        legacy_batch.mapping = {
            key: value for key, value in legacy_batch.mapping.items() if key != "product_group"
        }
        sheet_ids = set(db.scalars(select(QuoteSheet.id)))
        quote_ids = set(db.scalars(select(QuoteItem.id)))
    again = upload_same_file()
    result = confirm(crm, req["id"], again["id"])["result"]
    assert result["already_imported"] is True
    assert result["quote_sheet_ids"] == imported["quote_sheet_ids"]
    with crm["sessions"]() as db:
        assert set(db.scalars(select(QuoteSheet.id))) == sheet_ids
        assert set(db.scalars(select(QuoteItem.id))) == quote_ids


@pytest.mark.parametrize("invalid", ["unknown", "inactive", "ambiguous"])
def test_invalid_explicit_quote_group_blocks_whole_import_without_creating_groups(crm, invalid):
    req, supplier = demand(crm)
    if invalid == "unknown":
        reference = "Несуществующая группа"
    elif invalid == "inactive":
        selected = group(crm)
        with crm["sessions"].begin() as db:
            db.get(ProductGroup, selected["id"]).active = False
        reference = selected["name"]
    else:
        group(crm, "Стандарты", "standards")
        group(crm, "СТАНДАРТЫ", "standards_upper")
        reference = "стандарты"
    with crm["sessions"]() as db:
        initial_groups = db.scalar(select(func.count()).select_from(ProductGroup))
    batch = quote_preview(crm, req, supplier, reference)
    assert batch["summary"]["errors"] == 1
    assert batch["rows"][0]["errors"]
    assert confirm(crm, req["id"], batch["id"], 422)["code"] == "IMPORT_ERRORS"
    assert_no_quote_writes(crm)
    with crm["sessions"]() as db:
        assert db.scalar(select(func.count()).select_from(ProductGroup)) == initial_groups


def test_existing_nomenclature_keeps_group_when_quote_specifies_different_valid_group(crm):
    req, supplier = demand(crm)
    original = group(crm)
    different = group(crm, "Колонки", "columns")
    known = post(
        crm,
        "/nomenclatures",
        {
            "name": "Известный товар",
            "article": "0341",
            "product_group_id": original["id"],
            "packings": [{"value": "1", "unit": "pcs"}],
        },
    )
    batch = quote_preview(crm, req, supplier, different["name"])
    assert batch["summary"] == {"total": 1, "checked": 1, "create_nomenclature": 0, "errors": 0}
    assert batch["rows"][0]["data"]["product_group_name"] == original["name"]
    confirm(crm, req["id"], batch["id"])
    with crm["sessions"]() as db:
        assert db.get(Nomenclature, known["id"]).product_group_id == original["id"]
        assert db.scalar(select(RequestItem)).product_group_id == original["id"]
        assert db.scalar(select(func.count()).select_from(Nomenclature)) == 1


def test_existing_nomenclature_does_not_hide_invalid_explicit_group(crm):
    req, supplier = demand(crm)
    post(
        crm,
        "/nomenclatures",
        {"name": "Известный товар", "article": "0341", "packings": [{"value": "1", "unit": "pcs"}]},
    )
    batch = quote_preview(crm, req, supplier, "Опечатка в группе")
    assert batch["summary"]["errors"] == 1
    assert confirm(crm, req["id"], batch["id"], 422)["code"] == "IMPORT_ERRORS"
    assert_no_quote_writes(crm, catalogue_count=1)


def test_one_new_article_cannot_create_nomenclature_in_different_groups(crm):
    req, supplier = demand(crm, count=2)
    first = group(crm)
    second = group(crm, "Колонки", "columns")
    batch = preview(
        crm,
        req["id"],
        list(QUOTE_COLUMNS.values()),
        [
            quote_file(supplier["internal_code"], {"source_row": str(i), "product_group": choice["name"]})
            for i, choice in enumerate([first, second], 1)
        ],
        "quotes",
    )
    assert batch["summary"]["errors"] >= 1
    assert confirm(crm, req["id"], batch["id"], 422)["code"] == "IMPORT_ERRORS"
    assert_no_quote_writes(crm)


@pytest.mark.parametrize("change", ["name", "slug", "active"])
def test_product_group_changes_after_preview_require_new_preview(crm, change):
    req, supplier = demand(crm)
    selected = group(crm)
    batch = quote_preview(crm, req, supplier, selected["name"])
    assert batch["summary"]["errors"] == 0
    with crm["sessions"].begin() as db:
        row = db.get(ProductGroup, selected["id"])
        setattr(row, change, {"name": "Обновленная группа", "slug": "updated", "active": False}[change])
        row.version += 1
    assert confirm(crm, req["id"], batch["id"], 409)["code"] == "IMPORT_STALE"
    assert_no_quote_writes(crm)


def test_quote_template_has_group_column_directory_dropdown_and_prefilled_groups(crm):
    req, supplier = demand(crm)
    source_group = group(crm, "Колонки", "columns")
    known_group = group(crm)
    inactive = group(crm, "Архивные", "archived")
    known = post(
        crm,
        "/nomenclatures",
        {
            "name": "Известный товар",
            "article": "known-01",
            "product_group_id": known_group["id"],
            "packings": [{"value": "1", "unit": "pcs"}],
        },
    )
    post(
        crm,
        f"/requests/{req['id']}/items",
        {"nomenclature_id": known["id"], "packing_id": known["packings"][0]["id"], "quantity": "1"},
    )
    with crm["sessions"].begin() as db:
        raw = db.scalar(select(RequestItem).where(RequestItem.nomenclature_id.is_(None)))
        raw.product_group_id = source_group["id"]
        db.get(ProductGroup, inactive["id"]).active = False
    response = crm["client"].get(f"/api/v1/requests/{req['id']}/quote-import/template.xlsx")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.active is book.worksheets[0]
    sheet = book.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(QUOTE_COLUMNS.values())
    assert headers[-1] == "Товарная группа"
    group_column = headers.index("Товарная группа") + 1
    assert {sheet.cell(row, group_column).value for row in (2, 3)} == {
        source_group["name"],
        known_group["name"],
    }
    directory = book["Товарные группы"]
    directory_values = {str(cell.value) for row in directory for cell in row if cell.value is not None}
    assert {source_group["name"], known_group["name"], "Прочее"} <= directory_values
    assert inactive["name"] not in directory_values
    group_cells = [sheet.cell(row, group_column).coordinate for row in (2, 3)]
    assert any(
        validation.type == "list"
        and validation.formula1
        and all(coordinate in validation.sqref for coordinate in group_cells)
        for validation in sheet.data_validations.dataValidation
    )
    book.close()


def test_quote_template_import_uses_quotes_sheet_when_directory_was_last_open(crm):
    req, supplier = demand(crm)
    selected = group(crm)
    response = crm["client"].get(f"/api/v1/requests/{req['id']}/quote-import/template.xlsx")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    quote_sheet = book["Квоты"]
    values = quote_file(supplier["internal_code"], {"product_group": selected["name"]})
    for column, value in enumerate(values, 1):
        quote_sheet.cell(2, column).value = value
    # A user consults the lookup and saves Excel without returning to the quote sheet.
    book.active = book.sheetnames.index("Товарные группы")
    content = io.BytesIO()
    book.save(content)
    book.close()
    response = crm["client"].post(
        f"/api/v1/requests/{req['id']}/table-imports/preview",
        files={
            "file": (
                "filled-template.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"kind": "quotes"},
    )
    assert response.status_code == 200, response.text
    batch = response.json()
    assert batch["summary"] == {"total": 1, "checked": 0, "create_nomenclature": 1, "errors": 0}
    assert batch["rows"][0]["data"]["product_group_name"] == selected["name"]
    assert confirm(crm, req["id"], batch["id"])["result"]["imported"] == 1
    with crm["sessions"]() as db:
        assert db.scalar(select(Nomenclature)).product_group_id == selected["id"]

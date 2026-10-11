"""Payment calendar: planned and actual payments, payment days, balances and export."""

from datetime import date, timedelta
from decimal import Decimal
from typing import Literal

from fastapi import APIRouter, Query
from fastapi.responses import Response
from pydantic import Field
from sqlalchemy import select

from app.business.absences import ensure_available
from app.business.common import calendar_visible
from app.business.reviews import submit_review
from app.commerce.files import workbook
from app.commerce.models import CommercialDocument
from app.commerce.procurement import DB, Actor
from app.commerce.schemas import Command, Currency, Positive, VersionCommand
from app.core.errors import error
from app.core.models import AppSetting
from app.core.security import check_client, check_request, require_permission
from app.core.service import advisory, audit, check_version, idem, lock, plain, prefetch, serialize
from app.crm.models import Counterparty

from . import cashflow
from .models import CalendarBalance, CalendarEntry, SupplierOrder, WorkflowReview

router = APIRouter(tags=["Платёжный календарь и заказы поставщикам"])


def calendar_views(db, rows):
    """Calendar entries with counterparty names and their latest review, in three queries per page."""
    parties = prefetch(db, Counterparty, (row.counterparty_id for row in rows))
    latest = {}
    for review in db.scalars(select(WorkflowReview).where(
            WorkflowReview.kind == "calendar", WorkflowReview.entity_id.in_([row.id for row in rows]))
            .order_by(WorkflowReview.created_at.desc(), WorkflowReview.id.desc())):
        latest.setdefault(review.entity_id, review)
    views = []
    for row in rows:
        result = serialize(row)
        party = parties.get(row.counterparty_id) if row.counterparty_id else None
        result["counterparty_name"] = party.name if party else None
        review = latest.get(row.id)
        result["review_id"] = review.id if review and review.status == "pending" else None
        result["review_version"] = review.version if review and review.status == "pending" else None
        result["decision_reason"] = review.reason if review else None
        views.append(result)
    return views


def calendar_view(db, row):
    return calendar_views(db, [row])[0]


class CalendarIn(Command):
    direction: Literal["income", "expense"]
    planned_date: date
    amount: Positive
    currency: Currency = "RUB"
    purpose: str = Field(min_length=1, max_length=4000)
    counterparty_id: str | None = None
    request_id: str | None = None
    document_id: str | None = None
    supplier_order_id: str | None = None
    responsible_id: str | None = None
    payment_kind: Literal["prepayment", "deferred", "other"] = "other"
    recurrence: Literal["none", "weekly", "monthly"] = "none"
    outside_payment_days: bool = False


@router.get("/payment-calendar")
def calendar_list(
    db: DB, user: Actor, from_date: date | None = None, to_date: date | None = None,
    direction: str | None = None, status: str | None = None, counterparty_id: str | None = None,
    currency: str | None = None, payment_kind: str | None = None,
    date_basis: Literal["effective", "planned", "actual"] = "effective",
    page_number: int = Query(1, alias="page", ge=1), page_size: int = Query(50, ge=1, le=100),
):
    all_rows = db.scalars(cashflow.visible_query(db, user).order_by(CalendarEntry.planned_date, CalendarEntry.id)).all()
    dates = [row.actual_date or row.planned_date for row in all_rows]
    today = cashflow.company_today()
    end = to_date or max([today, *dates])
    start = from_date or max(min([today.replace(day=1), *dates]), end - timedelta(days=365))
    if end < start or (end - start).days > 365:
        error("CALENDAR_PERIOD", "Выберите период от 1 до 366 дней", field="to_date")
    rows = []
    for row in all_rows:
        day = row.planned_date if date_basis == "planned" else row.actual_date if date_basis == "actual" else (
            row.actual_date if row.status == "confirmed" else row.planned_date)
        if not day or not start <= day <= end:
            continue
        if any(value and getattr(row, key) != value for key, value in (
            ("direction", direction), ("status", status), ("counterparty_id", counterparty_id),
            ("currency", currency), ("payment_kind", payment_kind))):
            continue
        rows.append(row)
    totals, daily = cashflow.projection(db, user, all_rows, start, end, currency)
    paid = [row for row in all_rows if row.status == "confirmed" and row.actual_date and start <= row.actual_date <= end
            and row.counterparty_id and not (currency and currency != row.currency)]
    names = {key: party.name for key, party in prefetch(db, Counterparty, (row.counterparty_id for row in paid)).items()}
    parties = {}
    for row in paid:
        key = (row.counterparty_id, row.currency, row.direction)
        parties.setdefault(key, {"counterparty_name": names.get(row.counterparty_id),
                                 "currency": row.currency, "direction": row.direction, "amount": Decimal(0)})["amount"] += row.amount
    for row in parties.values():
        denominator = next(item[row["direction"]] for item in totals if item["currency"] == row["currency"])
        row["share_percent"] = row["amount"] / denominator * 100 if denominator else Decimal(0)
    return {"items": calendar_views(db, rows[(page_number-1)*page_size:page_number*page_size]),
            "total": len(rows), "page": page_number, "page_size": page_size,
            "summary": plain(totals), "daily": plain(daily), "parties": plain(list(parties.values())),
            "from_date": start.isoformat(), "to_date": end.isoformat(),
            "global_balance": cashflow.global_balance_allowed(db, user), "payment_days": cashflow.policy(db)}


@router.get("/payment-calendar/rules")
def calendar_rules(db: DB, user: Actor):
    return cashflow.policy(db)


class PaymentDaysIn(VersionCommand):
    weekdays: list[Literal[0, 1, 2, 3, 4, 5, 6]] = Field(min_length=1, max_length=7)


@router.put("/payment-calendar/rules")
def calendar_rules_save(data: PaymentDaysIn, db: DB, user: Actor):
    require_permission(db, user, "approvals.decide")
    def operation():
        advisory(db, "setting:payment_days")
        row = cashflow.policy_row(db)
        if row:
            check_version(row, data.version)
        elif data.version != 1:
            error("VERSION_CONFLICT", "Платёжные дни изменены. Обновите страницу", 409)
        if len(set(data.weekdays)) != len(data.weekdays):
            error("PAYMENT_DAYS_INVALID", "Выберите каждый день один раз", field="weekdays")
        before = cashflow.policy(db)
        previous = row
        if previous:
            previous.status = "archived"
            previous.effective_until = cashflow.company_today()
        row = AppSetting(key="payment_days", value={"weekdays": sorted(data.weekdays)}, status="published",
                         version=data.version + 1, author_id=user.id,
                         previous_id=previous.id if previous else None, effective_from=cashflow.company_today())
        db.add(row)
        db.flush()
        audit(db, user, "setting", row.id, "payment_days", before=before, after=cashflow.policy(db))
        return cashflow.policy(db)
    return idem(db, user, data.idempotency_key, "calendar.rules", data.model_dump(mode="json"), operation)


class CalendarBalanceIn(Command):
    balance_date: date
    currency: Currency = "RUB"
    amount: Decimal = Field(max_digits=24, decimal_places=8)
    reason: str = Field(min_length=1, max_length=4000)
    version: int | None = Field(default=None, ge=1)


def require_global_calendar(db, user):
    require_permission(db, user, "approvals.decide")
    if not cashflow.global_balance_allowed(db, user):
        error("BALANCE_SCOPE_REQUIRED", "Сальдо организации доступно руководителю с доступом ко всем заявкам", 403)


@router.get("/payment-calendar/balances")
def calendar_balances(db: DB, user: Actor):
    require_global_calendar(db, user)
    return {"items": [serialize(row) for row in db.scalars(select(CalendarBalance).order_by(CalendarBalance.balance_date.desc()))]}


@router.put("/payment-calendar/balances")
def calendar_balance_save(data: CalendarBalanceIn, db: DB, user: Actor):
    require_global_calendar(db, user)
    if data.balance_date > cashflow.company_today():
        error("BALANCE_DATE_FUTURE", "Сальдо фиксируется на текущую или прошедшую дату", field="balance_date")
    def operation():
        advisory(db, f"calendar-balance:{data.currency}:{data.balance_date}")
        row = db.scalar(select(CalendarBalance).where(CalendarBalance.currency == data.currency,
                        CalendarBalance.balance_date == data.balance_date).with_for_update())
        before = serialize(row) if row else None
        if row:
            if data.version is None:
                error("VERSION_REQUIRED", "Сальдо на эту дату уже существует. Откройте его для изменения", 409, field="version")
            check_version(row, data.version)
            row.amount, row.reason, row.author_id = data.amount, data.reason, user.id
            row.version += 1
        else:
            row = CalendarBalance(**data.model_dump(exclude={"idempotency_key", "version"}), author_id=user.id)
            db.add(row)
        db.flush()
        audit(db, user, "calendar_balance", row.id, "saved", before=before, after=serialize(row), reason=data.reason)
        return serialize(row)
    return idem(db, user, data.idempotency_key, "calendar.balance", data.model_dump(mode="json"), operation)


@router.get("/payment-calendar/export")
def calendar_export(on_date: date, db: DB, user: Actor, date_basis: Literal["planned", "actual", "effective"] = "effective"):
    require_permission(db, user, "exports.download")
    rows = db.scalars(cashflow.visible_query(db, user).order_by(CalendarEntry.planned_date, CalendarEntry.id)).all()
    selected = []
    for row in rows:
        day = row.planned_date if date_basis == "planned" else row.actual_date if date_basis == "actual" else (
            row.actual_date if row.status == "confirmed" else row.planned_date)
        if day == on_date:
            selected.append(row)
    statuses = {"draft": "Черновик", "pending": "На согласовании", "approved": "Согласовано",
                "confirmed": "Подтверждено", "rejected": "Отклонено", "cancelled": "Отменено"}
    views = calendar_views(db, selected)
    content = workbook(["Контрагент", "Сумма", "Плановая дата", "Фактическая дата", "Валюта", "Тип", "Статус", "Назначение"],
        [[view["counterparty_name"] or "—", row.amount, row.planned_date, row.actual_date,
          row.currency, "Приход" if row.direction == "income" else "Расход", statuses.get(row.status, row.status), row.purpose]
         for row, view in zip(selected, views, strict=True)], f"Реестр {on_date}")
    from io import BytesIO

    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    book = load_workbook(BytesIO(content))
    for index, row in enumerate(selected, 2):
        book.active.cell(index, 2).number_format = '#,##0.00########'
        for column in (3, 4):
            book.active.cell(index, column).number_format = 'DD.MM.YYYY'
        if row.status == "rejected":
            for cell in book.active[index]:
                cell.fill = PatternFill("solid", fgColor="FFF1F3")
    output = BytesIO()
    book.save(output)
    content = output.getvalue()
    audit(db, user, "calendar_export", on_date.isoformat(), "export", after={"rows": len(selected), "date_basis": date_basis})
    return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="payments-{on_date}.xlsx"', "Cache-Control": "private, no-store"})


@router.post("/payment-calendar")
def calendar_create(data: CalendarIn, db: DB, user: Actor):
    def operation():
        cashflow.validate_day(db, data.direction, data.planned_date, data.outside_payment_days, "planned_date")
        if data.request_id:
            check_request(db, user, data.request_id)
        if data.counterparty_id:
            check_client(db, user, data.counterparty_id)
        if data.document_id:
            doc = db.get(CommercialDocument, data.document_id)
            if not doc or doc.request_id != data.request_id or data.direction != "income":
                error("CALENDAR_SOURCE", "Документ не соответствует приходу по заявке")
        if data.supplier_order_id:
            require_permission(db, user, "procurement.write")
            order = db.get(SupplierOrder, data.supplier_order_id)
            if not order or data.direction != "expense":
                error("CALENDAR_SOURCE", "Заказ поставщику не соответствует расходу")
            for line in order.snapshot["lines"]:
                check_request(db, user, line["request_id"])
        responsible = data.responsible_id or user.id
        ensure_available(db, responsible)
        row = CalendarEntry(
            **data.model_dump(exclude={"idempotency_key", "responsible_id"}),
            responsible_id=responsible,
            author_id=user.id,
        )
        db.add(row)
        db.flush()
        audit(db, user, "calendar_entry", row.id, "created", after=serialize(row))
        return calendar_view(db, row)

    return idem(db, user, data.idempotency_key, "calendar.create", data.model_dump(mode="json"), operation)


class CalendarPatch(CalendarIn):
    version: int


@router.patch("/payment-calendar/{entry_id}")
def calendar_edit(entry_id: str, data: CalendarPatch, db: DB, user: Actor):
    def operation():
        row = lock(db, CalendarEntry, entry_id)
        calendar_visible(db, user, row)
        check_version(row, data.version)
        if row.status not in ("draft", "rejected"):
            error("CALENDAR_STATE", "Изменять можно только черновик или отклонённый платёж")
        cashflow.validate_day(db, data.direction, data.planned_date, data.outside_payment_days, "planned_date")
        # Source links are immutable; validate the changed counterparty and responsible person.
        if data.counterparty_id:
            check_client(db, user, data.counterparty_id)
        responsible = data.responsible_id or row.responsible_id
        ensure_available(db, responsible)
        before = serialize(row)
        for key in (
            "direction",
            "planned_date",
            "amount",
            "currency",
            "purpose",
            "counterparty_id",
            "payment_kind",
            "recurrence",
            "outside_payment_days",
        ):
            setattr(row, key, getattr(data, key))
        if (row.document_id and row.direction != "income") or (
            row.supplier_order_id and row.direction != "expense"
        ):
            error("CALENDAR_SOURCE", "Направление платежа не соответствует исходному документу")
        row.responsible_id, row.status, row.version = responsible, "draft", row.version + 1
        audit(db, user, "calendar_entry", row.id, "updated", before, serialize(row))
        return calendar_view(db, row)

    return idem(
        db, user, data.idempotency_key, f"calendar.edit:{entry_id}", data.model_dump(mode="json"), operation
    )


@router.post("/payment-calendar/{entry_id}/submit")
def calendar_submit(entry_id: str, data: VersionCommand, db: DB, user: Actor):
    def operation():
        row = lock(db, CalendarEntry, entry_id)
        calendar_visible(db, user, row)
        check_version(row, data.version)
        if row.status not in ("draft", "rejected"):
            error("CALENDAR_STATE", "На согласование можно отправить черновик или отклонённый платёж")
        cashflow.validate_day(db, row.direction, row.planned_date, row.outside_payment_days, "planned_date")
        cashflow.validate_invoice(db, row)
        row.status = "pending"
        row.version += 1
        submit_review(db, user, "calendar", row, f"Платёж: {row.purpose}", serialize(row), row.request_id)
        return calendar_view(db, row)

    return idem(
        db, user, data.idempotency_key, f"calendar.submit:{entry_id}", data.model_dump(mode="json"), operation
    )


class CalendarConfirm(VersionCommand):
    actual_date: date | None = None
    outside_payment_days: bool = False


@router.post("/payment-calendar/{entry_id}/confirm")
def calendar_confirm(entry_id: str, data: CalendarConfirm, db: DB, user: Actor):
    require_permission(db, user, "approvals.decide")

    def operation():
        row = lock(db, CalendarEntry, entry_id)
        calendar_visible(db, user, row)
        check_version(row, data.version)
        if row.status != "approved":
            error("CALENDAR_APPROVAL_REQUIRED", "Сначала согласуйте платёж")
        cashflow.confirm(db, user, row, data.actual_date, data.outside_payment_days or row.outside_payment_days)
        return calendar_view(db, row)

    return idem(
        db,
        user,
        data.idempotency_key,
        f"calendar.confirm:{entry_id}",
        data.model_dump(mode="json"),
        operation,
    )
